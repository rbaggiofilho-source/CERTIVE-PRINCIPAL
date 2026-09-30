#!/usr/bin/env python3
"""Gera as planilhas editaveis (.xlsx) do e-book "Antes de Fechar Negocio".

Uso:
    python3 build/build-planilhas.py            # gera em 03-materiais/planilhas/
    python3 build/build-planilhas.py --out DIR  # gera em outro diretorio

Publico: consumidor leigo que vai comprar carro usado.
Cores: protocolo-cautelar/02-design/tema.css (carbono #0F1318, ambar #F2A900).

Formulas compativeis com Excel 2010+ e LibreOffice (IF, AND, OR, N, MIN, MAX,
COUNT, COUNTIF, SUMIF, SUMPRODUCT, INDEX, MATCH, PMT, FIXED). Nada de
LET/LAMBDA/XLOOKUP/FILTER/IFS/MAXIFS.

Os helpers de estilo foram copiados (e adaptados) de
protocolo-cautelar/build/build-planilhas.py; nada e importado de la.

O modulo expoe build_comparativo(path, dados=...) e build_custo(path, dados=...)
para o verificador (build/verificar-planilhas.py) gerar cenarios de teste.
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent.parent
OUT_DEFAULT = BASE / "03-materiais" / "planilhas"

# ---------------------------------------------------------------- tema
CARBONO = "0F1318"
AMBAR = "F2A900"
AMBAR_ESCURO = "B87F00"
AMBAR_CLARO = "FFF4D6"
PAPEL2 = "F5F6F8"
LINHA = "DDE1E6"
TEXTO = "1A1F26"
TEXTO2 = "4E5864"
TEXTO3 = "7A8591"
VERDE = "2E9E5B"
AMARELO = "E0A100"
LARANJA = "E0661B"
VERMELHO = "C62828"
INC = "5A6472"
VERMELHO_CLARO = "FDECEA"
VERDE_CLARO = "E8F5EC"
AZUL_ENTRADA = "0000FF"
ENTRADA_FILL = "EEF3FF"

FONT = "Calibri"
MARCA = "ANTES DE FECHAR NEGÓCIO"


def f(size=11, bold=False, color=TEXTO, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


F_IN = Font(name=FONT, size=11, color=AZUL_ENTRADA)
F_IN_B = Font(name=FONT, size=11, color=AZUL_ENTRADA, bold=True)
F_FX = Font(name=FONT, size=11, color="000000")
F_FX_B = Font(name=FONT, size=11, color="000000", bold=True)
FILL = lambda c: PatternFill("solid", start_color=c, end_color=c)  # noqa: E731
THIN = Side(style="thin", color=LINHA)
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
ESQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
DIR = Alignment(horizontal="right", vertical="center")

MOEDA = 'R$ #,##0.00;-R$ #,##0.00;"-"'
MOEDA0 = 'R$ #,##0;-R$ #,##0;"-"'
PCT = '0.0%;-0.0%;"0,0%"'
KM = '#,##0'

AVISO_CAPA = ("Valores de exemplo fictícios. Esta planilha organiza informações para a SUA decisão: não é recomendação "
              "de compra nem conselho financeiro e não substitui a vistoria cautelar profissional, a vistoria oficial "
              "do órgão de trânsito nem a orientação de um advogado.")


# ---------------------------------------------------------------- helpers
def header_row(ws, row, labels, col0=1, height=30):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col0 + i, value=lab)
        c.font = f(11, True, "FFFFFF")
        c.fill = FILL(CARBONO)
        c.alignment = CENTRO
        c.border = BORDA
    ws.row_dimensions[row].height = height


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def entrada(c, value=None, bold=False, fmt=None, align=None):
    if value is not None:
        c.value = value
    c.font = F_IN_B if bold else F_IN
    c.fill = FILL(ENTRADA_FILL)
    c.border = BORDA
    if fmt:
        c.number_format = fmt
    c.alignment = align or CENTRO
    return c


def formula(c, value, bold=False, fmt=None, align=None):
    c.value = value
    c.font = F_FX_B if bold else F_FX
    c.border = BORDA
    if fmt:
        c.number_format = fmt
    c.alignment = align or CENTRO
    return c


def rotulo(c, value, bold=False, fill=None, size=10.5):
    c.value = value
    c.font = f(size, bold, TEXTO)
    if fill:
        c.fill = FILL(fill)
    c.border = BORDA
    c.alignment = ESQ
    return c


def dica(c, value):
    c.value = value
    c.font = f(8.5, False, TEXTO2, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    c.border = BORDA
    return c


def secao(ws, row, texto, ncols):
    for col in range(1, ncols + 1):
        cc = ws.cell(row=row, column=col)
        cc.fill = FILL(AMBAR_CLARO)
        cc.border = Border(top=Side(style="thin", color=AMBAR), bottom=Side(style="thin", color=AMBAR))
    c = ws.cell(row=row, column=1, value=texto)
    c.font = f(10.5, True, AMBAR_ESCURO)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20


def titulo_aba(ws, texto, ncols, sub=None):
    """Faixa carbono com a marca em ambar + titulo da aba (linhas 1-2)."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=MARCA + "  |  " + texto)
    c.font = f(14, True, AMBAR)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = FILL(CARBONO)
    if sub:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        c = ws.cell(row=2, column=1, value=sub)
        c.font = f(9.5, False, TEXTO2, italic=True)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        ws.row_dimensions[2].height = 32


def impressao(ws, paisagem=False, area=None, titulos=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if paisagem else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = "Antes de Fechar Negócio"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Página &P de &N"
    ws.oddFooter.right.size = 8
    if area:
        ws.print_area = area
    if titulos:
        ws.print_title_rows = titulos


def nota(ws, cell_range, texto, cor=AMBAR_CLARO, altura=None, cor_texto=TEXTO):
    ws.merge_cells(cell_range)
    first = cell_range.split(":")[0]
    c = ws[first]
    c.value = texto
    c.font = f(9.5, False, cor_texto, italic=True)
    c.fill = FILL(cor)
    c.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
    if altura:
        ws.row_dimensions[c.row].height = altura


def dv_lista(ws, itens, rng, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(itens) + '"', allow_blank=True, showDropDown=False)
    dv.error = "Escolha um valor da lista."
    dv.errorTitle = "Valor inválido"
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def dv_numero(ws, rng, minimo=0, maximo=None, inteiro=False, msg=None):
    if maximo is None:
        dv = DataValidation(type="whole" if inteiro else "decimal", operator="greaterThanOrEqual",
                            formula1=str(minimo), allow_blank=True)
    else:
        dv = DataValidation(type="whole" if inteiro else "decimal", operator="between",
                            formula1=str(minimo), formula2=str(maximo), allow_blank=True)
    dv.error = msg or "Digite um número válido."
    dv.errorTitle = "Valor inválido"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)
    return dv


def capa(wb, titulo, instrucoes, extra=None):
    ws = wb.active
    ws.title = "Capa"
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 3, "B": 22, "C": 22, "D": 22, "E": 22, "F": 3})
    for r in range(1, 45):
        for col in range(1, 7):
            ws.cell(row=r, column=col).fill = FILL("FFFFFF")
    for r in range(1, 8):
        for col in range(1, 7):
            ws.cell(row=r, column=col).fill = FILL(CARBONO)
    ws.merge_cells("B2:E2")
    c = ws["B2"]
    c.value = "ANTES DE FECHAR NEGÓCIO"
    c.font = f(22, True, AMBAR)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 34
    ws.merge_cells("B3:E3")
    c = ws["B3"]
    c.value = "O guia do vistoriador para quem vai comprar carro usado  ·  Protocolo Cautelar"
    c.font = f(10, False, "C9CED6", italic=True)
    ws.merge_cells("B5:E6")
    c = ws["B5"]
    c.value = titulo
    c.font = f(18, True, "FFFFFF")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 26
    for col in range(1, 7):
        ws.cell(row=8, column=col).fill = FILL(AMBAR)
    ws.row_dimensions[8].height = 5

    r = 10
    ws.cell(row=r, column=2, value="COMO USAR (passo a passo)").font = f(12, True, AMBAR_ESCURO)
    r += 1
    for i, linha in enumerate(instrucoes, 1):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
        c = ws.cell(row=r, column=2, value=f"{i}. {linha}")
        c.font = f(11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 18 if len(linha) <= 85 else (32 if len(linha) <= 170 else 46)
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="LEGENDA DE CORES").font = f(12, True, AMBAR_ESCURO)
    r += 1
    entrada(ws.cell(row=r, column=2), "58.900,00")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws.cell(row=r, column=3, value="AZUL (fundo azul-claro) = você PREENCHE: digite ou escolha na lista.").font = f(11)
    r += 1
    formula(ws.cell(row=r, column=2), "=2*617", fmt="#,##0.00")
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws.cell(row=r, column=3, value="PRETO = CALCULADO automaticamente. Não digite por cima.").font = f(11)
    r += 1
    c = ws.cell(row=r, column=2, value="Alerta")
    c.font = f(11, True, VERMELHO)
    c.fill = FILL(VERMELHO_CLARO)
    c.border = BORDA
    c.alignment = CENTRO
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    ws.cell(row=r, column=3, value="VERMELHO = sinal de alerta: pare, confira e, se preciso, chame um profissional.").font = f(11)
    r += 1
    if extra:
        r += 1
        for linha in extra:
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
            c = ws.cell(row=r, column=2, value=linha)
            c.font = f(9.5, False, TEXTO2, italic=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 if len(linha) <= 95 else (30 if len(linha) <= 190 else 44)
            r += 1
    r += 1
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    c = ws.cell(row=r, column=2, value=AVISO_CAPA)
    c.font = f(8.5, False, TEXTO3, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 36
    impressao(ws, paisagem=False, area=f"A1:F{r + 1}")
    return ws


def cor_lista(ws, rng, pares):
    """pares: [(texto, cor_fundo, cor_fonte)] -> formatacao condicional por igualdade."""
    for txt, cor, fc in pares:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=FILL(cor),
                                                      font=Font(name=FONT, bold=True, color=fc)))


# ================================================================ 1. COMPARATIVO
CARROS = ["Carro 1", "Carro 2", "Carro 3", "Carro 4"]
COLS = ["B", "C", "D", "E"]
CONSULTA = ["OK", "Pendente", "Alerta"]
LAUDO = ["Aprovado", "Aprovado com apontamentos", "Reprovado", "Inconclusivo", "Não feito"]
NAO_REC = "Não recomendado sem avaliação adicional"
ALERTA_PRECO = "Preço muito abaixo da referência: redobre a atenção"

# Exemplo ficticio (3 carros preenchidos, o 4o em branco).
EXEMPLO_COMP = {
    "apelido": ["Hatch A 1.0 Flex", "Sedã B 1.6 Flex", "Hatch C 1.3 Flex", None],
    "fonte": ["Anúncio on-line (cole o link)", "Anúncio on-line (cole o link)", "Loja do bairro", None],
    "vendedor": ["Particular", "Particular", "Loja", None],
    "pp": [58900, 45900, 64900, None],
    "pn": [56500, None, 62900, None],
    "fab": [2020, 2019, 2021, None],
    "mod": [2021, 2019, 2022, None],
    "km": [62000, 98000, 41000, None],
    "cor": ["Prata", "Branco", "Cinza", None],
    "cambio": ["Manual", "Automático", "Automático", None],
    "comb": ["Flex", "Flex", "Flex", None],
    "consumo": [12.5, 10.8, 12.0, None],
    "ref": [59000, 61000, 63500, None],
    "donos": [1, 3, 1, None],
    "rev": ["S", "N", "S", None],
    "c1": ["OK", "Pendente", "OK", None],
    "c2": ["OK", "OK", "OK", None],
    "c3": ["OK", "OK", "OK", None],
    "c4": ["OK", "OK", "Pendente", None],
    "c5": ["OK", "Alerta", "OK", None],
    "laudo": ["Aprovado com apontamentos", "Reprovado", "Aprovado", None],
    "rep": [1200, 6500, 400, None],
    "transf": [1000, 1000, 1000, None],
    "ct2": [None, None, None, None],
    "s6": [7, 9, 8, None],
}

# (tipo, chave, rotulo, formato, dica)
#  tipo: sec | txt | num | lista:<nome> | fx | nota (entrada 0-10) | score
LINHAS_COMP = [
    ("sec", None, "1. O ANÚNCIO", None, None),
    ("txt", "apelido", "Carro (marca / modelo / versão)", None, "Um nome curto para você reconhecer o carro."),
    ("txt", "fonte", "Onde viu (link ou fonte do anúncio)", None, "Cole o link ou anote onde viu. Tire print do anúncio."),
    ("lista:vendedor", "vendedor", "Vendedor", None, "Particular ou loja. Na loja vale também o Código de Defesa do Consumidor."),
    ("num", "pp", "Preço pedido (R$)", MOEDA, "Preço do anúncio. Sem preço pedido, a coluna fica em branco."),
    ("num", "pn", "Preço negociado (R$)", MOEDA, "Se já negociou, anote aqui. Em branco = usa o preço pedido."),
    ("fx", "pc", "Preço considerado (R$)", MOEDA, "Negociado, se houver; senão, o pedido."),
    ("sec", None, "2. O CARRO", None, None),
    ("num", "fab", "Ano de fabricação", "0", "Confira no documento do veículo, não só no anúncio."),
    ("num", "mod", "Ano do modelo", "0", ""),
    ("num", "km", "Quilometragem (km)", KM, "O que aparece no painel, conferido na visita."),
    ("fx", "kmano", "Km por ano (aprox.)", KM, "Km ÷ anos de uso (ano de referência − fabricação). Muito baixo ou muito alto merece pergunta."),
    ("txt", "cor", "Cor", None, ""),
    ("lista:cambio", "cambio", "Câmbio", None, ""),
    ("lista:comb", "comb", "Combustível", None, ""),
    ("num", "consumo", "Consumo informado (km/l)", "0.0", "Informado pelo vendedor ou pelo fabricante. Use na planilha de custo."),
    ("sec", None, "3. PREÇO x MERCADO", None, None),
    ("num", "ref", "Valor de referência de mercado (R$)", MOEDA,
     "Valor que você consultou numa tabela de referência de preços de usados (mesmo ano/modelo/versão)."),
    ("fx", "dif", "Diferença do preço vs referência (%)", PCT, "Negativo = mais barato que a referência."),
    ("fx", "alerta", "Alerta de preço", None, "Muito abaixo da referência é um sinal clássico de golpe ou de problema escondido."),
    ("sec", None, "4. HISTÓRICO", None, None),
    ("num", "donos", "Nº de donos (1 = único dono)", "0", "Pergunte e confira no documento/histórico quando possível."),
    ("lista:sn", "rev", "Revisões comprovadas? (S/N)", None, "S = tem notas fiscais ou manual carimbado para conferir."),
    ("sec", None, "5. CONSULTAS FEITAS (OK / Pendente / Alerta)", None, None),
    ("lista:consulta", "c1", "Débitos (IPVA, licenciamento, multas)", None, "OK = consultou e está tudo certo. Alerta = achou problema."),
    ("lista:consulta", "c2", "Restrições (judicial, administrativa, furto/roubo)", None, ""),
    ("lista:consulta", "c3", "Gravame (financiamento/alienação em aberto)", None, ""),
    ("lista:consulta", "c4", "Recall pendente", None, "Consulte pelo chassi no canal oficial do fabricante."),
    ("lista:consulta", "c5", "Histórico de leilão / sinistro", None, ""),
    ("fx", "cres", "Resumo das consultas", None, "Qualquer Alerta tira o carro da disputa até ser esclarecido."),
    ("sec", None, "6. LAUDO CAUTELAR E REPAROS", None, None),
    ("lista:laudo", "laudo", "Resultado do laudo cautelar", None,
     "Classificação do laudo feito por vistoriador. Reprovado/Inconclusivo tira o carro da disputa."),
    ("num", "rep", "Custo estimado de reparos (R$)", MOEDA, "Orçamento de oficina de confiança para o que o laudo apontou."),
    ("sec", None, "7. CUSTO TOTAL DA COMPRA", None, None),
    ("num", "transf", "Transferência e taxas (R$)", MOEDA, "Estimativa. Valores reais: órgão de trânsito do seu estado (veja a planilha 2)."),
    ("fx", "resv", "Reserva de imprevistos (R$)", MOEDA, "Preço considerado × % de reserva (parâmetro abaixo)."),
    ("num", "ct2", "Custo total da planilha 2 (opcional)", MOEDA, "Se fez a planilha 2 para este carro, cole o total aqui: ele passa a valer."),
    ("fx", "ct", "CUSTO TOTAL DA COMPRA (R$)", MOEDA, "Planilha 2, se preenchido; senão: preço + transferência + reparos + reserva."),
    ("sec", None, "8. PONTUAÇÃO (0 a 10) — ajuste os pesos na coluna F", None, None),
    ("score", "s1", "Custo total (menor = melhor)", "0.0", "10 para o menor custo total; os demais, proporcional."),
    ("score", "s2", "Km por ano", "0.0", "10 até 10.000 km/ano; cai 1 ponto a cada 2.000 km/ano a mais (0 em 30.000)."),
    ("score", "s3", "Histórico (donos + revisões)", "0.0", "Donos: 1 = 5 pts, 2 = 3, 3 ou mais = 1. Revisões comprovadas = +5."),
    ("score", "s4", "Laudo cautelar", "0.0", "Aprovado 10 · Com apontamentos 6 · Não feito/em branco 2 · Reprovado/Inconclusivo 0."),
    ("score", "s5", "Consultas", "0.0", "2 pontos por consulta OK; qualquer Alerta = 0."),
    ("nota", "s6", "Sua nota pessoal (conforto, espaço, test drive)", "0", "Você dá a nota de 0 a 10."),
    ("fx", "pont", "PONTUAÇÃO FINAL (0 a 10)", "0.0", "Média ponderada pelas notas e pesos acima."),
    ("fx", "restr", "Situação", None, "Laudo Reprovado/Inconclusivo ou Alerta nas consultas = não recomendado sem avaliação adicional."),
    ("fx", "valida", "Pontuação que vale para o ranking", "0.0", "Carros com restrição ficam de fora da escolha."),
]
PESOS = {"s1": 30, "s2": 15, "s3": 15, "s4": 20, "s5": 10, "s6": 10}
LISTAS = {
    "vendedor": ["Particular", "Loja"],
    "cambio": ["Manual", "Automático", "Automatizado", "CVT", "Outro"],
    "comb": ["Flex", "Gasolina", "Etanol", "Diesel", "Híbrido", "Elétrico", "GNV"],
    "sn": ["S", "N"],
    "consulta": CONSULTA,
    "laudo": LAUDO,
}
HDR_COMP = 4


def _layout_comp():
    R = {}
    r = HDR_COMP + 1
    for tipo, chave, *_ in LINHAS_COMP:
        if chave:
            R[chave] = r
        r += 1
    R["_fim"] = r - 1
    R["melhor"] = r + 1
    R["sit_melhor"] = r + 2
    R["ressalva"] = r + 3
    R["par_titulo"] = r + 6
    R["anoref"] = r + 7
    R["lim_baixo"] = r + 8
    R["lim_alto"] = r + 9
    R["reserva"] = r + 10
    return R


def _fx_comp(chave, L, R):
    g = lambda k: f"{L}{R[k]}"  # noqa: E731
    vazio = f'{g("pp")}=""'
    cons = f'{L}{R["c1"]}:{L}{R["c5"]}'
    if chave == "pc":
        return f'=IF({vazio},"",IF({g("pn")}<>"",{g("pn")},{g("pp")}))'
    if chave == "kmano":
        return f'=IF(OR({vazio},{g("km")}="",{g("fab")}=""),"",{g("km")}/MAX(1,$B${R["anoref"]}-{g("fab")}))'
    if chave == "dif":
        return f'=IF(OR({vazio},N({g("ref")})=0),"",{g("pc")}/{g("ref")}-1)'
    if chave == "alerta":
        return (f'=IF({g("dif")}="","",IF(ROUND({g("dif")},4)<$B${R["lim_baixo"]},"{ALERTA_PRECO}",'
                f'IF(ROUND({g("dif")},4)>$B${R["lim_alto"]},"Acima da referência: argumento para negociar",'
                f'"Dentro da faixa da referência")))')
    if chave == "cres":
        return (f'=IF({vazio},"",IF(COUNTIF({cons},"Alerta")>0,"ALERTA em "&COUNTIF({cons},"Alerta")&" consulta(s)",'
                f'IF(COUNTIF({cons},"OK")=5,"Todas OK","Faltam "&(5-COUNTIF({cons},"OK"))&" consulta(s)")))')
    if chave == "resv":
        return f'=IF({vazio},"",{g("pc")}*$B${R["reserva"]})'
    if chave == "ct":
        return (f'=IF({vazio},"",IF(N({g("ct2")})>0,{g("ct2")},'
                f'{g("pc")}+N({g("transf")})+N({g("rep")})+{g("resv")}))')
    if chave == "s1":
        return f'=IF({vazio},"",IF(N({g("ct")})=0,0,10*MIN($B${R["ct"]}:$E${R["ct"]})/{g("ct")}))'
    if chave == "s2":
        return f'=IF({vazio},"",IF({g("kmano")}="",0,MAX(0,MIN(10,10-({g("kmano")}-10000)/2000))))'
    if chave == "s3":
        return (f'=IF({vazio},"",IF({g("donos")}="",0,IF({g("donos")}<=1,5,IF({g("donos")}=2,3,1)))'
                f'+IF({g("rev")}="S",5,0))')
    if chave == "s4":
        return (f'=IF({vazio},"",IF({g("laudo")}="Aprovado",10,IF({g("laudo")}="Aprovado com apontamentos",6,'
                f'IF(OR({g("laudo")}="Não feito",{g("laudo")}=""),2,0))))')
    if chave == "s5":
        return f'=IF({vazio},"",IF(COUNTIF({cons},"Alerta")>0,0,2*COUNTIF({cons},"OK")))'
    if chave == "pont":
        s1, s6 = R["s1"], R["s6"]
        return (f'=IF(OR({vazio},SUM($F${s1}:$F${s6})=0),"",'
                f'SUMPRODUCT({L}{s1}:{L}{s6},$F${s1}:$F${s6})/SUM($F${s1}:$F${s6}))')
    if chave == "restr":
        return (f'=IF({vazio},"",IF(OR({g("laudo")}="Reprovado",{g("laudo")}="Inconclusivo",COUNTIF({cons},"Alerta")>0),'
                f'"{NAO_REC}",IF(OR({g("laudo")}="",{g("laudo")}="Não feito"),"Sem laudo cautelar: faça antes de decidir",'
                f'IF(COUNTIF({cons},"OK")<5,"Consultas incompletas: termine antes de pagar","Sem restrições registradas"))))')
    if chave == "valida":
        return f'=IF({vazio},"",IF({g("restr")}="{NAO_REC}","Fora",{g("pont")}))'
    raise KeyError(chave)


def build_comparativo(path, dados=None):
    dados = dict(EXEMPLO_COMP if dados is None else dados)
    R = _layout_comp()
    wb = Workbook()
    capa(wb, "Comparativo de veículos", [
        "Abra a aba Comparar. Cada coluna (Carro 1 a 4) é um carro que você está considerando.",
        "Preencha só as células AZUIS. As pretas se calculam sozinhas. Comece pelo preço pedido: sem ele, a coluna fica vazia.",
        "Nas consultas, marque OK (consultou e está certo), Pendente (ainda não consultou) ou Alerta (achou problema).",
        "Informe o resultado do laudo cautelar, se já tiver. Laudo Reprovado/Inconclusivo ou qualquer Alerta tira o carro da escolha.",
        "Se quiser, ajuste os pesos da pontuação (coluna F) conforme o que é mais importante para você.",
        "Na aba Anúncio x Realidade, compare o que o anúncio promete com o que você viu na visita.",
        "Apague os exemplos (células azuis) antes de usar com os seus carros.",
    ], extra=[
        "Preço muito abaixo da referência de mercado (mais de 15% abaixo, ajustável) aparece como alerta: é um dos sinais mais comuns de golpe ou de problema escondido.",
        "A \"melhor opção\" é só a soma dos critérios e pesos que você escolheu. Ela não enxerga o que não foi informado.",
    ])

    ws = wb.create_sheet("Comparar")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 42, "B": 22, "C": 22, "D": 22, "E": 22, "F": 8, "G": 46})
    titulo_aba(ws, "COMPARATIVO DE VEÍCULOS (ATÉ 4)", 7,
               "Preencha as células azuis. Coluna sem preço pedido fica em branco. A pontuação é um apoio para organizar "
               "a decisão, não uma recomendação de compra.")
    header_row(ws, HDR_COMP, ["Critério"] + CARROS + ["Peso", "Dica"])

    for tipo, chave, lab, fmt, dic in LINHAS_COMP:
        if tipo == "sec":
            continue
        r = R[chave]
        destaque = chave in ("ct", "pont", "restr", "valida")
        rotulo(ws.cell(row=r, column=1), lab, bold=destaque, fill=AMBAR_CLARO if destaque else None)
        dica(ws.cell(row=r, column=7), dic)
        ws.cell(row=r, column=6).border = BORDA
        ws.cell(row=r, column=6).fill = FILL(PAPEL2)
        for i, L in enumerate(COLS):
            c = ws[f"{L}{r}"]
            if tipo in ("txt", "num", "nota") or tipo.startswith("lista"):
                v = dados.get(chave, [None] * 4)[i]
                entrada(c, v, fmt=fmt)
            else:
                formula(c, _fx_comp(chave, L, R), bold=destaque, fmt=fmt)
            if tipo.startswith("lista"):
                pass
        if tipo.startswith("lista"):
            nome = tipo.split(":")[1]
            dv_lista(ws, LISTAS[nome], f"B{r}:E{r}")
        if tipo in ("score", "nota"):
            entrada(ws.cell(row=r, column=6), PESOS[chave], fmt="0")
        if tipo == "nota":
            dv_numero(ws, f"B{r}:E{r}", 0, 10, msg="Nota de 0 a 10.")
        ws.row_dimensions[r].height = 30 if len(dic or "") > 60 or len(lab) > 44 else 20
    # secoes
    r = HDR_COMP + 1
    for tipo, chave, lab, *_ in LINHAS_COMP:
        if tipo == "sec":
            secao(ws, r, lab, 7)
        r += 1
    # validacoes numericas
    for k in ("pp", "pn", "ref", "rep", "transf", "ct2", "km"):
        dv_numero(ws, f"B{R[k]}:E{R[k]}", 0)
    for k in ("fab", "mod"):
        dv_numero(ws, f"B{R[k]}:E{R[k]}", 1950, 2100, inteiro=True, msg="Ano com 4 dígitos (ex.: 2020).")
    dv_numero(ws, f"B{R['donos']}:E{R['donos']}", 1, 99, inteiro=True, msg="Número inteiro (1 = único dono).")
    dv_numero(ws, f"F{R['s1']}:F{R['s6']}", 0, 100, msg="Peso de 0 a 100.")

    # formatacao condicional
    cor_lista(ws, f"B{R['c1']}:E{R['c5']}", [("OK", VERDE, "FFFFFF"), ("Pendente", AMARELO, TEXTO),
                                             ("Alerta", VERMELHO, "FFFFFF")])
    cor_lista(ws, f"B{R['laudo']}:E{R['laudo']}", [("Aprovado", VERDE, "FFFFFF"),
                                                   ("Aprovado com apontamentos", AMARELO, TEXTO),
                                                   ("Reprovado", VERMELHO, "FFFFFF"), ("Inconclusivo", INC, "FFFFFF"),
                                                   ("Não feito", LINHA, TEXTO)])
    ra = R["alerta"]
    ws.conditional_formatting.add(f"B{ra}:E{ra}", FormulaRule(formula=[f'LEFT(B{ra},10)="Preço muit"'],
                                                              fill=FILL(VERMELHO_CLARO), font=Font(name=FONT, bold=True, color=VERMELHO)))
    rd = R["dif"]
    ws.conditional_formatting.add(f"B{rd}:E{rd}", FormulaRule(formula=[f'AND(ISNUMBER(B{rd}),ROUND(B{rd},4)<$B${R["lim_baixo"]})'],
                                                              font=Font(name=FONT, bold=True, color=VERMELHO)))
    rc = R["cres"]
    ws.conditional_formatting.add(f"B{rc}:E{rc}", FormulaRule(formula=[f'LEFT(B{rc},6)="ALERTA"'],
                                                              fill=FILL(VERMELHO_CLARO), font=Font(name=FONT, bold=True, color=VERMELHO)))
    rr_ = R["restr"]
    ws.conditional_formatting.add(f"B{rr_}:E{rr_}", FormulaRule(formula=[f'B{rr_}="{NAO_REC}"'],
                                                                fill=FILL(VERMELHO), font=Font(name=FONT, bold=True, color="FFFFFF")))
    ws.conditional_formatting.add(f"B{rr_}:E{rr_}", FormulaRule(formula=[f'B{rr_}="Sem restrições registradas"'],
                                                                fill=FILL(VERDE_CLARO), font=Font(name=FONT, bold=True, color=VERDE)))
    ws.row_dimensions[rr_].height = 44
    ws.row_dimensions[R["alerta"]].height = 44
    ws.row_dimensions[R["cres"]].height = 30

    # resultado
    rv, ap = R["valida"], R["apelido"]
    faixa = f"B{rv}:E{rv}"
    pos = f"MATCH(MAX({faixa}),{faixa},0)"
    rm = R["melhor"]
    rotulo(ws.cell(row=rm, column=1), "MELHOR OPÇÃO PELOS SEUS CRITÉRIOS", bold=True, fill=AMBAR)
    ws.merge_cells(start_row=rm, start_column=2, end_row=rm, end_column=7)
    formula(ws.cell(row=rm, column=2),
            f'=IF(COUNT({faixa})=0,IF(COUNTA(B{R["pp"]}:E{R["pp"]})=0,"Preencha os dados dos carros.",'
            f'"Nenhum carro sem restrição: busque avaliação adicional antes de decidir."),'
            f'INDEX($B${HDR_COMP}:$E${HDR_COMP},{pos})&IF(INDEX(B{ap}:E{ap},{pos})="",""," – "&INDEX(B{ap}:E{ap},{pos}))'
            f'&"  ("&FIXED(MAX({faixa}),1)&" de 10)")', bold=True, align=ESQ)
    ws.cell(row=rm, column=2).font = f(12, True, CARBONO)
    ws.cell(row=rm, column=2).fill = FILL(AMBAR_CLARO)
    ws.row_dimensions[rm].height = 26
    rs = R["sit_melhor"]
    rotulo(ws.cell(row=rs, column=1), "Situação da melhor opção", bold=True)
    ws.merge_cells(start_row=rs, start_column=2, end_row=rs, end_column=7)
    formula(ws.cell(row=rs, column=2),
            f'=IF(COUNT({faixa})=0,"",INDEX(B{R["restr"]}:E{R["restr"]},{pos})'
            f'&IF(LEFT(INDEX(B{R["alerta"]}:E{R["alerta"]},{pos}),10)="Preço muit"," · Atenção: preço muito abaixo da referência.",""))',
            align=ESQ)
    nota(ws, f"A{R['ressalva']}:G{R['ressalva'] + 1}",
         "Ressalva: a \"melhor opção\" resulta apenas dos dados e pesos que você informou. Carro com laudo Reprovado ou "
         "Inconclusivo, ou com Alerta em qualquer consulta, fica marcado \"Não recomendado sem avaliação adicional\" e não "
         "pode ser a melhor opção. Nenhuma pontuação substitui a vistoria cautelar profissional, a vistoria de transferência "
         "e a conferência de documentos.", altura=30)
    ws.row_dimensions[R["ressalva"] + 1].height = 30

    # parametros
    rp = R["par_titulo"]
    secao(ws, rp, "PARÂMETROS (você pode mudar)", 7)
    params = [
        ("anoref", "Ano de referência (para km por ano)", 2026, "0", "Normalmente o ano atual."),
        ("lim_baixo", "Alerta quando o preço estiver abaixo da referência em", -0.15, PCT,
         "Padrão −15%. Abaixo disso aparece o alerta de preço."),
        ("lim_alto", "Aviso de preço acima da referência a partir de", 0.10, PCT, "Padrão +10%. Serve de argumento para negociar."),
        ("reserva", "Reserva para imprevistos (% do preço)", 0.10, PCT, "Padrão 10%. Dinheiro separado para o que aparecer."),
    ]
    for k, lab, v, fmt, dic in params:
        r = R[k]
        rotulo(ws.cell(row=r, column=1), lab)
        entrada(ws.cell(row=r, column=2), v, bold=True, fmt=fmt)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=7)
        dica(ws.cell(row=r, column=3), dic)
        ws.row_dimensions[r].height = 20
    dv_numero(ws, f"B{R['anoref']}", 1990, 2100, inteiro=True)
    dv_numero(ws, f"B{R['lim_baixo']}", -1, 0, msg="Digite um percentual negativo, por exemplo −15%.")
    dv_numero(ws, f"B{R['lim_alto']}", 0, 1, msg="Digite um percentual, por exemplo 10%.")
    dv_numero(ws, f"B{R['reserva']}", 0, 1, msg="Digite um percentual, por exemplo 10%.")
    ws["A" + str(R["lim_baixo"])].comment = Comment(
        "Limite do alerta de preço. Padrão: −15% em relação ao valor de referência de mercado.", "Antes de Fechar Negócio")

    ws.freeze_panes = f"B{HDR_COMP + 1}"
    impressao(ws, paisagem=True, area=f"A1:G{R['reserva']}", titulos=f"{HDR_COMP}:{HDR_COMP}")

    _anuncio_realidade(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(path)
    return R


def _anuncio_realidade(wb):
    ws = wb.create_sheet("Anúncio x Realidade")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 30, "B": 30, "C": 32, "D": 16, "E": 34})
    titulo_aba(ws, "ANÚNCIO x REALIDADE", 5,
               "Anote o que o anúncio (ou o vendedor) afirma e, ao lado, o que você viu ou conferiu. "
               "Para outro carro, duplique esta aba (botão direito na guia > Mover ou copiar > Criar uma cópia).")
    rotulo(ws["A4"], "Carro analisado", bold=True)
    ws.merge_cells("B4:C4")
    entrada(ws["B4"], "Carro 1 – Hatch A 1.0 Flex – placa ABC1D23", align=ESQ)
    rotulo(ws["A5"], "Link / data do anúncio", bold=True)
    ws.merge_cells("B5:C5")
    entrada(ws["B5"], "(cole o link) – 10/09/2026", align=ESQ)
    header_row(ws, 7, ["O que conferir", "O anúncio diz", "O que eu vi / conferi", "Situação", "Observação"])
    linhas = [
        ("Preço", "R$ 58.900", "R$ 58.900 (negociado R$ 56.500)", "Confere", ""),
        ("Ano de fabricação / modelo", "2020/2021", "2020/2021 no documento", "Confere", ""),
        ("Quilometragem", "55.000 km", "62.000 km no painel", "Diverge", "Perguntar o motivo da diferença"),
        ("Único dono", "Sim", "Vendedor não é o dono do documento", "Diverge", "Só pagar ao titular do documento"),
        ("Revisões em dia", "Todas feitas", "Notas fiscais de 3 revisões", "Confere", ""),
        ("Sem batidas / nunca bateu", "Nunca bateu", "", "Não informado", "Laudo cautelar vai responder"),
        ("Pneus", "Novos", "Meia-vida", "Diverge", "Argumento para negociar"),
        ("Chave reserva e manual", "", "", "", ""),
        ("Documentação e débitos em dia", "", "", "", ""),
        ("Câmbio / motor / ar-condicionado", "", "", "", ""),
        ("Opcionais (multimídia, sensores...)", "", "", "", ""),
        ("Fotos batem com o carro visto", "", "", "", ""),
        ("", "", "", "", ""),
        ("", "", "", "", ""),
    ]
    L1 = 8
    for k, vals in enumerate(linhas):
        r = L1 + k
        for j, v in enumerate(vals, 1):
            c = entrada(ws.cell(row=r, column=j), v if v != "" else None, align=WRAP if j != 4 else CENTRO)
            if j == 1:
                c.font = F_IN_B
        ws.row_dimensions[r].height = 30
    L2 = L1 + len(linhas) - 1
    dv_lista(ws, ["Confere", "Diverge", "Não informado"], f"D{L1}:D{L2}")
    cor_lista(ws, f"D{L1}:D{L2}", [("Confere", VERDE, "FFFFFF"), ("Diverge", VERMELHO, "FFFFFF"),
                                   ("Não informado", LINHA, TEXTO)])
    rs = L2 + 2
    header_row(ws, rs, ["Resumo", "Quantidade"], height=22)
    D = f"$D${L1}:$D${L2}"
    res = [("Confere", f'=COUNTIF({D},"Confere")'), ("Diverge", f'=COUNTIF({D},"Diverge")'),
           ("Não informado", f'=COUNTIF({D},"Não informado")'), ("Itens avaliados", f"=SUM(B{rs + 1}:B{rs + 3})")]
    for k, (lab, fx) in enumerate(res, 1):
        rotulo(ws.cell(row=rs + k, column=1), lab)
        formula(ws.cell(row=rs + k, column=2), fx, bold=True)
    rr = rs + 5
    rotulo(ws.cell(row=rr, column=1), "Conclusão", bold=True, fill=AMBAR_CLARO)
    ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=5)
    formula(ws.cell(row=rr, column=2),
            f'=IF(B{rs + 4}=0,"Preencha a coluna Situação.",IF(B{rs + 2}>0,B{rs + 2}&" divergência(s) entre o anúncio e o carro. '
            f'Peça explicação por escrito, renegocie ou desista.",IF(B{rs + 3}>0,"Sem divergências, mas "&B{rs + 3}'
            f'&" item(ns) sem informação: confira antes de pagar.","O que foi anunciado confere com o que você viu.")))',
            bold=True, align=ESQ)
    ws.conditional_formatting.add(f"B{rr}", FormulaRule(formula=[f"$B${rs + 2}>0"], fill=FILL(VERMELHO_CLARO),
                                                        font=Font(name=FONT, bold=True, color=VERMELHO)))
    ws.row_dimensions[rr].height = 34
    nota(ws, f"A{rr + 2}:E{rr + 2}",
         "Dica: tire print do anúncio antes da visita. Anúncios podem ser editados ou apagados depois.", altura=22)
    ws.freeze_panes = "A8"
    impressao(ws, paisagem=False, area=f"A1:E{rr + 2}", titulos="7:7")
    return {"L1": L1, "L2": L2, "rs": rs}


# ================================================================ 2. CUSTO TOTAL
QUEM = ["Vendedor", "Comprador"]
VERIF = "Consulte o órgão de trânsito do seu estado. [VERIFICAR]"
EXEMPLO_CUSTO = {
    "preco": 56500,
    "vist_transf": None, "taxa_orgao": None, "firma": None, "despachante": None,
    "ipva_deb": 0, "ipva_deb_q": "Vendedor",
    "lic_deb": 0, "lic_deb_q": "Vendedor",
    "multas": 195.23, "multas_q": "Vendedor",
    "cautelar": 350,
    "seg1": 480, "seg_anual": 3600,
    "oleo": 600, "correia": 0, "pneus": 1600, "freios": 0, "reparos": 1200,
    "res_pct": 0.10,
    "fin": "Sim", "entrada": 20000, "n": 48, "taxa": 0.019,
    "ipva_aliq": 0.04, "lic_anual": None, "manut": 150,
    "km_mes": 1000, "comb_preco": 6.0, "consumo": 12.5,
    # orcamento mensal
    "renda": 6500, "limite": 0.20, "estac": 150, "outros": 80,
}

# (tipo, chave, rotulo, formato, observacao)   tipo: sec | in | quem | fx | lista:fin
LINHAS_CUSTO = [
    ("sec", None, "1. PREÇO", None, None),
    ("in", "preco", "Preço negociado (R$)", MOEDA, "O valor que você vai pagar pelo carro."),
    ("sec", None, "2. TRANSFERÊNCIA (valores variam por estado)", None, None),
    ("in", "vist_transf", "Vistoria de transferência (R$)", MOEDA, VERIF),
    ("in", "taxa_orgao", "Taxa do órgão estadual de trânsito (R$)", MOEDA, VERIF),
    ("in", "firma", "Reconhecimento de firma / assinatura, se aplicável (R$)", MOEDA,
     "Depende da forma de transferência do seu estado (papel ou eletrônica). [VERIFICAR]"),
    ("in", "despachante", "Despachante (opcional) (R$)", MOEDA, "Só se você optar por contratar."),
    ("fx", "sub_transf", "Subtotal transferência", MOEDA, None),
    ("sec", None, "3. DÉBITOS PENDENTES DO VEÍCULO (combine quem paga)", None, None),
    ("quem", "ipva_deb", "IPVA em aberto (R$)", MOEDA, "Consulte antes. Se o vendedor paga, exija quitação antes da transferência ou desconto no preço."),
    ("quem", "lic_deb", "Licenciamento em aberto (R$)", MOEDA, ""),
    ("quem", "multas", "Multas em aberto (R$)", MOEDA, ""),
    ("fx", "sub_deb", "Débitos que ficam com você (comprador)", MOEDA, "Soma só o que está marcado \"Comprador\"."),
    ("sec", None, "4. PROTEÇÃO E DEIXAR O CARRO EM DIA", None, None),
    ("in", "cautelar", "Vistoria cautelar (R$)", MOEDA, "Feita por vistoriador antes de pagar."),
    ("in", "seg1", "Seguro: 1ª parcela (R$)", MOEDA, "Contrate para valer a partir da retirada do carro."),
    ("in", "seg_anual", "Seguro: valor anual (R$)", MOEDA, "Usado no custo do 1º ano e no orçamento mensal."),
    ("in", "oleo", "Revisão pós-compra: óleo e filtros (R$)", MOEDA, "Orçamento de oficina de confiança."),
    ("in", "correia", "Correias, se a km pedir (R$)", MOEDA, "Siga o manual do fabricante para o intervalo."),
    ("in", "pneus", "Pneus, se o laudo apontar (R$)", MOEDA, ""),
    ("in", "freios", "Freios, se o laudo apontar (R$)", MOEDA, ""),
    ("in", "reparos", "Outros reparos apontados (R$)", MOEDA, "Itens do laudo cautelar orçados em oficina."),
    ("fx", "sub_pos", "Subtotal logo após a compra (revisão e reparos)", MOEDA, None),
    ("sec", None, "5. RESERVA DE IMPREVISTOS", None, None),
    ("in", "res_pct", "Reserva (% do preço)", PCT, "Padrão 10%. Ajuste como preferir."),
    ("fx", "reserva", "Reserva de imprevistos (R$)", MOEDA, "Dinheiro separado, não é gasto certo."),
    ("sec", None, "6. SE FOR FINANCIAR", None, None),
    ("lista:fin", "fin", "Vai financiar? (Sim/Não)", None, "\"Não\" zera os cálculos do financiamento."),
    ("in", "entrada", "Entrada (R$)", MOEDA, ""),
    ("fx", "financiado", "Valor financiado (R$)", MOEDA, "Preço − entrada."),
    ("in", "n", "Número de parcelas", "0", ""),
    ("in", "taxa", "Taxa de juros ao mês (%)", '0.00%', "Use a taxa informada por escrito. Compare propostas pelo CET (Custo Efetivo Total)."),
    ("fx", "parcela", "Parcela mensal (R$)", MOEDA, "Tabela Price (função PMT). A parcela real pode incluir tarifas, seguro e impostos."),
    ("fx", "total_fin", "Total pago nas parcelas (R$)", MOEDA, "Parcela × número de parcelas."),
    ("fx", "custo_fin", "Custo do financiamento (juros) (R$)", MOEDA, "Total pago nas parcelas − valor financiado."),
    ("sec", None, "7. CUSTOS DE USO", None, None),
    ("in", "ipva_aliq", "Alíquota de IPVA do seu estado (%)", '0.0%', "Varia por estado e tipo de veículo. [VERIFICAR]"),
    ("fx", "ipva_est", "IPVA anual estimado (R$)", MOEDA, "Alíquota × preço (aproximação: o estado usa tabela própria de valores)."),
    ("in", "lic_anual", "Licenciamento anual (taxa) (R$)", MOEDA, VERIF),
    ("in", "manut", "Manutenção mensal estimada (R$)", MOEDA, "Média mensal de revisões, pneus e desgaste."),
    ("in", "km_mes", "Quantos km você roda por mês", KM, ""),
    ("in", "comb_preco", "Preço do combustível (R$/litro)", MOEDA, ""),
    ("in", "consumo", "Consumo do carro (km/l)", "0.0", "Use um valor realista (cidade), não o melhor caso."),
    ("fx", "comb_mes", "Combustível por mês (R$)", MOEDA, "km/mês × preço ÷ consumo."),
    ("sec", None, "8. RESUMO", None, None),
    ("fx", "dia", "DINHEIRO NECESSÁRIO NO DIA DA COMPRA (R$)", MOEDA,
     "Preço (ou entrada) + transferência + débitos seus + vistoria cautelar + 1ª parcela do seguro."),
    ("fx", "logo", "Logo após a compra: revisão e reparos (R$)", MOEDA, "Subtotal do bloco 4 (sem o seguro e a cautelar)."),
    ("fx", "com_res", "Para comprar, deixar em dia e ter reserva (R$)", MOEDA, "Dia da compra + logo após + reserva."),
    ("fx", "ano1", "CUSTO TOTAL NO 1º ANO (R$)", MOEDA,
     "Compra (ou entrada + 12 parcelas) + transferência + débitos + cautelar + reparos + seguro anual + IPVA + licenciamento "
     "+ manutenção × 12 + combustível × 12. Não inclui a reserva."),
    ("fx", "ano1_mes", "Equivale a, por mês, no 1º ano (R$)", MOEDA, "Custo do 1º ano ÷ 12."),
]


def _layout_custo():
    R = {}
    r = 5
    for tipo, chave, *_ in LINHAS_CUSTO:
        if chave:
            R[chave] = r
        r += 1
    R["_fim"] = r - 1
    return R


def _fx_custo(chave, R):
    b = lambda k: f"$B${R[k]}"  # noqa: E731
    fin = f'{b("fin")}="Sim"'
    if chave == "sub_transf":
        return f'=SUM(B{R["vist_transf"]}:B{R["despachante"]})'
    if chave == "sub_deb":
        return f'=SUMIF(C{R["ipva_deb"]}:C{R["multas"]},"Comprador",B{R["ipva_deb"]}:B{R["multas"]})'
    if chave == "sub_pos":
        return f'=SUM(B{R["oleo"]}:B{R["reparos"]})'
    if chave == "reserva":
        return f'=N({b("preco")})*N({b("res_pct")})'
    if chave == "financiado":
        return f'=IF({fin},MAX(0,N({b("preco")})-N({b("entrada")})),0)'
    if chave == "parcela":
        return (f'=IF(OR(NOT({fin}),{b("financiado")}=0,N({b("n")})<=0),0,'
                f'IF(N({b("taxa")})=0,{b("financiado")}/{b("n")},PMT({b("taxa")},{b("n")},-{b("financiado")})))')
    if chave == "total_fin":
        return f'={b("parcela")}*N({b("n")})'
    if chave == "custo_fin":
        return f'=IF({b("parcela")}=0,0,{b("total_fin")}-{b("financiado")})'
    if chave == "ipva_est":
        return f'=N({b("preco")})*N({b("ipva_aliq")})'
    if chave == "comb_mes":
        return f'=IF(N({b("consumo")})=0,0,N({b("km_mes")})*N({b("comb_preco")})/{b("consumo")})'
    aquis_dia = f'IF({fin},N({b("entrada")}),N({b("preco")}))'
    if chave == "dia":
        return (f'={aquis_dia}+{b("sub_transf")}+{b("sub_deb")}+N({b("cautelar")})+N({b("seg1")})')
    if chave == "logo":
        return f'={b("sub_pos")}'
    if chave == "com_res":
        return f'={b("dia")}+{b("logo")}+{b("reserva")}'
    if chave == "ano1":
        aquis_ano = (f'IF(AND({fin},{b("parcela")}>0),N({b("entrada")})+{b("parcela")}*MIN(12,N({b("n")})),'
                     f'N({b("preco")}))')
        return (f'={aquis_ano}+{b("sub_transf")}+{b("sub_deb")}+N({b("cautelar")})+{b("sub_pos")}'
                f'+MAX(N({b("seg_anual")}),N({b("seg1")}))+{b("ipva_est")}+N({b("lic_anual")})'
                f'+N({b("manut")})*12+{b("comb_mes")}*12')
    if chave == "ano1_mes":
        return f'={b("ano1")}/12'
    raise KeyError(chave)


def build_custo(path, dados=None):
    d = dict(EXEMPLO_CUSTO)
    if dados:
        d.update(dados)
    R = _layout_custo()
    wb = Workbook()
    capa(wb, "Custo total da compra", [
        "Abra a aba Custo total e preencha as células AZUIS, de cima para baixo. As pretas se calculam sozinhas.",
        "Taxas de transferência e licenciamento mudam de estado para estado: consulte o órgão de trânsito do seu estado e digite os valores.",
        "Nos débitos do carro, marque quem vai pagar (vendedor ou comprador). Só entra na sua conta o que for do comprador.",
        "Se for financiar, informe entrada, número de parcelas e taxa de juros ao mês: a planilha calcula parcela, total pago e juros.",
        "No resumo você vê o dinheiro necessário no dia da compra e o custo total no 1º ano.",
        "Na aba Orçamento mensal, compare quanto o carro vai custar por mês com a sua renda disponível.",
        "Apague os exemplos (células azuis) antes de usar com os seus números.",
    ], extra=[
        "O preço é só uma parte: transferência, débitos, seguro, revisão, reparos, IPVA e combustível entram na conta.",
        "Os percentuais (reserva, alíquota de IPVA, limite do orçamento) são editáveis. Nada aqui é conselho financeiro.",
    ])
    ws = wb.create_sheet("Custo total")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 50, "B": 18, "C": 15, "D": 62})
    titulo_aba(ws, "CUSTO TOTAL DA COMPRA", 4,
               "Preencha as células azuis. Valores de transferência ficam em branco de propósito: "
               "cada estado cobra valores diferentes.")
    header_row(ws, 4, ["Item", "Valor", "Quem paga", "Observação"])
    for tipo, chave, lab, fmt, obs in LINHAS_CUSTO:
        if tipo == "sec":
            continue
        r = R[chave]
        destaque = tipo == "fx" and (chave.startswith("sub") or chave in ("dia", "com_res", "ano1", "parcela", "reserva"))
        forte = chave in ("dia", "ano1")
        rotulo(ws.cell(row=r, column=1), lab, bold=destaque, fill=AMBAR if forte else (AMBAR_CLARO if destaque else None))
        if tipo == "fx":
            formula(ws.cell(row=r, column=2), _fx_custo(chave, R), bold=destaque, fmt=fmt, align=DIR)
            if forte:
                ws.cell(row=r, column=2).font = f(12, True, CARBONO)
                ws.cell(row=r, column=2).fill = FILL(AMBAR_CLARO)
        else:
            entrada(ws.cell(row=r, column=2), d.get(chave), fmt=fmt, align=DIR if fmt != PCT else CENTRO)
        cq = ws.cell(row=r, column=3)
        cq.border = BORDA
        if tipo == "quem":
            entrada(cq, d.get(chave + "_q"))
            dv_lista(ws, QUEM, cq.coordinate)
        dica(ws.cell(row=r, column=4), obs or "")
        if obs and "[VERIFICAR]" in obs:
            ws.cell(row=r, column=4).font = f(8.5, True, AMBAR_ESCURO, italic=True)
        ws.row_dimensions[r].height = 34 if len(obs or "") > 95 else (22 if not forte else 26)
        if tipo.startswith("lista"):
            dv_lista(ws, ["Sim", "Não"], f"B{r}")
    r = 5
    for tipo, chave, lab, *_ in LINHAS_CUSTO:
        if tipo == "sec":
            secao(ws, r, lab, 4)
        r += 1
    ws.row_dimensions[R["ano1"]].height = 46
    cor_lista(ws, f"C{R['ipva_deb']}:C{R['multas']}", [("Comprador", AMBAR_CLARO, AMBAR_ESCURO)])
    for k in ("preco", "vist_transf", "taxa_orgao", "firma", "despachante", "ipva_deb", "lic_deb", "multas", "cautelar",
              "seg1", "seg_anual", "oleo", "correia", "pneus", "freios", "reparos", "entrada", "lic_anual", "manut",
              "km_mes", "comb_preco", "consumo"):
        dv_numero(ws, f"B{R[k]}", 0)
    dv_numero(ws, f"B{R['n']}", 0, 120, inteiro=True, msg="Número inteiro de parcelas (0 a 120).")
    for k in ("res_pct", "taxa", "ipva_aliq"):
        dv_numero(ws, f"B{R[k]}", 0, 1, msg="Digite um percentual (ex.: 1,9% ou 0,019).")
    ws[f"A{R['parcela']}"].comment = Comment(
        "Parcela = PMT(taxa; nº de parcelas; −valor financiado) — sistema de parcelas iguais (Price). "
        "Com taxa 0%, parcela = valor financiado ÷ nº de parcelas.", "Antes de Fechar Negócio")
    rn = R["_fim"] + 2
    nota(ws, f"A{rn}:D{rn + 1}",
         "Valores de exemplo fictícios. Taxas públicas, alíquotas e prazos mudam por estado e ao longo do tempo: confirme "
         "no órgão de trânsito e na secretaria da fazenda do seu estado. Em financiamento, peça a proposta por escrito com "
         "o CET (Custo Efetivo Total). Esta planilha é uma ferramenta de organização, não conselho financeiro.", altura=26)
    ws.row_dimensions[rn + 1].height = 26
    ws.freeze_panes = "B5"
    impressao(ws, paisagem=False, area=f"A1:D{rn + 1}", titulos="4:4")

    Ro = _orcamento(wb, R, d)
    wb.calculation.fullCalcOnLoad = True
    wb.active = 0
    wb.save(path)
    return {"custo": R, "orc": Ro}


def _orcamento(wb, R, d):
    ws = wb.create_sheet("Orçamento mensal")
    ws.sheet_view.showGridLines = False
    widths(ws, {"A": 46, "B": 18, "C": 58})
    titulo_aba(ws, "ORÇAMENTO MENSAL DO CARRO", 3,
               "Quanto o carro vai custar por mês comparado com a sua renda disponível. É um indicador para você "
               "organizar as contas, não uma regra.")
    header_row(ws, 4, ["Item", "Por mês (R$)", "De onde vem"])
    CT = "'Custo total'!"
    Ro = {}
    linhas = [
        ("in", "renda", "Renda mensal disponível (líquida) (R$)", MOEDA, "Quanto entra por mês, já descontados impostos."),
        ("in", "limite", "Limite que você define (% da renda)", PCT,
         "Padrão 20%. É só uma referência pessoal; ajuste à sua realidade."),
        ("sec", None, "CUSTOS MENSAIS DO CARRO", None, None),
        ("fx", "o_parcela", "Parcela do financiamento", MOEDA, "Da aba Custo total (0 se não financiar)."),
        ("fx", "o_seguro", "Seguro (valor anual ÷ 12)", MOEDA, "Da aba Custo total."),
        ("fx", "o_ipva", "IPVA estimado (÷ 12)", MOEDA, "Da aba Custo total."),
        ("fx", "o_lic", "Licenciamento (÷ 12)", MOEDA, "Da aba Custo total."),
        ("fx", "o_manut", "Manutenção", MOEDA, "Da aba Custo total."),
        ("fx", "o_comb", "Combustível", MOEDA, "Da aba Custo total."),
        ("in", "estac", "Estacionamento", MOEDA, "Garagem, rotativo, condomínio."),
        ("in", "outros", "Pedágio, lavagem e outros", MOEDA, ""),
        ("fx", "total", "CUSTO MENSAL DO CARRO", MOEDA, "Soma dos itens acima."),
        ("fx", "pct", "Quanto isso representa da sua renda", PCT, ""),
        ("fx", "sobra", "Sobra da renda depois do carro", MOEDA, "Renda − custo mensal do carro."),
        ("fx", "indic", "Indicador", None, "Compara com o limite que você definiu."),
    ]
    fx = {
        "o_parcela": f"={CT}$B${R['parcela']}",
        "o_seguro": f"=MAX(N({CT}$B${R['seg_anual']}),N({CT}$B${R['seg1']}))/12",
        "o_ipva": f"={CT}$B${R['ipva_est']}/12",
        "o_lic": f"=N({CT}$B${R['lic_anual']})/12",
        "o_manut": f"=N({CT}$B${R['manut']})",
        "o_comb": f"={CT}$B${R['comb_mes']}",
    }
    r = 5
    for tipo, k, *_ in linhas:
        if k:
            Ro[k] = r
        r += 1
    fx["total"] = f"=SUM(B{Ro['o_parcela']}:B{Ro['outros']})"
    fx["pct"] = f'=IF(N(B{Ro["renda"]})=0,"",B{Ro["total"]}/B{Ro["renda"]})'
    fx["sobra"] = f"=N(B{Ro['renda']})-B{Ro['total']}"
    fx["indic"] = (f'=IF(B{Ro["pct"]}="","Informe sua renda disponível.",IF(B{Ro["pct"]}>B{Ro["limite"]},'
                   f'"Acima do limite que você definiu: reveja o orçamento antes de fechar.",'
                   f'"Dentro do limite que você definiu."))')
    for tipo, k, lab, fmt, obs in linhas:
        r = Ro.get(k)
        if tipo == "sec":
            continue
        destaque = k in ("total", "indic", "pct")
        rotulo(ws.cell(row=r, column=1), lab, bold=destaque or tipo == "in" and k in ("renda", "limite"),
               fill=AMBAR_CLARO if destaque else None)
        if tipo == "in":
            entrada(ws.cell(row=r, column=2), d.get(k), bold=k in ("renda", "limite"), fmt=fmt,
                    align=CENTRO if fmt == PCT else DIR)
        else:
            formula(ws.cell(row=r, column=2), fx[k], bold=destaque, fmt=fmt, align=DIR if fmt else ESQ)
        dica(ws.cell(row=r, column=3), obs)
        ws.row_dimensions[r].height = 22
    # secao
    secao(ws, Ro["limite"] + 1, "CUSTOS MENSAIS DO CARRO", 3)
    ri = Ro["indic"]
    ws.merge_cells(f"B{ri}:C{ri}")
    ws.row_dimensions[ri].height = 30
    ws.conditional_formatting.add(f"B{ri}", FormulaRule(formula=[f'LEFT(B{ri},5)="Acima"'], fill=FILL(VERMELHO_CLARO),
                                                        font=Font(name=FONT, bold=True, color=VERMELHO)))
    ws.conditional_formatting.add(f"B{ri}", FormulaRule(formula=[f'LEFT(B{ri},6)="Dentro"'], fill=FILL(VERDE_CLARO),
                                                        font=Font(name=FONT, bold=True, color=VERDE)))
    ws.conditional_formatting.add(f"B{Ro['sobra']}", CellIsRule(operator="lessThan", formula=["0"],
                                                                font=Font(name=FONT, bold=True, color=VERMELHO)))
    dv_numero(ws, f"B{Ro['renda']}", 0)
    dv_numero(ws, f"B{Ro['limite']}", 0, 1, msg="Digite um percentual (ex.: 20%).")
    dv_numero(ws, f"B{Ro['estac']}:B{Ro['outros']}", 0)
    rn = ri + 2
    nota(ws, f"A{rn}:C{rn + 1}",
         "Quando o financiamento terminar, a parcela sai da conta. Manutenção tende a subir com a idade e a quilometragem "
         "do carro. O limite de % da renda é uma referência que você escolhe, não uma recomendação.", altura=24)
    ws.row_dimensions[rn + 1].height = 24
    ws.freeze_panes = "B5"
    impressao(ws, paisagem=False, area=f"A1:C{rn + 1}", titulos="4:4")
    return Ro


# ================================================================ main
ARQUIVOS = [
    ("01-comparativo-de-veiculos.xlsx", build_comparativo),
    ("02-custo-total-da-compra.xlsx", build_custo),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(OUT_DEFAULT))
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    for nome, fn in ARQUIVOS:
        fn(out / nome)
        print("ok", out / nome)


if __name__ == "__main__":
    main()
