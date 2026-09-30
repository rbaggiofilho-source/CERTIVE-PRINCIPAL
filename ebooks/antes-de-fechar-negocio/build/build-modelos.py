#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-modelos.py — gera os modelos editáveis (.docx) do e-book "Antes de Fechar Negócio".

Saída: ebooks/antes-de-fechar-negocio/03-materiais/modelos/*.docx
  01-recibo-de-sinal.docx
  02-contrato-particular-compra-e-venda-veiculo.docx
  03-roteiro-de-perguntas-ao-vendedor.docx

Público: consumidor leigo que compra carro usado de particular.
Visual neutro (Calibri, títulos carbono com fio âmbar, A4, margens de 2 cm, páginas numeradas).
Placeholders entre colchetes: [NOME COMPLETO], [ABC1D23]... Pontos a confirmar: [VERIFICAR].

Os utilitários de XML/tabela/texto foram copiados de protocolo-cautelar/build/build-modelos.py
(nada é importado de lá).
Uso: python3 build/build-modelos.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "03-materiais/modelos"

# Paleta (protocolo-cautelar/02-design/tema.css) — uso sóbrio: carbono nos títulos, âmbar só no fio.
CARBONO = "0F1318"
AMBAR = "F2A900"
AMBAR_ESC = "B87F00"
AMBAR_CLARO = "FFF4D6"
LINHA = "C9CED4"
PAPEL2 = "F2F3F5"
TEXTO = "1A1F26"
TEXTO2 = "4E5864"
TEXTO3 = "7A8591"

AVISO = "Modelo de referência para consulta. Não substitui a orientação de um advogado."
LARGURA = 17.0  # cm úteis (A4 21 cm − 2 × 2 cm)


# --------------------------------------------------------------------------------------
# Utilidades XML (ordem dos elementos respeita o schema OOXML)
# --------------------------------------------------------------------------------------
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _local(el):
    return el.tag.split("}")[-1]


def insert_ordered(parent, el, order):
    tag = _local(el)
    for old in parent.findall(qn("w:" + tag)):
        parent.remove(old)
    idx = order.index(tag)
    for child in parent:
        name = _local(child)
        if name in order and order.index(name) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def rgb(h):
    return RGBColor.from_string(h)


def cm2tw(v):
    return int(round(v / 2.54 * 1440))


def set_fonts(rpr_parent_el, name="Calibri"):
    rpr = rpr_parent_el.get_or_add_rPr()
    for old in rpr.findall(qn("w:rFonts")):
        rpr.remove(old)
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    rpr.insert(0, rf)


def p_border(p, side="bottom", sz=6, color=AMBAR, space=2, val="single"):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = insert_ordered(pPr, OxmlElement("w:pBdr"), PPR_ORDER)
    el = OxmlElement("w:" + side)
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    # ordem dentro de pBdr: top, left, bottom, right, between, bar
    order = ["top", "left", "bottom", "right", "between", "bar"]
    insert_ordered(bdr, el, order)


def p_shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd, PPR_ORDER)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(tcPr, shd, TCPR_ORDER)


def cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (val, sz, color)"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            val, sz, color = sides[side]
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), val)
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            b.append(e)
    insert_ordered(tcPr, b, TCPR_ORDER)


def tbl_borders(tbl, val="single", sz=4, color=LINHA, inside=True):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side)
        if side.startswith("inside") and not inside:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), val)
            if val != "nil":
                e.set(qn("w:sz"), str(sz))
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), color)
        b.append(e)
    insert_ordered(tblPr, b, TBLPR_ORDER)


def tbl_cell_margins(tbl, top=0.06, left=0.15, bottom=0.06, right=0.15):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(cm2tw(v)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    insert_ordered(tblPr, m, TBLPR_ORDER)


def set_widths(tbl, widths):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(sum(cm2tw(x) for x in widths)))
    w.set(qn("w:type"), "dxa")
    insert_ordered(tblPr, w, TBLPR_ORDER)
    for gc, wd in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(cm2tw(wd)))
    for row in tbl.rows:
        for c, wd in zip(row.cells, widths):
            c.width = Cm(wd)


def keep_table(t):
    """Mantém a tabela inteira na mesma página (keep-with-next em todas as linhas menos a última)."""
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t


def row_flags(row, header=False, cant_split=True):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def add_field(p, instr, size=8, color=TEXTO3):
    def run():
        r = p.add_run()
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(color)
        return r
    r = run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r = run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    r = run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
    r = run(); r.text = "1"
    r = run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)


# --------------------------------------------------------------------------------------
# Texto rico: **negrito**, [PLACEHOLDER] destacado, ☐ com fonte de símbolo
# --------------------------------------------------------------------------------------
TOK = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]|☐)")


def add_rich(p, text, size=None, color=None, bold=False, italic=False):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
            if color:
                r.font.color.rgb = rgb(color)
        elif part.startswith("[") and part.endswith("]"):
            r = p.add_run(part); r.font.color.rgb = rgb(AMBAR_ESC)
            if bold:
                r.bold = True
        elif part == "☐":
            r = p.add_run(part); set_fonts(r._r, "Segoe UI Symbol")
            r.font.color.rgb = rgb(TEXTO)
        else:
            r = p.add_run(part)
            if bold:
                r.bold = True
            if color:
                r.font.color.rgb = rgb(color)
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
    return p


def _target_par(container):
    """Em células, reaproveita o primeiro parágrafo vazio."""
    if hasattr(container, "_tc"):
        ps = container.paragraphs
        if len(ps) == 1 and not ps[0].runs and not ps[0].text:
            return ps[0]
    return container.add_paragraph()


def para(container, text="", size=None, color=None, bold=False, italic=False, align=None,
         after=None, before=None, indent=None, hanging=None, keep=False, style=None):
    p = _target_par(container)
    if style:
        p.style = style
    add_rich(p, text, size=size, color=color, bold=bold, italic=italic)
    pf = p.paragraph_format
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if after is not None:
        pf.space_after = Pt(after)
    if before is not None:
        pf.space_before = Pt(before)
    if indent is not None:
        pf.left_indent = Cm(indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    if keep:
        pf.keep_with_next = True
    return p


def bullets(container, items, size=None, align="j"):
    for it in items:
        p = para(container, it, size=size, align=align, after=2, style="List Bullet")
    return p


def numbered(container, items, prefix="", size=None, start=1, indent=0.9):
    for i, it in enumerate(items, start):
        para(container, f"{prefix}{i}.\t{it}", size=size, align="j", indent=indent, hanging=indent, after=3)


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def spacer(container, pt=4):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(); r.font.size = Pt(pt)
    p.paragraph_format.line_spacing = Pt(pt)
    return p


def page_break(doc):
    doc.add_page_break()


# --------------------------------------------------------------------------------------
# Tabelas
# --------------------------------------------------------------------------------------
def table(container, headers, rows, widths, size=9, header_fill=PAPEL2, row_h=None, aligns=None,
          repeat_header=True, first_bold=False):
    n = len(widths)
    nrows = len(rows) + (1 if headers else 0)
    t = container.add_table(rows=nrows, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, widths)
    tbl_borders(t)
    tbl_cell_margins(t)
    aligns = aligns or ["l"] * n
    r0 = 0
    if headers:
        row = t.rows[0]
        row_flags(row, header=repeat_header)
        for i, htxt in enumerate(headers):
            c = row.cells[i]
            shade(c, header_fill)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para(c, htxt, size=size - 0.5, bold=True, color=CARBONO, after=0,
                 align={"c": "c", "r": "r"}.get(aligns[i]))
        r0 = 1
    for ri, data in enumerate(rows):
        row = t.rows[r0 + ri]
        row_flags(row)
        if row_h:
            row.height = Cm(row_h)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i in range(n):
            c = row.cells[i]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            txt = data[i] if i < len(data) else ""
            para(c, txt, size=size, after=0, bold=(first_bold and i == 0),
                 align={"c": "c", "r": "r"}.get(aligns[i]))
    return t


def form(container, rows, label_w=3.3, total=LARGURA, size=9.5, h=0.72, cols=2):
    """rows: lista de linhas; cada linha é lista de (rótulo, valor). 1 par = valor ocupa a linha toda."""
    value_w = (total - label_w * cols) / cols
    widths = [label_w, value_w] * cols
    t = container.add_table(rows=len(rows), cols=2 * cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, widths)
    tbl_borders(t)
    tbl_cell_margins(t)
    for ri, pairs in enumerate(rows):
        row = t.rows[ri]
        row_flags(row)
        row.height = Cm(h)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cells = row.cells
        if len(pairs) == 1:
            merged = cells[1].merge(cells[2 * cols - 1])
            cells = [cells[0], merged]
        for pi, (lab, val) in enumerate(pairs):
            lc, vc = cells[2 * pi], cells[2 * pi + 1]
            shade(lc, PAPEL2)
            for c in (lc, vc):
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para(lc, lab, size=size - 1, bold=True, color=TEXTO2, after=0)
            para(vc, val, size=size, after=0)
    return t


def boxed(container, lines, fill=PAPEL2, size=10, border=LINHA, width=LARGURA, bar=None):
    """Caixa de texto (tabela 1×1). lines: lista de strings (cada uma vira parágrafo)."""
    t = container.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, [width])
    tbl_borders(t, color=border)
    tbl_cell_margins(t, top=0.2, left=0.3, bottom=0.2, right=0.3)
    c = t.cell(0, 0)
    shade(c, fill)
    if bar:
        cell_borders(c, left=("single", 24, bar), top=("single", 4, border),
                     bottom=("single", 4, border), right=("single", 4, border))
    row_flags(t.rows[0])
    for ln in lines:
        para(c, ln, size=size, after=3)
    return t


def assinaturas(container, blocos, cols=2, gap=1.0):
    """blocos: lista de listas de linhas (primeira linha = papel)."""
    if cols == 1:
        t = container.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_widths(t, [4.5, 8.0, 4.5])
        tbl_borders(t, val="nil")
        row_flags(t.rows[0])
        cell = t.cell(0, 1)
        p = para(cell, "", after=0)
        p.paragraph_format.space_before = Pt(30)
        p2 = para(cell, blocos[0][0], size=9, bold=True, color=CARBONO, after=0, align="c")
        p_border(p2, "top", sz=6, color=TEXTO, space=4)
        for ln in blocos[0][1:]:
            para(cell, ln, size=8.5, color=TEXTO2, after=0, align="c")
        return t
    rows = (len(blocos) + cols - 1) // cols
    w = (LARGURA - gap * (cols - 1)) / cols
    widths = []
    for i in range(cols):
        widths.append(w)
        if i < cols - 1:
            widths.append(gap)
    t = container.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, widths)
    tbl_borders(t, val="nil")
    for bi, bloco in enumerate(blocos):
        r, c = divmod(bi, cols)
        row_flags(t.rows[r])
        cell = t.cell(r, c * 2)
        p = para(cell, "", after=0)
        p.paragraph_format.space_before = Pt(30)
        p2 = para(cell, bloco[0], size=9, bold=True, color=CARBONO, after=0, align="c")
        p_border(p2, "top", sz=6, color=TEXTO, space=4)
        for ln in bloco[1:]:
            para(cell, ln, size=8.5, color=TEXTO2, after=0, align="c")
    return t



# --------------------------------------------------------------------------------------
# Documento base, título, rodapé (adaptado)
# --------------------------------------------------------------------------------------
def novo_doc(rotulo_rodape):
    doc = Document()
    for z in doc.settings.element.findall(qn("w:zoom")):  # template padrão traz <w:zoom> sem w:percent
        z.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.0))
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.9)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = rgb(TEXTO)
    set_fonts(st.element)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (("Heading 1", 13.5, 14, 6), ("Heading 2", 11, 10, 4), ("Heading 3", 10.5, 8, 3)):
        s = doc.styles[name]
        set_fonts(s.element)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = rgb(CARBONO)
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        if name == "Heading 1":
            ppr = s.element.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            e = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "3"), ("w:color", AMBAR)):
                e.set(qn(k), v)
            bdr.append(e)
            insert_ordered(ppr, bdr, PPR_ORDER)
    lb = doc.styles["List Bullet"]
    set_fonts(lb.element)
    lb.paragraph_format.left_indent = Cm(0.6)
    lb.paragraph_format.first_line_indent = Cm(-0.4)

    ft = sec.footer
    p = ft.paragraphs[0]
    p.style = doc.styles["Normal"]
    p_border(p, "top", sz=4, color=LINHA, space=4)
    r = p.add_run(AVISO)
    r.italic = True
    r.font.size = Pt(8)
    r.font.color.rgb = rgb(TEXTO2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p2 = ft.add_paragraph()
    p2.style = doc.styles["Normal"]
    p2.paragraph_format.tab_stops.add_tab_stop(Cm(LARGURA), WD_TAB_ALIGNMENT.RIGHT)
    p2.paragraph_format.space_after = Pt(0)
    r = p2.add_run(rotulo_rodape); r.font.size = Pt(8); r.font.color.rgb = rgb(TEXTO3)
    r = p2.add_run("\tPágina "); r.font.size = Pt(8); r.font.color.rgb = rgb(TEXTO3)
    add_field(p2, "PAGE")
    r = p2.add_run(" de "); r.font.size = Pt(8); r.font.color.rgb = rgb(TEXTO3)
    add_field(p2, "NUMPAGES")
    return doc


def titulo(doc, text, sub=None):
    p = para(doc, text.upper(), size=16, bold=True, color=CARBONO, before=0, after=2 if sub else 6)
    if sub:
        p = para(doc, sub, size=9.5, color=TEXTO2, after=6)
    p_border(p, "bottom", sz=8, color=AMBAR, space=4)
    spacer(doc, 6)
    return p


def secao(doc, text):
    """Título de seção numerada (Heading 2 com fio âmbar curto)."""
    p = h2(doc, text)
    p_border(p, "bottom", sz=4, color=AMBAR, space=2)
    return p


def clausula(doc, n, nome, itens):
    """Cláusula numerada: itens str | (str, [subitens])."""
    h2(doc, f"CLÁUSULA {n}ª – {nome.upper()}")
    for i, it in enumerate(itens, 1):
        subs = []
        if isinstance(it, tuple):
            it, subs = it
        para(doc, f"{n}.{i}.\t{it}", align="j", indent=1.0, hanging=1.0, after=3)
        for j, s in enumerate(subs):
            para(doc, f"{chr(97 + j)})\t{s}", align="j", indent=1.6, hanging=0.6, after=2)


def orientacao(doc, titulo_box, linhas, apagar=True):
    """Caixa de orientação ao leitor (fundo âmbar-claro, barra âmbar)."""
    cab = f"**{titulo_box}**" + ("  (apague esta caixa antes de assinar)" if apagar else "")
    t = boxed(doc, [cab] + linhas, fill=AMBAR_CLARO, size=9, border="E8D9A8", bar=AMBAR)
    spacer(doc, 6)
    return t


def local_data(doc):
    return para(doc, "[CIDADE/UF], [DIA] de [MÊS] de [ANO].", align="r", before=10, after=4, keep=True)


def fechamento(doc, texto, blocos):
    """Frase final + local/data + assinaturas, mantidos na mesma página."""
    if texto:
        para(doc, texto, align="j", before=6, keep=True)
    local_data(doc)
    keep_table(assinaturas(doc, blocos))


def dados_pessoa(papel):
    return [
        [(f"{papel} – nome completo", "[NOME COMPLETO]")],
        [("CPF", "[000.000.000-00]"), ("RG / órgão", "[00.000.000-0 / ÓRGÃO-UF]")],
        [("Endereço", "[RUA, Nº, BAIRRO, CIDADE/UF, CEP]")],
        [("Telefone", "[(00) 00000-0000]"), ("E-mail", "[E-MAIL]")],
    ]


DADOS_VEICULO = [
    [("Marca / modelo / versão", "[MARCA MODELO VERSÃO]")],
    [("Ano fabricação / modelo", "[0000] / [0000]"), ("Cor", "[COR]")],
    [("Placa", "[ABC1D23]"), ("RENAVAM", "[00000000000]")],
    [("Chassi", "[17 CARACTERES]"), ("Combustível", "[FLEX/GASOLINA/...]")],
]


# ======================================================================================
# 01 — RECIBO DE SINAL (ARRAS)
# ======================================================================================
def recibo_sinal():
    doc = novo_doc("Recibo de sinal (arras) – veículo usado")
    titulo(doc, "Recibo de sinal (arras)", "Compra e venda de veículo usado entre particulares")
    orientacao(doc, "Como usar este modelo", [
        "Preencha os campos entre colchetes [ ] e marque as opções ☐ que valem para o seu caso. Faça duas vias: uma para cada parte.",
        "**O que é o sinal (arras):** um valor pago para confirmar o negócio. As regras estão no Código Civil (Lei 10.406/2002), "
        "**arts. 417 a 420**. Resumo: nas **arras confirmatórias** (padrão deste modelo), o sinal é descontado do preço quando o "
        "negócio é concluído (art. 417). Se quem pagou o sinal desistir sem motivo, perde o sinal; se quem recebeu desistir, "
        "devolve o sinal mais o equivalente, com atualização monetária, juros e honorários de advogado (art. 418). A parte "
        "prejudicada pode pedir indenização maior se provar prejuízo maior (art. 419). Se as partes combinarem **direito de "
        "arrependimento** (arras penitenciais), o sinal serve só como indenização e não cabe indenização suplementar (art. 420).",
        "**Condição (opcional, item 5):** permite desfazer o negócio com devolução integral do sinal se o laudo cautelar ou as "
        "consultas mostrarem problema. Recomendada quando você ainda não fez a vistoria e as consultas.",
        "**Cuidados:** sinal pequeno, pago só depois de ver o carro e o documento, por meio rastreável (PIX ou transferência) e "
        "**somente para a conta do proprietário que consta no documento do veículo**. Nunca pague sinal a intermediário ou terceiro.",
    ])

    secao(doc, "1. Partes")
    form(doc, dados_pessoa("VENDEDOR(A)"))
    spacer(doc, 4)
    form(doc, dados_pessoa("COMPRADOR(A)"))

    secao(doc, "2. Veículo")
    form(doc, DADOS_VEICULO + [[("Quilometragem no dia", "[000.000] km"), ("Documento (CRLV) exercício", "[0000]")]])

    secao(doc, "3. Valores")
    form(doc, [
        [("Preço total do veículo", "R$ [0.000,00] ([VALOR POR EXTENSO])")],
        [("Valor do sinal", "R$ [0.000,00] ([VALOR POR EXTENSO])")],
        [("Saldo a pagar", "R$ [0.000,00] ([VALOR POR EXTENSO])")],
        [("Sinal pago por", "☐ PIX   ☐ transferência bancária   ☐ outro: [ ]")],
        [("Data do sinal", "[DD/MM/AAAA]")],
        [("Conta do vendedor (titular do documento)", "Banco [ ]  ·  agência [ ]  ·  conta [ ]  ·  chave PIX [ ]")],
    ], label_w=4.0)

    secao(doc, "4. Declaração de recebimento e prazo")
    para(doc, "Eu, **VENDEDOR(A)** acima qualificado(a), declaro que recebi do(a) **COMPRADOR(A)** a quantia de "
              "**R$ [0.000,00]** ([VALOR POR EXTENSO]), a título de **sinal e princípio de pagamento** (arras confirmatórias, "
              "art. 417 do Código Civil), referente à venda do veículo descrito no item 2, pelo preço total de R$ [0.000,00].",
         align="j")
    para(doc, "O saldo de **R$ [0.000,00]** será pago até **[DD/MM/AAAA]**, na mesma data em que as partes assinarão o contrato de "
              "compra e venda e o documento de transferência. O valor do sinal será **descontado do preço** na conclusão do negócio.",
         align="j")
    para(doc, "Até essa data, o(a) VENDEDOR(A) se compromete a não vender, prometer ou entregar o veículo a outra pessoa, a mantê-lo "
              "no estado em que foi visto e a permitir a vistoria cautelar prevista no item 5.", align="j")

    secao(doc, "5. Condição para concluir o negócio (opcional)")
    para(doc, "☐ **Marque se quiser usar.** Este negócio fica condicionado a:", align="j")
    para(doc, "a)\tlaudo de **vistoria cautelar**, feito por profissional escolhido pelo(a) COMPRADOR(A) até [DD/MM/AAAA], "
              "**sem classificação \"Reprovado\"** [ou \"Inconclusivo\"]; e", align="j", indent=1.0, hanging=0.6)
    para(doc, "b)\t**consultas** de débitos, restrições, gravame (financiamento) e histórico do veículo **sem pendências que não "
              "tenham sido informadas por escrito** pelo(a) VENDEDOR(A) antes deste recibo.", align="j", indent=1.0, hanging=0.6)
    para(doc, "Se qualquer dessas condições não se cumprir, o negócio fica desfeito e o(a) VENDEDOR(A) **devolverá o sinal "
              "integralmente, sem qualquer desconto, em até [X] dias**, para a conta de onde o valor saiu. O(A) VENDEDOR(A) autoriza "
              "a vistoria cautelar no local [ENDEREÇO] ou em local combinado. (Código Civil, arts. 121 a 130 – condição.)",
         align="j")

    secao(doc, "6. Desistência")
    para(doc, "Marque **uma** opção:", after=2)
    para(doc, "☐ **Arras confirmatórias (padrão):** sem direito de arrependimento. Se o(a) COMPRADOR(A) desistir sem justa causa, "
              "perde o sinal; se o(a) VENDEDOR(A) desistir, devolve o sinal mais o equivalente, com atualização monetária "
              "(Código Civil, arts. 418 e 419).", indent=0.4)
    para(doc, "☐ **Arras penitenciais:** as partes podem se arrepender até [DD/MM/AAAA]. Quem desistir perde o sinal (se pagou) ou "
              "devolve o sinal mais o equivalente (se recebeu), sem outra indenização (Código Civil, art. 420).",
         indent=0.4)
    para(doc, "Em qualquer caso, o descumprimento da condição do item 5 **não é desistência**: gera a devolução integral do sinal.",
         align="j", italic=True, size=9.5)

    fechamento(doc, "E, por estarem de acordo, assinam este recibo em 2 (duas) vias de igual teor, na presença de duas "
                    "testemunhas.", [["VENDEDOR(A)", "[NOME COMPLETO] · CPF [000.000.000-00]"],
                      ["COMPRADOR(A)", "[NOME COMPLETO] · CPF [000.000.000-00]"],
                      ["TESTEMUNHA 1", "[NOME] · CPF [000.000.000-00]"],
                      ["TESTEMUNHA 2", "[NOME] · CPF [000.000.000-00]"]])
    para(doc, "Base legal: Código Civil – Lei 10.406/2002, arts. 121 a 130 (condição) e 417 a 420 (arras).",
         size=8, italic=True, color=TEXTO3, before=8, after=0)
    doc.save(OUT / "01-recibo-de-sinal.docx")


# ======================================================================================
# 02 — CONTRATO PARTICULAR DE COMPRA E VENDA
# ======================================================================================
def contrato():
    doc = novo_doc("Contrato particular de compra e venda de veículo usado")
    titulo(doc, "Contrato particular de compra e venda de veículo usado", "Entre particulares (pessoas físicas)")
    orientacao(doc, "Como usar este modelo", [
        "Preencha os campos entre colchetes [ ], marque as opções ☐ e apague o que não se aplica. Assinem em 2 vias, com 2 testemunhas.",
        "Este modelo é para venda **entre particulares**. Se o vendedor for loja ou revenda, vale também o Código de Defesa do "
        "Consumidor (Lei 8.078/1990), que não pode ser afastado por contrato.",
        "O contrato **não substitui** o documento de transferência exigido pelo órgão de trânsito (CRV assinado com firma "
        "reconhecida ou autorização de transferência eletrônica, conforme o estado) [VERIFICAR no órgão de trânsito do seu estado].",
        "**Pagamento:** somente para a conta do proprietário que consta no documento do veículo, por meio rastreável. "
        "Entregue o carro e o documento de transferência só com o dinheiro confirmado na conta.",
    ])

    para(doc, "Pelo presente instrumento particular, as partes abaixo qualificadas:", align="j")
    form(doc, dados_pessoa("VENDEDOR(A)"))
    spacer(doc, 4)
    form(doc, dados_pessoa("COMPRADOR(A)"))
    para(doc, "têm entre si justo e acordado o presente contrato de compra e venda de veículo usado, que se regerá pelas "
              "cláusulas seguintes.", align="j", before=6, after=6)

    clausula(doc, 1, "Do objeto", [
        "O(A) VENDEDOR(A) vende ao(à) COMPRADOR(A) o veículo abaixo descrito (\"Veículo\"), de sua propriedade, conforme o "
        "documento do veículo (CRLV) em seu nome:",
    ])
    form(doc, DADOS_VEICULO + [
        [("Nº do motor", "[CONFORME DOCUMENTO, SE CONSTAR]"), ("Quilometragem", "[000.000] km")],
        [("Acessórios incluídos", "[DESCREVER OU \"NENHUM\"]")],
    ])
    spacer(doc, 4)

    clausula(doc, 2, "Do preço e da forma de pagamento", [
        ("O preço total é de **R$ [0.000,00]** ([VALOR POR EXTENSO]), pago da seguinte forma (marque uma opção):", [
            "☐ à vista, em [DD/MM/AAAA], no ato da entrega;",
            "☐ sinal de R$ [0.000,00], já pago conforme recibo de [DD/MM/AAAA], e saldo de R$ [0.000,00] em [DD/MM/AAAA];",
            "☐ entrada de R$ [0.000,00] e saldo de R$ [0.000,00] por financiamento contratado pelo(a) COMPRADOR(A) junto a "
            "[INSTITUIÇÃO FINANCEIRA], sendo a entrega condicionada à liberação do crédito ao(à) VENDEDOR(A).",
        ]),
        "Todos os pagamentos serão feitos **exclusivamente por meio rastreável** (PIX ou transferência bancária) para conta de "
        "titularidade do(a) VENDEDOR(A), **que é o(a) proprietário(a) que consta no documento do Veículo**: banco [ ], agência [ ], "
        "conta [ ], chave PIX [ ].",
        "O(A) COMPRADOR(A) não fará pagamento a intermediários, parentes, empresas ou terceiros, **ainda que o(a) VENDEDOR(A) peça**. "
        "Pagamento feito a terceiro não quita o preço.",
        "A quitação só ocorre com o crédito **confirmado** na conta do(a) VENDEDOR(A). A entrega do Veículo e do documento de "
        "transferência assinado ocorre após essa confirmação.",
    ])

    clausula(doc, 3, "Das declarações do(a) vendedor(a)", [
        "O(A) VENDEDOR(A) declara que é o(a) legítimo(a) proprietário(a) do Veículo e presta as informações abaixo, que o(a) "
        "COMPRADOR(A) considera essenciais para decidir a compra:",
    ])
    decl = [
        "Há débitos de IPVA, licenciamento ou multas em aberto?",
        "Há restrição judicial, administrativa, de furto ou roubo?",
        "Há financiamento, alienação fiduciária ou outro gravame?",
        "O Veículo já passou por leilão ou foi recuperado de furto/roubo?",
        "O Veículo já teve sinistro (batida, alagamento, incêndio) ou indenização de seguro?",
        "O Veículo tem alteração de características (suspensão, motor, GNV, rodas, etc.)?",
        "O Veículo já foi usado como táxi, transporte por aplicativo, locadora ou frota?",
        "Há defeito, reparo ou problema conhecido pelo(a) vendedor(a)?",
        "A quilometragem do painel é, até onde sabe, a real?",
    ]
    table(doc, ["Declaração do(a) VENDEDOR(A)", "Sim", "Não", "Se sim, qual / detalhes"],
          [[d, "☐", "☐", "[ ]"] for d in decl], [7.6, 1.1, 1.1, 7.2], size=9, aligns=["l", "c", "c", "l"], row_h=0.8)
    spacer(doc, 4)
    para(doc, "3.2.\tA declaração falsa ou a omissão de informação relevante que o(a) VENDEDOR(A) conhecia permite ao(à) COMPRADOR(A) "
              "pedir o desfazimento do negócio ou o abatimento do preço, além de perdas e danos, nos termos da lei.",
         align="j", indent=1.0, hanging=1.0, after=3)
    para(doc, "3.3.\tDébitos existentes até a entrega, se houver, serão quitados por [VENDEDOR(A) / COMPRADOR(A)] até [DD/MM/AAAA] "
              "ou descontados do preço, conforme combinado: [DESCREVER].", align="j", indent=1.0, hanging=1.0, after=3)

    clausula(doc, 4, "Do estado do veículo", [
        "O(A) COMPRADOR(A) examinou o Veículo e fez **test drive** em [DD/MM/AAAA], no trajeto [DESCREVER].",
        "☐ Foi feita **vistoria cautelar** por profissional escolhido pelo(a) COMPRADOR(A) em [DD/MM/AAAA], laudo nº [ ], com "
        "classificação [ ], cuja cópia integra este contrato como **Anexo II**. ☐ Não foi feita vistoria cautelar.",
        "Apontamentos conhecidos e aceitos pelo(a) COMPRADOR(A), já considerados no preço: [LISTAR OU \"NENHUM\"].",
        "O laudo cautelar descreve o que foi observado no momento da vistoria e não garante o Veículo; ele não afasta as "
        "declarações da Cláusula 3 nem a garantia da Cláusula 8.",
    ])

    clausula(doc, 5, "Da entrega", [
        "O Veículo será entregue em [DD/MM/AAAA], às [00h00], em [LOCAL], com o registro de **data, hora e quilometragem** no "
        "**Termo de Entrega (Anexo I)**, assinado pelas partes.",
        "Até a entrega, os riscos do Veículo correm por conta do(a) VENDEDOR(A); a partir dela, por conta do(a) COMPRADOR(A) "
        "(Código Civil, art. 492).",
    ])

    clausula(doc, 6, "Das multas, débitos e responsabilidades", [
        "Multas, débitos e responsabilidades cujo fato tenha ocorrido **até a data e hora da entrega** são do(a) VENDEDOR(A); os "
        "ocorridos **depois** são do(a) COMPRADOR(A).",
        "Se chegar notificação de infração ocorrida depois da entrega em nome do(a) VENDEDOR(A), o(a) COMPRADOR(A) se obriga a "
        "assinar a indicação de condutor no prazo informado na notificação e a pagar os valores correspondentes.",
        "IPVA e licenciamento do exercício de [ANO]: ☐ pagos pelo(a) VENDEDOR(A)  ☐ a cargo do(a) COMPRADOR(A)  ☐ divididos "
        "proporcionalmente: [DESCREVER].",
    ])

    clausula(doc, 7, "Da transferência e da comunicação de venda", [
        "O(A) VENDEDOR(A) entregará, na data da entrega, o documento de transferência preenchido e assinado conforme a exigência "
        "do órgão de trânsito do estado (CRV com firma reconhecida ou autorização de transferência eletrônica) [VERIFICAR].",
        "O(A) COMPRADOR(A) deverá providenciar a transferência de propriedade no prazo de **30 (trinta) dias** (Código de Trânsito "
        "Brasileiro – Lei 9.503/1997, art. 123, § 1º), arcando com vistoria de transferência, taxas e demais custos.",
        "O(A) VENDEDOR(A) fará a **comunicação de venda** ao órgão executivo de trânsito do estado, com cópia do comprovante de "
        "transferência, nos termos do art. 134 do CTB [VERIFICAR prazo vigente], sob pena de responder solidariamente pelas "
        "penalidades impostas até a data da comunicação. O(A) COMPRADOR(A) fornecerá a cópia necessária.",
        "O(A) COMPRADOR(A) responde pelas consequências de não transferir o Veículo no prazo.",
    ])

    clausula(doc, 8, "Dos vícios ocultos", [
        "Este contrato **não exclui** a garantia legal contra vícios ocultos prevista nos **arts. 441 a 446 do Código Civil**.",
        "Para informação das partes: o(a) COMPRADOR(A) pode rejeitar o Veículo ou pedir abatimento do preço (arts. 441 e 442); se "
        "o(a) VENDEDOR(A) conhecia o vício, responde também por perdas e danos (art. 443); o prazo para bens móveis é de 30 dias "
        "contados da entrega e, quando o vício só puder ser conhecido mais tarde, conta-se de quando dele se tiver ciência, até o "
        "limite de 180 dias (art. 445) [VERIFICAR redação e contagem no caso concreto].",
        "Problemas **informados** pelo(a) VENDEDOR(A) na Cláusula 3 ou apontados no laudo da Cláusula 4 não são vícios ocultos.",
    ])

    clausula(doc, 9, "Do descumprimento", [
        "Havendo sinal, aplicam-se as regras do recibo de sinal e dos arts. 417 a 420 do Código Civil.",
        "A parte que descumprir este contrato responde pelas perdas e danos que causar, na forma da lei.",
    ])

    clausula(doc, 10, "Disposições gerais", [
        "Avisos entre as partes serão feitos por escrito, pelos e-mails ou telefones (mensagem) indicados na qualificação.",
        "Integram este contrato: Anexo I – Termo de Entrega; Anexo II – laudo de vistoria cautelar (se houver); "
        "[OUTROS: recibo de sinal, comprovantes de pagamento].",
        "Fica eleito o foro da comarca de [CIDADE/UF] para resolver questões deste contrato.",
    ])

    fechamento(doc, "E, por estarem justas e contratadas, as partes assinam este contrato em 2 (duas) vias de igual teor, na "
                    "presença de duas testemunhas.", [["VENDEDOR(A)", "[NOME COMPLETO] · CPF [000.000.000-00]"],
                      ["COMPRADOR(A)", "[NOME COMPLETO] · CPF [000.000.000-00]"],
                      ["TESTEMUNHA 1", "[NOME] · CPF [000.000.000-00]"],
                      ["TESTEMUNHA 2", "[NOME] · CPF [000.000.000-00]"]])
    para(doc, "Base legal citada: Código Civil – Lei 10.406/2002 (arts. 417 a 420, 441 a 446 e 492); Código de Trânsito Brasileiro – "
              "Lei 9.503/1997 (arts. 123, § 1º, e 134); Código de Defesa do Consumidor – Lei 8.078/1990 (vendas por loja/revenda).",
         size=8, italic=True, color=TEXTO3, before=8, after=0)

    # ---------------- Anexo I
    page_break(doc)
    titulo(doc, "Anexo I – Termo de entrega do veículo", "Parte integrante do contrato particular de compra e venda")
    form(doc, [
        [("Veículo / placa", "[MARCA MODELO] · [ABC1D23]"), ("RENAVAM", "[00000000000]")],
        [("Data da entrega", "[DD/MM/AAAA]"), ("HORA da entrega", "[00h00]")],
        [("Quilometragem no painel", "[000.000] km"), ("Combustível", "☐ res. ☐ 1/4 ☐ 1/2 ☐ 3/4 ☐ cheio")],
        [("Local da entrega", "[ENDEREÇO, CIDADE/UF]")],
    ], label_w=3.6)
    secao(doc, "Itens entregues")
    itens = [
        "Chave principal", "Chave reserva", "Manual do proprietário", "Manual / comprovantes de revisão",
        "Estepe", "Macaco e chave de roda", "Triângulo de sinalização",
        "Documento do veículo (CRLV) – impresso ou digital",
        "Documento de transferência assinado (CRV ou autorização eletrônica) [VERIFICAR]",
        "Comprovantes de quitação de débitos", "Notas fiscais de peças e serviços", "Cópia do laudo de vistoria cautelar",
        "[OUTRO ITEM]",
    ]
    table(doc, ["Item", "Entregue", "Não se aplica", "Observação"],
          [[i, "☐", "☐", ""] for i in itens], [7.2, 2.0, 2.0, 5.8], size=9, aligns=["l", "c", "c", "l"], row_h=0.6)
    secao(doc, "Estado do veículo na entrega")
    para(doc, "Avarias, luzes de alerta no painel ou observações: [DESCREVER OU \"NENHUMA\"]", align="j")
    for _ in range(2):
        ln = para(doc, "", before=10, after=0)
        p_border(ln, "bottom", sz=4, color=TEXTO3, space=1)
    para(doc, "Pelo presente termo, o(a) COMPRADOR(A) declara ter recebido o Veículo e os itens marcados acima na data, hora e "
              "quilometragem indicadas. A partir deste momento, o(a) COMPRADOR(A) assume a posse e a responsabilidade pelo Veículo, "
              "nos termos das Cláusulas 5 e 6 do contrato.", align="j", before=6, keep=True)
    fechamento(doc, None, [["VENDEDOR(A)", "[NOME COMPLETO] · CPF [000.000.000-00]"],
                      ["COMPRADOR(A)", "[NOME COMPLETO] · CPF [000.000.000-00]"]])
    doc.save(OUT / "02-contrato-particular-compra-e-venda-veiculo.docx")


# ======================================================================================
# 03 — ROTEIRO DE PERGUNTAS AO VENDEDOR
# ======================================================================================
PERGUNTAS = [
    ("Antes de ir ver o carro (por mensagem ou telefone)", [
        ("O carro está no seu nome? O documento está com você?",
         "Só negocie e pague a quem consta como proprietário no documento."),
        ("Por que você está vendendo?", "Resposta vaga ou que muda depois merece atenção."),
        ("Qual a quilometragem atual? Tem como comprovar?", "Notas de revisão e manual carimbado ajudam a conferir."),
        ("Há quanto tempo o carro é seu? Quantos donos ele já teve?", "Muitos donos em pouco tempo pede mais cuidado."),
        ("Já bateu, alagou, pegou fogo ou teve funilaria e pintura? Em quais partes?",
         "Anote a resposta: ela vai para as declarações do contrato."),
        ("Tem débitos (IPVA, licenciamento, multas) ou financiamento em aberto?",
         "Débito e gravame precisam estar resolvidos antes da transferência."),
        ("Pode me passar a placa e o RENAVAM para eu fazer as consultas?", "Quem não quer informar dá um sinal de alerta."),
        ("O carro já foi de leilão, recuperado de furto/roubo ou indenizado pelo seguro?",
         "Isso afeta o valor de revenda e às vezes a aceitação de seguro."),
        ("Você aceita que eu leve um vistoriador para uma vistoria cautelar antes de fechar?",
         "Recusar a vistoria é um dos sinais de alerta mais fortes."),
    ]),
    ("Na visita", [
        ("Posso ver o documento original e conferir com a placa e o chassi do carro?",
         "Os dados do documento devem bater com os do carro."),
        ("Tem manual, chave reserva e comprovantes das revisões?", "Faltando itens, use como argumento de negociação."),
        ("Quando foram trocados óleo, correia (ou corrente), pneus e bateria?",
         "O que está vencido entra no custo total da compra."),
        ("Alguma luz de alerta já acendeu no painel? Conhece algum barulho ou defeito?",
         "Anote: problema informado e aceito não é vício oculto."),
        ("Ar-condicionado, vidros, travas, faróis e multimídia funcionam?", "Teste cada um na hora, não aceite só a palavra."),
        ("O carro tem alguma modificação (suspensão, rodas, som, GNV, motor)? Está regularizada?",
         "Algumas alterações precisam constar no documento."),
        ("O carro já foi usado como táxi, transporte por aplicativo, locadora ou frota?", "Uso intenso muda o desgaste esperado."),
        ("Posso ver o motor ser ligado frio e fazer um test drive?", "Motor já aquecido pode esconder falhas na partida."),
        ("Tem recall pendente ou já feito? Tem o comprovante?", "Recall se consulta pelo chassi no canal oficial do fabricante."),
    ]),
    ("Antes de pagar", [
        ("Os débitos serão quitados antes da transferência ou descontados do preço? Pode ser por escrito?",
         "Combine no recibo ou no contrato."),
        ("O pagamento vai para a conta de quem está no documento?",
         "Nunca pague a intermediário, parente ou terceiro, nem com pedido urgente."),
        ("Quando e onde vamos assinar o documento de transferência?",
         "Confirme no órgão de trânsito do seu estado como a assinatura é feita (cartório ou eletrônica)."),
        ("Você vai fazer a comunicação de venda ao órgão de trânsito? Quando?",
         "Protege os dois lados sobre multas e responsabilidades."),
        ("O que entra no negócio: chaves, manual, estepe, acessórios?", "Liste tudo no termo de entrega."),
        ("Aceita registrar no contrato a data, a hora e a quilometragem da entrega?",
         "Isso define quem responde por multas e débitos."),
        ("Se o laudo cautelar apontar problema, podemos renegociar ou desfazer com devolução do sinal?",
         "Use a condição do recibo de sinal."),
    ]),
]


def roteiro():
    doc = novo_doc("Roteiro de perguntas ao vendedor")
    titulo(doc, "Roteiro de perguntas ao vendedor", "25 perguntas para fazer antes de fechar negócio, com espaço para anotar")
    form(doc, [
        [("Carro / versão", "[MARCA MODELO VERSÃO]"), ("Placa", "[ABC1D23]")],
        [("Vendedor", "[NOME]"), ("Contato", "[(00) 00000-0000]")],
        [("Onde viu o anúncio", "[LINK OU FONTE]"), ("Data", "[DD/MM/AAAA]")],
    ])
    spacer(doc, 4)
    para(doc, "Faça as perguntas com calma e anote as respostas com as palavras do vendedor. Respostas importantes (débitos, "
              "batidas, defeitos conhecidos) devem ir **por escrito** para o recibo ou o contrato.", align="j", size=9.5,
         color=TEXTO2)
    n = 1
    for grupo, itens in PERGUNTAS:
        secao(doc, f"{'ABC'[PERGUNTAS.index((grupo, itens))]}. {grupo}")
        t = table(doc, ["Nº", "Pergunta", "Resposta do vendedor"], [], [0.9, 8.6, 7.5], size=9.5,
                  aligns=["c", "l", "l"])
        for q, porque in itens:
            row = t.add_row()
            row_flags(row)
            row.height = Cm(1.45)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            c0, c1, c2 = row.cells
            for c, w in zip(row.cells, [0.9, 8.6, 7.5]):
                c.width = Cm(w)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            para(c0, str(n), size=10, bold=True, color=CARBONO, align="c", after=0)
            para(c1, q, size=9.5, bold=True, after=1)
            para(c1, "Por quê: " + porque, size=8.5, italic=True, color=TEXTO2, after=0)
            n += 1
    assert n - 1 == 25, n - 1
    spacer(doc, 8)
    boxed(doc, [
        "**Sinais de alerta nas respostas**",
        "Pressa para fechar · recusa da vistoria cautelar · pedido de pagamento para outra pessoa ou conta · preço muito abaixo "
        "da referência de mercado · documento que não está no nome de quem vende · história que muda a cada conversa · "
        "recusa em colocar as informações por escrito.",
        "Diante de qualquer um deles: pare, confira e, se preciso, chame um profissional antes de pagar.",
    ], fill=AMBAR_CLARO, size=9.5, border="E8D9A8", bar=AMBAR)
    doc.save(OUT / "03-roteiro-de-perguntas-ao-vendedor.docx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for fn in (recibo_sinal, contrato, roteiro):
        fn()
    for p in sorted(OUT.glob("*.docx")):
        print("ok", p.relative_to(BASE))


if __name__ == "__main__":
    main()
