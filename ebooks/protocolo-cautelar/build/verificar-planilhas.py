#!/usr/bin/env python3
"""Verifica as planilhas do Protocolo Cautelar.

1. Gera os 6 arquivos numa pasta temporaria, recalcula com LibreOffice
   (soffice headless) e procura erros (#REF!, #NAME?, #DIV/0!, #VALUE!, #N/A...).
2. Gera cenarios de teste do Registro de Apontamentos, recalcula e compara a
   classificacao da planilha com (a) o resultado esperado e (b) uma
   reimplementacao da matriz em Python.
3. Confere valores-chave dos demais arquivos contra calculo em Python.

Uso: python3 build/verificar-planilhas.py
Requer LibreOffice (soffice) no PATH.
"""
import importlib.util
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from statistics import median

from openpyxl import load_workbook

AQUI = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("bp", AQUI / "build-planilhas.py")
bp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bp)

ERROS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")


def recalc(paths, outdir):
    """Abre cada xlsx no LibreOffice (que recalcula ao carregar arquivo sem
    valores em cache) e regrava como xlsx em outdir."""
    prof = Path(tempfile.mkdtemp(prefix="lo-prof-"))
    cmd = ["soffice", f"-env:UserInstallation={prof.as_uri()}", "--headless", "--calc",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir)] + [str(p) for p in paths]
    subprocess.run(cmd, check=True, capture_output=True, timeout=300)
    shutil.rmtree(prof, ignore_errors=True)
    return [outdir / p.name for p in paths]


def varrer_erros(path):
    wbf = load_workbook(path)  # formulas
    wbv = load_workbook(path, data_only=True)
    achados, n_form = [], 0
    for ws in wbf.worksheets:
        wv = wbv[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    n_form += 1
                v = wv[c.coordinate].value
                if isinstance(v, str) and any(v.startswith(e) for e in ERROS):
                    achados.append(f"{ws.title}!{c.coordinate}={v}")
    return n_form, achados


def classif_python(niveis_main, niveis_anexos, idv_ids, n_total_main, anexos_on):
    """Reimplementacao da matriz. niveis_*: dict id->nivel (ausente = branco)."""
    brancos = n_total_main - len(niveis_main)
    todos = list(niveis_main.values())
    for ax, (ids_ax, niv_ax) in niveis_anexos.items():
        if anexos_on.get(ax) == "SIM":
            brancos += len(ids_ax) - len(niv_ax)
            todos += list(niv_ax.values())
    if brancos > 0:
        return f"INCOMPLETO – faltam {brancos} " + ("item" if brancos == 1 else "itens")
    if any(niveis_main.get(i) == "NV" for i in idv_ids):
        return "INCONCLUSIVO"
    if "N4" in todos:
        return "REPROVADO"
    if "N2" in todos or "N3" in todos:
        return "APROVADO COM APONTAMENTOS"
    return "APROVADO"


def main():
    d = bp.carregar()
    ids_main = [it["id"] for s in d["sistemas"] for it in s["itens"]]
    idv_ids = [it["id"] for s in d["sistemas"] if s["id"] == "IDV" for it in s["itens"]]
    anexos_ids = {ax["id"]: [it["id"] for it in ax["itens"]] for ax in d["anexos"]}
    base_ok = {i: ("N1" if k % 7 == 0 else "N0") for k, i in enumerate(ids_main)}
    base_ok["CAR-05"] = "N/A" if "CAR-05" in base_ok else base_ok.get("CAR-05")

    def com(**mods):
        x = dict(base_ok)
        x.update(mods)
        return x

    todos_moto = {i: "N0" for i in anexos_ids["MOTO"]}
    cenarios = [
        ("1 Só N0/N1/N/A", com(), {"MOTO": "NÃO", "UTL": "NÃO"}, "APROVADO"),
        ("2 Um N2 e um N3", com(**{"CAR-03": "N2", "RSP-02": "N3"}), {"MOTO": "NÃO", "UTL": "NÃO"},
         "APROVADO COM APONTAMENTOS"),
        ("3 N4 estrutural + N2", com(**{"EST-04": "N4", "CAR-03": "N2"}), {"MOTO": "NÃO", "UTL": "NÃO"}, "REPROVADO"),
        ("4 NV em IDV + N4", com(**{"IDV-05": "NV", "EST-04": "N4"}), {"MOTO": "NÃO", "UTL": "NÃO"}, "INCONCLUSIVO"),
        ("5 Três itens em branco", {k: v for k, v in com().items() if k not in ("DOC-04", "MOT-02", "INT-08")},
         {"MOTO": "NÃO", "UTL": "NÃO"}, "INCOMPLETO – faltam 3 itens"),
        ("6 Anexo MOTO ativado e vazio", com(), {"MOTO": "SIM", "UTL": "NÃO"},
         f"INCOMPLETO – faltam {len(anexos_ids['MOTO'])} itens"),
        ("7 Anexo MOTO preenchido com N4", {**com(), **todos_moto, anexos_ids["MOTO"][0]: "N4"},
         {"MOTO": "SIM", "UTL": "NÃO"}, "REPROVADO"),
        ("8 N4 só no anexo UTL desativado", {**com(), anexos_ids["UTL"][0]: "N4"}, {"MOTO": "NÃO", "UTL": "NÃO"},
         "APROVADO"),
    ]
    for nome, niv, _, _ in cenarios:
        for k in niv:
            assert k in ids_main or any(k in v for v in anexos_ids.values()), (nome, k)

    tmp = Path(tempfile.mkdtemp(prefix="pc-verif-"))
    src, out = tmp / "src", tmp / "out"
    src.mkdir()
    out.mkdir()
    falhas = 0

    # --- 1. arquivos entregues
    arquivos = [("01-registro-de-apontamentos.xlsx", bp.build_registro), ("02-mapa-de-pintura.xlsx", bp.build_pintura),
                ("03-precificacao.xlsx", bp.build_precificacao), ("04-controle-de-vistorias.xlsx", bp.build_controle),
                ("05-comparativo-anuncio-x-veiculo.xlsx", bp.build_comparativo), ("06-agenda-e-rota.xlsx", bp.build_agenda)]
    meta = {}
    for nome, fn in arquivos:
        meta[nome] = fn(src / nome)
    cen_paths = []
    for k, (nome, niv, ax, _) in enumerate(cenarios, 1):
        p = src / f"cenario-{k}.xlsx"
        bp.build_registro(p, niveis=niv, anexos=ax)
        cen_paths.append(p)
    rec = recalc([src / n for n, _ in arquivos] + cen_paths, out)

    print("== Erros de fórmula (após recálculo LibreOffice)")
    for p in rec:
        n, err = varrer_erros(p)
        status = "OK" if not err else "ERRO"
        falhas += bool(err)
        print(f"  [{status}] {p.name}: {n} fórmulas, {len(err)} erros {err[:5]}")

    # --- 2. cenarios
    print("== Cenários de classificação (Registro)")
    for (nome, niv, ax, esperado), p in zip(cenarios, rec[len(arquivos):]):
        wv = load_workbook(p, data_only=True)
        obtido = wv["Resumo"]["B6"].value
        vist = wv["Vistoria"]["G8"].value
        main_niv = {k: v for k, v in niv.items() if k in ids_main}
        anx = {a: (ids, {k: v for k, v in niv.items() if k in ids}) for a, ids in anexos_ids.items()}
        py = classif_python(main_niv, anx, idv_ids, len(ids_main), ax)
        ok = obtido == esperado == py == vist
        falhas += not ok
        print(f"  [{'OK' if ok else 'FALHA'}] {nome}: planilha='{obtido}' python='{py}' esperado='{esperado}'")

    # entregue (em branco) -> INCOMPLETO – faltam 95 itens
    wv = load_workbook(rec[0], data_only=True)
    esp = f"INCOMPLETO – faltam {len(ids_main)} itens"
    ok = wv["Resumo"]["B6"].value == esp
    falhas += not ok
    print(f"  [{'OK' if ok else 'FALHA'}] Arquivo entregue (em branco): '{wv['Resumo']['B6'].value}'")

    # --- 3. valores-chave
    print("== Valores-chave")
    # pintura
    wv = load_workbook(rec[1], data_only=True)
    ws = wv["Mapa"]
    m = meta["02-mapa-de-pintura.xlsx"]
    medias = []
    for r in range(m["first"], m["last"] + 1):
        leit = [ws.cell(row=r, column=c).value for c in range(3, 8)]
        leit = [x for x in leit if isinstance(x, (int, float))]
        if ws.cell(row=r, column=2).value != "Plástico" and leit:
            medias.append(sum(leit) / len(leit))
    ref = median(medias)
    ok = abs(ws["B7"].value - ref) < 1e-9
    esperados_status = {}
    for r in range(m["first"], m["last"] + 1):
        h = ws.cell(row=r, column=8).value
        mat = ws.cell(row=r, column=2).value
        if mat == "Plástico":
            e = "NÃO MEDIR (plástico)"
        elif not isinstance(h, (int, float)):
            e = "— sem leitura"
        else:
            rz = h / ref
            e = "Compatível" if rz <= 1.3 else ("Provável repintura" if rz <= 2 else "Provável repintura c/ massa/reparo")
        esperados_status[r] = e
        ok &= ws.cell(row=r, column=11).value == e
    falhas += not ok
    print(f"  [{'OK' if ok else 'FALHA'}] Pintura: referência={ref:.1f} µm; status de {len(esperados_status)} peças conferem; "
          f"porta tras. dir.='{ws['K18'].value}', lateral tras. dir.='{ws['K20'].value}'")

    # precificacao (tipo Basica, coluna B)
    wv = load_workbook(rec[2], data_only=True)
    ws = wv["Precificação"]
    m = meta["03-precificacao.xlsx"]
    cf = 800 + 150 + 300 + 120 + 90 + 150 + 300 + 4000 + 100
    ch = cf / 120
    custo = 1.0 * ch + 15 + 20 * 1.10 + 8
    pmin = custo / (1 - 0.06 - 0.035)
    psug = custo / (1 - 0.06 - 0.035 - 0.20)
    mc = psug * (1 - 0.06 - 0.035) - (15 + 20 * 1.10 + 8)
    import math
    pe = math.ceil(cf / mc)
    got = (ws.cell(row=m["ct_row"], column=2).value, ws.cell(row=m["pmin_row"], column=2).value,
           ws.cell(row=m["psug_row"], column=2).value, ws.cell(row=m["pe_row"], column=2).value)
    ok = all(abs(a - b) < 1e-6 for a, b in zip(got, (custo, pmin, psug, pe)))
    falhas += not ok
    print(f"  [{'OK' if ok else 'FALHA'}] Precificação (Básica): custo={got[0]:.2f} mín={got[1]:.2f} sugerido={got[2]:.2f} "
          f"PE={got[3]} (python: {custo:.2f}/{pmin:.2f}/{psug:.2f}/{pe})")

    # controle: mes 9/2026 -> exemplos 1-4
    wv = load_workbook(rec[3], data_only=True)
    ws = wv["Painel"]
    fat, n, tk, rec_, arec, com_ = (ws[c].value for c in ("B9", "B10", "B11", "B12", "B13", "B14"))
    exp = (350 + 180 + 250 + 150, 4, (930) / 4, 350 + 250 + 150, 180, 180 * 0.10 + 250 * 0.10)
    ok = all(abs(a - b) < 1e-6 for a, b in zip((fat, n, tk, rec_, arec, com_), exp))
    falhas += not ok
    print(f"  [{'OK' if ok else 'FALHA'}] Painel 09/2026: fat={fat} n={n} ticket={tk} recebido={rec_} a receber={arec} "
          f"comissões={com_} (esperado {exp})")

    # comparativo
    wv = load_workbook(rec[4], data_only=True)
    ws = wv["Anúncio x Veículo"]
    vals = [ws.cell(row=r, column=2).value for r in range(22, 26)]
    ok = vals == [2, 3, 2, 7]
    wc = wv["Comparar veículos"]
    ok2 = wc["B15"].value == 62800 and wc["D15"].value == 61200 and str(wc["B21"].value).startswith("Veículo C") \
        and str(wc["B22"].value).startswith("Veículo C")
    falhas += not (ok and ok2)
    print(f"  [{'OK' if ok and ok2 else 'FALHA'}] Comparativo: confere/diverge/n.inf/total={vals}; "
          f"menor custo='{wc['B21'].value}'; sem N4='{wc['B22'].value}'")

    # agenda
    wv = load_workbook(rec[5], data_only=True)
    ws = wv["Agenda"]
    tot_row = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "Somatório":
            tot_row = r
    g, h, i, j, k = (ws.cell(row=tot_row, column=c).value for c in range(7, 12))
    if hasattr(j, "total_seconds"):  # LibreOffice grava [h]:mm como duracao
        j = j.total_seconds() / 86400
    exp_min = (150 + 120 + 90 + 60) + (20 + 25 + 30 + 15)
    ok = g == 420 and h == 90 and i == 41 and abs(j * 1440 - exp_min) < 1e-6 and k == 4
    falhas += not ok
    print(f"  [{'OK' if ok else 'FALHA'}] Agenda: vistoria={g} min, desloc={h} min, {i} km, total={j * 24:.2f} h, "
          f"{k} vistorias")

    print(f"\nPasta de trabalho: {tmp}")
    print("RESULTADO:", "TUDO OK" if falhas == 0 else f"{falhas} FALHA(S)")
    sys.exit(1 if falhas else 0)


if __name__ == "__main__":
    main()
