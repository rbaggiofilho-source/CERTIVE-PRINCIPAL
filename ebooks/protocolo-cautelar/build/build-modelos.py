#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build-modelos.py — gera os modelos editáveis (.docx) para o vistoriador cautelar.

Saída: ebooks/protocolo-cautelar/03-materiais/modelos/*.docx
Fonte de dados: 01-conteudo/dados/sistemas.json (níveis, classificação e os 10 sistemas).

Os modelos são do VISTORIADOR: visual neutro, sem a marca do kit.
Placeholders entre colchetes: [NOME DA EMPRESA], [CNPJ], [CIDADE/UF]...
Uso: python3 build/build-modelos.py
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = Path(__file__).resolve().parent.parent
DADOS = json.loads((BASE / "01-conteudo/dados/sistemas.json").read_text(encoding="utf-8"))
OUT = BASE / "03-materiais/modelos"

# Paleta (tema.css) — uso sóbrio: carbono nos títulos, âmbar só no fio.
CARBONO = "0F1318"
AMBAR = "F2A900"
AMBAR_ESC = "B87F00"
LINHA = "C9CED4"
PAPEL2 = "F2F3F5"
TEXTO = "1A1F26"
TEXTO2 = "4E5864"
TEXTO3 = "7A8591"
NIVEL_COR = {n["id"]: n["cor"].lstrip("#") for n in DADOS["niveis"]}
CLASS_COR = {c["id"]: c["cor"].lstrip("#") for c in DADOS["classificacao"]}

AVISO = "Modelo de referência. Revise com um advogado de sua confiança antes de usar."
LARGURA = 17.0  # cm úteis (A4 21 cm − 2 × 2 cm)

# --------------------------------------------------------------------------------------
# Utilidades XML (ordem dos elementos respeita o schema OOXML)
# --------------------------------------------------------------------------------------
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
             "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
             "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
             "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def _local(el):
    return el.tag.split("}")[-1]


def insert_ordered(parent, el, order):
    tag = _local(el)
    for old in parent.findall(qn("w:" + tag)):
        parent.remove(old)
    idx = order.index(tag)
    for child in parent:
        name = _local(child)
        if name in order and order.index(name) > idx:
            child.addprevious(el)
            return el
    parent.append(el)
    return el


def rgb(h):
    return RGBColor.from_string(h)


def cm2tw(v):
    return int(round(v / 2.54 * 1440))


def set_fonts(rpr_parent_el, name="Calibri"):
    rpr = rpr_parent_el.get_or_add_rPr()
    for old in rpr.findall(qn("w:rFonts")):
        rpr.remove(old)
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    rpr.insert(0, rf)


def p_border(p, side="bottom", sz=6, color=AMBAR, space=2, val="single"):
    pPr = p._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = insert_ordered(pPr, OxmlElement("w:pBdr"), PPR_ORDER)
    el = OxmlElement("w:" + side)
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    # ordem dentro de pBdr: top, left, bottom, right, between, bar
    order = ["top", "left", "bottom", "right", "between", "bar"]
    insert_ordered(bdr, el, order)


def p_shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pPr, shd, PPR_ORDER)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(tcPr, shd, TCPR_ORDER)


def cell_borders(cell, **sides):
    """sides: top/left/bottom/right = (val, sz, color)"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        if side in sides:
            val, sz, color = sides[side]
            e = OxmlElement("w:" + side)
            e.set(qn("w:val"), val)
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
            b.append(e)
    insert_ordered(tcPr, b, TCPR_ORDER)


def tbl_borders(tbl, val="single", sz=4, color=LINHA, inside=True):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side)
        if side.startswith("inside") and not inside:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), val)
            if val != "nil":
                e.set(qn("w:sz"), str(sz))
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), color)
        b.append(e)
    insert_ordered(tblPr, b, TBLPR_ORDER)


def tbl_cell_margins(tbl, top=0.06, left=0.15, bottom=0.06, right=0.15):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + side)
        e.set(qn("w:w"), str(cm2tw(v)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    insert_ordered(tblPr, m, TBLPR_ORDER)


def set_widths(tbl, widths):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(sum(cm2tw(x) for x in widths)))
    w.set(qn("w:type"), "dxa")
    insert_ordered(tblPr, w, TBLPR_ORDER)
    for gc, wd in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(cm2tw(wd)))
    for row in tbl.rows:
        for c, wd in zip(row.cells, widths):
            c.width = Cm(wd)


def keep_table(t):
    """Mantém a tabela inteira na mesma página (keep-with-next em todas as linhas menos a última)."""
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    return t


def row_flags(row, header=False, cant_split=True):
    trPr = row._tr.get_or_add_trPr()
    if cant_split:
        trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def add_field(p, instr, size=8, color=TEXTO3):
    def run():
        r = p.add_run()
        r.font.size = Pt(size)
        r.font.color.rgb = rgb(color)
        return r
    r = run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r = run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    r = run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
    r = run(); r.text = "1"
    r = run()
    fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)


# --------------------------------------------------------------------------------------
# Texto rico: **negrito**, [PLACEHOLDER] destacado, ☐ com fonte de símbolo
# --------------------------------------------------------------------------------------
TOK = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]|☐)")


def add_rich(p, text, size=None, color=None, bold=False, italic=False):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2]); r.bold = True
            if color:
                r.font.color.rgb = rgb(color)
        elif part.startswith("[") and part.endswith("]"):
            r = p.add_run(part); r.font.color.rgb = rgb(AMBAR_ESC)
            if bold:
                r.bold = True
        elif part == "☐":
            r = p.add_run(part); set_fonts(r._r, "Segoe UI Symbol")
            r.font.color.rgb = rgb(TEXTO)
        else:
            r = p.add_run(part)
            if bold:
                r.bold = True
            if color:
                r.font.color.rgb = rgb(color)
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
    return p


def _target_par(container):
    """Em células, reaproveita o primeiro parágrafo vazio."""
    if hasattr(container, "_tc"):
        ps = container.paragraphs
        if len(ps) == 1 and not ps[0].runs and not ps[0].text:
            return ps[0]
    return container.add_paragraph()


def para(container, text="", size=None, color=None, bold=False, italic=False, align=None,
         after=None, before=None, indent=None, hanging=None, keep=False, style=None):
    p = _target_par(container)
    if style:
        p.style = style
    add_rich(p, text, size=size, color=color, bold=bold, italic=italic)
    pf = p.paragraph_format
    if align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "r":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if after is not None:
        pf.space_after = Pt(after)
    if before is not None:
        pf.space_before = Pt(before)
    if indent is not None:
        pf.left_indent = Cm(indent)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    if keep:
        pf.keep_with_next = True
    return p


def bullets(container, items, size=None, align="j"):
    for it in items:
        p = para(container, it, size=size, align=align, after=2, style="List Bullet")
    return p


def numbered(container, items, prefix="", size=None, start=1, indent=0.9):
    for i, it in enumerate(items, start):
        para(container, f"{prefix}{i}.\t{it}", size=size, align="j", indent=indent, hanging=indent, after=3)


def h1(doc, text):
    return doc.add_heading(text, level=1)


def h2(doc, text):
    return doc.add_heading(text, level=2)


def spacer(container, pt=4):
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(); r.font.size = Pt(pt)
    p.paragraph_format.line_spacing = Pt(pt)
    return p


def page_break(doc):
    doc.add_page_break()


# --------------------------------------------------------------------------------------
# Tabelas
# --------------------------------------------------------------------------------------
def table(container, headers, rows, widths, size=9, header_fill=PAPEL2, row_h=None, aligns=None,
          repeat_header=True, first_bold=False):
    n = len(widths)
    nrows = len(rows) + (1 if headers else 0)
    t = container.add_table(rows=nrows, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, widths)
    tbl_borders(t)
    tbl_cell_margins(t)
    aligns = aligns or ["l"] * n
    r0 = 0
    if headers:
        row = t.rows[0]
        row_flags(row, header=repeat_header)
        for i, htxt in enumerate(headers):
            c = row.cells[i]
            shade(c, header_fill)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para(c, htxt, size=size - 0.5, bold=True, color=CARBONO, after=0,
                 align={"c": "c", "r": "r"}.get(aligns[i]))
        r0 = 1
    for ri, data in enumerate(rows):
        row = t.rows[r0 + ri]
        row_flags(row)
        if row_h:
            row.height = Cm(row_h)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i in range(n):
            c = row.cells[i]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            txt = data[i] if i < len(data) else ""
            para(c, txt, size=size, after=0, bold=(first_bold and i == 0),
                 align={"c": "c", "r": "r"}.get(aligns[i]))
    return t


def form(container, rows, label_w=3.3, total=LARGURA, size=9.5, h=0.72, cols=2):
    """rows: lista de linhas; cada linha é lista de (rótulo, valor). 1 par = valor ocupa a linha toda."""
    value_w = (total - label_w * cols) / cols
    widths = [label_w, value_w] * cols
    t = container.add_table(rows=len(rows), cols=2 * cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, widths)
    tbl_borders(t)
    tbl_cell_margins(t)
    for ri, pairs in enumerate(rows):
        row = t.rows[ri]
        row_flags(row)
        row.height = Cm(h)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        cells = row.cells
        if len(pairs) == 1:
            merged = cells[1].merge(cells[2 * cols - 1])
            cells = [cells[0], merged]
        for pi, (lab, val) in enumerate(pairs):
            lc, vc = cells[2 * pi], cells[2 * pi + 1]
            shade(lc, PAPEL2)
            for c in (lc, vc):
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            para(lc, lab, size=size - 1, bold=True, color=TEXTO2, after=0)
            para(vc, val, size=size, after=0)
    return t


def boxed(container, lines, fill=PAPEL2, size=10, border=LINHA, width=LARGURA, bar=None):
    """Caixa de texto (tabela 1×1). lines: lista de strings (cada uma vira parágrafo)."""
    t = container.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, [width])
    tbl_borders(t, color=border)
    tbl_cell_margins(t, top=0.2, left=0.3, bottom=0.2, right=0.3)
    c = t.cell(0, 0)
    shade(c, fill)
    if bar:
        cell_borders(c, left=("single", 24, bar), top=("single", 4, border),
                     bottom=("single", 4, border), right=("single", 4, border))
    row_flags(t.rows[0])
    for ln in lines:
        para(c, ln, size=size, after=3)
    return t


def assinaturas(container, blocos, cols=2, gap=1.0):
    """blocos: lista de listas de linhas (primeira linha = papel)."""
    if cols == 1:
        t = container.add_table(rows=1, cols=3)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_widths(t, [4.5, 8.0, 4.5])
        tbl_borders(t, val="nil")
        row_flags(t.rows[0])
        cell = t.cell(0, 1)
        p = para(cell, "", after=0)
        p.paragraph_format.space_before = Pt(30)
        p2 = para(cell, blocos[0][0], size=9, bold=True, color=CARBONO, after=0, align="c")
        p_border(p2, "top", sz=6, color=TEXTO, space=4)
        for ln in blocos[0][1:]:
            para(cell, ln, size=8.5, color=TEXTO2, after=0, align="c")
        return t
    rows = (len(blocos) + cols - 1) // cols
    w = (LARGURA - gap * (cols - 1)) / cols
    widths = []
    for i in range(cols):
        widths.append(w)
        if i < cols - 1:
            widths.append(gap)
    t = container.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, widths)
    tbl_borders(t, val="nil")
    for bi, bloco in enumerate(blocos):
        r, c = divmod(bi, cols)
        row_flags(t.rows[r])
        cell = t.cell(r, c * 2)
        p = para(cell, "", after=0)
        p.paragraph_format.space_before = Pt(30)
        p2 = para(cell, bloco[0], size=9, bold=True, color=CARBONO, after=0, align="c")
        p_border(p2, "top", sz=6, color=TEXTO, space=4)
        for ln in bloco[1:]:
            para(cell, ln, size=8.5, color=TEXTO2, after=0, align="c")
    return t


# --------------------------------------------------------------------------------------
# Documento base, timbre, título, rodapé
# --------------------------------------------------------------------------------------
def novo_doc(aviso=True, rodape_esq="[NOME DA EMPRESA]", cabecalho=None):
    doc = Document()
    # settings.xml do template padrão traz <w:zoom> sem w:percent (inválido no schema)
    for z in doc.settings.element.findall(qn("w:zoom")):
        z.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.0))
    sec.header_distance = Cm(0.9)
    sec.footer_distance = Cm(0.9)

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = rgb(TEXTO)
    set_fonts(st.element)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.12

    for name, size, before, after in (("Heading 1", 13.5, 14, 6), ("Heading 2", 11, 10, 4), ("Heading 3", 10.5, 8, 3)):
        s = doc.styles[name]
        set_fonts(s.element)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = False
        s.font.color.rgb = rgb(CARBONO)
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        if name == "Heading 1":
            ppr = s.element.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            e = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "3"), ("w:color", AMBAR)):
                e.set(qn(k), v)
            bdr.append(e)
            insert_ordered(ppr, bdr, PPR_ORDER)
    lb = doc.styles["List Bullet"]
    set_fonts(lb.element)
    lb.paragraph_format.left_indent = Cm(0.6)
    lb.paragraph_format.first_line_indent = Cm(-0.4)

    # Rodapé: aviso (opcional) + nome à esquerda e "Página X de Y" à direita
    def monta_rodape(ft):
        p = ft.paragraphs[0]
        p.style = doc.styles["Normal"]
        p_border(p, "top", sz=4, color=LINHA, space=4)
        if aviso:
            r = p.add_run(AVISO)
            r.italic = True
            r.font.size = Pt(7.5)
            r.font.color.rgb = rgb(TEXTO3)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            p2 = ft.add_paragraph()
            p2.style = doc.styles["Normal"]
        else:
            p2 = p
        p2.paragraph_format.tab_stops.add_tab_stop(Cm(LARGURA), WD_TAB_ALIGNMENT.RIGHT)
        p2.paragraph_format.space_after = Pt(0)
        add_rich(p2, rodape_esq, size=8, color=TEXTO3)
        r = p2.add_run("\tPágina "); r.font.size = Pt(8); r.font.color.rgb = rgb(TEXTO3)
        add_field(p2, "PAGE")
        r = p2.add_run(" de "); r.font.size = Pt(8); r.font.color.rgb = rgb(TEXTO3)
        add_field(p2, "NUMPAGES")

    if cabecalho:
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_rich(hp, cabecalho, size=8, color=TEXTO3)
        p_border(hp, "bottom", sz=4, color=LINHA, space=3)
        monta_rodape(sec.first_page_footer)
    monta_rodape(sec.footer)
    return doc


def timbre(doc, compacto=False):
    """Cabeçalho do vistoriador: área de logo tracejada + dados."""
    t = doc.add_table(rows=1, cols=2)
    set_widths(t, [5.0, 12.0])
    tbl_borders(t, val="nil")
    left, right = t.cell(0, 0), t.cell(0, 1)
    logo = left.add_table(rows=1, cols=1)
    set_widths(logo, [4.6])
    tbl_borders(logo, val="dashed", sz=8, color="9AA3AD")
    lr = logo.rows[0]
    lr.height = Cm(2.0 if compacto else 2.4)
    lr.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    lc = logo.cell(0, 0)
    lc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para(lc, "SUA LOGO", size=11, bold=True, color=TEXTO3, align="c", after=0)
    para(lc, "(substitua por sua imagem)", size=7, color=TEXTO3, align="c", after=0)
    left._tc.remove(left.paragraphs[0]._p)  # parágrafo vazio acima da tabela aninhada
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para(right, "[NOME DA EMPRESA / NOME DO VISTORIADOR]", size=13, bold=True, color=CARBONO, align="r", after=1)
    para(right, "CNPJ/CPF: [00.000.000/0000-00]  ·  [REGISTRO PROFISSIONAL, SE HOUVER]", size=8.5, color=TEXTO2, align="r", after=0)
    para(right, "[ENDEREÇO COMPLETO] · [CIDADE/UF] · CEP [00000-000]", size=8.5, color=TEXTO2, align="r", after=0)
    para(right, "WhatsApp [(00) 00000-0000] · [E-MAIL] · [SITE OU PERFIL]", size=8.5, color=TEXTO2, align="r", after=0)
    return t


def titulo(doc, text, sub=None):
    p = para(doc, text.upper(), size=17, bold=True, color=CARBONO, before=10, after=2 if sub else 6)
    if sub:
        p = para(doc, sub, size=9.5, color=TEXTO2, after=6)
    p_border(p, "bottom", sz=8, color=AMBAR, space=4)
    spacer(doc, 6)
    return p


def clausula(doc, n, nome, itens):
    """Cláusula numerada: itens str | (str, [subitens])."""
    h2(doc, f"CLÁUSULA {n}ª – {nome.upper()}")
    for i, it in enumerate(itens, 1):
        subs = []
        if isinstance(it, tuple):
            it, subs = it
        para(doc, f"{n}.{i}.\t{it}", align="j", indent=1.0, hanging=1.0, after=3)
        for j, s in enumerate(subs):
            para(doc, f"{chr(97 + j)})\t{s}", align="j", indent=1.6, hanging=0.6, after=2)


def partes_contrato(doc, a_nome, b_nome, b_desc):
    para(doc, f"Pelo presente instrumento particular, as partes abaixo qualificadas:", align="j")
    para(doc, f"**{a_nome}:** [NOME DA EMPRESA], inscrita no CNPJ sob o nº [00.000.000/0000-00] (ou [NOME COMPLETO], "
              f"CPF nº [000.000.000-00]), com endereço em [ENDEREÇO COMPLETO], [CIDADE/UF], neste ato representada por "
              f"[NOME DO REPRESENTANTE], [CARGO], doravante denominada **{a_nome}**;", align="j")
    para(doc, f"**{b_nome}:** {b_desc}, doravante denominado(a) **{b_nome}**;", align="j")
    para(doc, "têm entre si justo e acordado o presente contrato, que se regerá pelas cláusulas e condições seguintes.",
         align="j", after=6)


def local_data(doc):
    para(doc, "[CIDADE/UF], [DIA] de [MÊS] de [ANO].", align="r", before=10, after=4)


# ======================================================================================
# 01 — LAUDO
# ======================================================================================
def laudo():
    doc = novo_doc(rodape_esq="[NOME DA EMPRESA]  ·  Laudo nº [0000/AAAA]",
                   cabecalho="[NOME DA EMPRESA]  ·  Laudo de vistoria cautelar nº [0000/AAAA]  ·  Placa [ABC1D23]")
    timbre(doc)
    titulo(doc, "Laudo de Vistoria Cautelar", "Exame técnico visual e não destrutivo de veículo usado, para apoio à decisão de compra ou venda.")
    form(doc, [
        [("Laudo nº", "[0000/AAAA]"), ("Data", "[DD/MM/AAAA]")],
        [("Início / término", "[00h00] / [00h00]"), ("Validade sugerida", "[30] dias (ver item 9.10)")],
        [("Local da vistoria", "[ENDEREÇO / PÁTIO / LOJA] – [CIDADE/UF]")],
    ])

    h1(doc, "1. Dados do solicitante")
    form(doc, [
        [("Nome / razão social", "[NOME DO SOLICITANTE]")],
        [("CPF / CNPJ", "[000.000.000-00]"), ("Telefone", "[(00) 00000-0000]")],
        [("E-mail", "[E-MAIL]"), ("Relação com o veículo", "☐ comprador  ☐ vendedor  ☐ loja  ☐ outro")],
        [("Proprietário/possuidor presente", "[NOME] – autorização assinada: ☐ sim  ☐ não")],
    ])

    h1(doc, "2. Dados do veículo e condições da vistoria")
    form(doc, [
        [("Placa", "[ABC1D23]"), ("Marca / modelo", "[MARCA/MODELO]")],
        [("Versão", "[VERSÃO]"), ("Ano fab. / mod.", "[0000/0000]")],
        [("Cor (documento)", "[COR]"), ("Cor (observada)", "[COR]")],
        [("Combustível", "[COMBUSTÍVEL]"), ("Câmbio", "☐ manual  ☐ automático  ☐ CVT")],
        [("Chassi (VIN)", "[17 CARACTERES]"), ("RENAVAM", "[00000000000]")],
        [("Nº do motor", "[NÚMERO]"), ("Hodômetro", "[000.000] km")],
        [("Iluminação", "☐ luz natural  ☐ coberto  ☐ artificial"), ("Clima", "☐ seco  ☐ chuva  ☐ outro")],
        [("Veículo", "☐ limpo  ☐ sujo  ☐ molhado"), ("Motor na chegada", "☐ frio  ☐ morno/quente")],
        [("Recursos usados", "☐ medidor de espessura  ☐ scanner OBD2  ☐ testador de bateria  ☐ lanterna/espelho  ☐ elevador  ☐ test drive")],
    ])

    h1(doc, "3. Resultado da vistoria")
    para(doc, "Marque **uma** classificação. A regra é aplicada nesta ordem, de cima para baixo (ver seção 10).", size=9.5, color=TEXTO2, keep=True)
    t = doc.add_table(rows=len(DADOS["classificacao"]), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, [0.35, 5.6, 11.05])
    tbl_borders(t)
    tbl_cell_margins(t, top=0.12, bottom=0.12)
    for i, c in enumerate(DADOS["classificacao"]):
        row = t.rows[i]
        row_flags(row)
        row.height = Cm(0.95)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        shade(row.cells[0], c["cor"].lstrip("#"))
        for cc in row.cells:
            cc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para(row.cells[1], f"☐  {c['id']}", size=11, bold=True, color=CARBONO, after=0)
        para(row.cells[2], c["regra"], size=9.5, color=TEXTO2, after=0)
    keep_table(t)
    spacer(doc)
    form(doc, [
        [("Motivo (se INCONCLUSIVO)", "[DESCREVER O ITEM NÃO VERIFICÁVEL E O MOTIVO]")],
        [("Itens N3 em destaque (Atenção)", "[CÓDIGOS E DESCRIÇÃO RESUMIDA]")],
        [("Recomendação", "[EX.: AVALIAÇÃO COMPLEMENTAR EM OFICINA / PERÍCIA OFICIAL / NOVA CONSULTA ANTES DO FECHAMENTO]")],
    ], label_w=4.2, cols=1, h=0.9)
    spacer(doc)
    boxed(doc, ["**Importante:** a classificação refere-se ao resultado do exame técnico visual realizado neste veículo, "
                "na data e nas condições descritas. \"Aprovado\" não aprova a compra, não garante o veículo e não substitui "
                "a vistoria oficial exigida pelo órgão de trânsito."], size=9, bar=AMBAR)

    h1(doc, "4. Resumo dos achados")
    para(doc, "Liste aqui todos os achados N2, N3 e N4, e os N1 que mereçam menção. Use o código do item (seção 5) e o número da foto (seção 7).",
         size=9.5, color=TEXTO2)
    table(doc, ["Nº", "Código", "Descrição objetiva do achado", "Nível", "Foto(s)"],
          [[str(i), "", "", "", ""] for i in range(1, 9)],
          [1.0, 2.0, 10.2, 1.7, 2.1], row_h=0.75, aligns=["c", "c", "l", "c", "c"])
    spacer(doc)
    para(doc, "Escala de níveis", size=10, bold=True, color=CARBONO, keep=True, after=2)
    t = table(doc, None, [[n["id"], n["nome"], n["definicao"]] for n in DADOS["niveis"]],
              [1.2, 3.6, 12.2], size=9, first_bold=True, aligns=["c", "l", "l"])
    for i, n in enumerate(DADOS["niveis"]):
        c = t.rows[i].cells[0]
        shade(c, n["cor"].lstrip("#"))
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = rgb("FFFFFF")
    keep_table(t)
    para(doc, "NA = não aplicável ao veículo.  NV = não verificado (informar o motivo na observação).", size=8.5, color=TEXTO2, before=3)

    h1(doc, "5. Verificação por sistema")
    para(doc, "Registre para cada item o nível (N0 a N4, NA ou NV) e uma observação objetiva: o que foi visto, onde e em qual foto. "
              "Itens marcados com ◆ são críticos: se não puderem ser verificados na identificação veicular, o resultado é INCONCLUSIVO.",
         size=9.5, color=TEXTO2, align="j")
    for s in DADOS["sistemas"]:
        h2(doc, f"5.{s['n']}  {s['id']} — {s['nome']}")
        rows = []
        for it in s["itens"]:
            nome = it["item"] + (" ◆" if it.get("critico") else "")
            rows.append([it["id"], nome, "", "", ""])
        table(doc, ["Código", "Item verificado", "Nível", "Observação", "Foto"], rows,
              [1.75, 6.6, 1.3, 6.15, 1.2], size=8.5, row_h=0.62, aligns=["l", "l", "c", "l", "c"])

    page_break(doc)
    h1(doc, "6. Mapa de pintura")
    para(doc, "Medições em micrômetros (µm), no mínimo três pontos por peça. Valores de referência são **orientativos** e variam "
              "por fabricante, modelo, ano e equipamento: compare peças vizinhas do mesmo veículo e considere o conjunto de evidências "
              "(tonalidade, textura, overspray, parafusos). Peças plásticas ou de materiais não metálicos exigem equipamento compatível; "
              "se não medidas, registre NV.", size=9.5, color=TEXTO2, align="j")
    form(doc, [[("Equipamento", "[MARCA/MODELO DO MEDIDOR – OPCIONAL]"), ("Faixa de referência adotada", "[000–000] µm")]])
    spacer(doc)
    pecas = ["Capô", "Teto", "Para-lama dianteiro esquerdo", "Para-lama dianteiro direito", "Porta dianteira esquerda",
             "Porta dianteira direita", "Porta traseira esquerda", "Porta traseira direita", "Lateral traseira esquerda",
             "Lateral traseira direita", "Tampa traseira / porta-malas", "Coluna A esquerda / direita",
             "Coluna B esquerda / direita", "Soleira esquerda / direita", "Para-choque dianteiro", "Para-choque traseiro",
             "Retrovisores", "[OUTRA PEÇA]"]
    table(doc, ["Peça", "Leitura 1", "Leitura 2", "Leitura 3", "Status", "Nível"],
          [[p, "", "", "", "", ""] for p in pecas],
          [5.4, 1.75, 1.75, 1.75, 4.8, 1.55], size=8.5, row_h=0.6, aligns=["l", "c", "c", "c", "l", "c"])
    para(doc, "Status sugerido: Original · Repintura · Massa/reparo · Peça substituída · Não medido (NV).", size=8.5, color=TEXTO2, before=3)

    h1(doc, "7. Relação de fotos")
    para(doc, "Numere as fotos na ordem do relatório fotográfico. Placas e dados de terceiros visíveis em segundo plano devem ser ocultados.",
         size=9.5, color=TEXTO2)
    table(doc, ["Nº", "Descrição da foto", "Código do item", "Arquivo / página"],
          [[f"{i:02d}", "", "", ""] for i in range(1, 25)],
          [1.1, 9.4, 2.6, 3.9], size=8.5, row_h=0.55, aligns=["c", "l", "c", "l"])

    h1(doc, "8. Consultas realizadas")
    para(doc, "As informações abaixo refletem o conteúdo das bases consultadas na data e hora indicadas. Anexe os comprovantes.",
         size=9.5, color=TEXTO2)
    consultas = ["Restrições (judicial, administrativa, roubo/furto)", "Gravame / alienação fiduciária", "Leilão e sinistro",
                 "Débitos (IPVA, licenciamento, multas)", "Recall", "Histórico de quilometragem", "[OUTRA CONSULTA]"]
    table(doc, ["Consulta", "Fonte / provedor", "Data e hora", "Resultado resumido", "Anexo"],
          [[c, "", "", "", "☐"] for c in consultas],
          [4.6, 3.4, 2.4, 5.0, 1.6], size=8.5, row_h=0.7, aligns=["l", "l", "c", "l", "c"])

    page_break(doc)
    h1(doc, "9. Metodologia e limites")
    limites = [
        ("Natureza do exame", "Esta vistoria cautelar é um exame técnico **visual, sensitivo e não destrutivo**, realizado a pedido do solicitante, com "
         "o auxílio dos instrumentos portáteis indicados na seção 2. Não houve desmontagem de peças, remoção de componentes fixados por "
         "ferramentas, abertura de motor, câmbio ou outros conjuntos, nem ensaios de laboratório."),
        ("Estado e momento", "O resultado retrata exclusivamente as condições observadas **na data, no horário e no local indicados**, no estado "
         "em que o veículo foi apresentado (limpeza, iluminação, clima, temperatura do motor). Uso, desgaste, intervenções e eventos "
         "posteriores não são alcançados por este laudo."),
        ("Itens não acessíveis", "Áreas ocultas por revestimentos, forrações, protetores, carenagens, sujeira, adesivos, películas, "
         "acessórios ou carga, e as que exigiriam elevador ou desmontagem, podem não ter sido examinadas no todo ou em parte. Itens não "
         "verificados estão identificados como **NV**, com o motivo."),
        ("Vícios ocultos", "O exame não garante a inexistência de vícios ocultos, defeitos intermitentes ou falhas que só se manifestam em "
         "uso prolongado ou em condições específicas (carga, temperatura, tipo de via), nem estima a vida útil futura de componentes."),
        ("Não é vistoria oficial nem perícia", "Este laudo **não é** a vistoria de identificação veicular exigida pelo órgão executivo de "
         "trânsito para transferência, regularização ou outros atos, realizada pelo próprio órgão ou por empresa por ele credenciada, nos "
         "termos do CTB – Lei 9.503/1997 e das Resoluções do CONTRAN [VERIFICAR resolução vigente]. Também não é perícia oficial nem "
         "perícia judicial, e não substitui nenhuma delas."),
        ("Identificação veicular", "O vistoriador confere a presença, a legibilidade e a coerência dos elementos de identificação com o "
         "documento apresentado. Divergências e indícios são **registrados de forma descritiva**; o vistoriador cautelar não conclui sobre "
         "adulteração, o que depende de perícia oficial. Havendo divergência, recomenda-se não concluir o negócio e procurar a autoridade competente."),
        ("Consultas", "Os resultados de consultas (restrições, gravame, leilão, sinistro, débitos, recall e outras) refletem o conteúdo das "
         "bases consultadas na data e hora registradas, **não são produzidos pelo vistoriador** e dependem da atualização e da abrangência "
         "das próprias fontes. Recomenda-se repetir as consultas imediatamente antes da conclusão do negócio."),
        ("Valores de referência", "Espessura de pintura, tensão de bateria, profundidade de sulco e demais valores são **orientativos** e "
         "variam por fabricante, modelo, ano e condições de medição. A conclusão considera o conjunto de evidências, não uma medição isolada."),
        ("Teste de rodagem", "Quando realizado e autorizado, o teste de rodagem foi curto, conduzido por pessoa habilitada, em trajeto "
         "seguro e dentro dos limites legais, e não reproduz todas as condições de uso do veículo."),
        ("Validade da informação", "A informação deste laudo perde atualidade com o tempo e o uso. Para decisões tomadas após [30] dias da "
         "vistoria, ou após qualquer evento relevante (colisão, alagamento, reparo, troca de peças), recomenda-se nova vistoria."),
        ("Finalidade e uso", "Este laudo destina-se ao solicitante identificado, para apoiar sua decisão. Não é garantia, certificado de "
         "qualidade, aprovação da compra nem avaliação de preço de mercado. A decisão de comprar, vender ou recusar é exclusiva do solicitante."),
        ("Integridade", "Este laudo contém [00] páginas e o relatório fotográfico com [00] fotos. Só é válido completo, sem alterações e com a "
         "assinatura do vistoriador. Correções geram nova versão numerada, com referência à anterior."),
    ]
    for i, (tit, txt) in enumerate(limites, 1):
        para(doc, f"9.{i}.\t**{tit}.** {txt}", align="j", indent=1.0, hanging=1.0, after=4)

    h1(doc, "10. Política de classificação")
    para(doc, "Cada item recebe um nível. O resultado final segue a matriz abaixo, **aplicada nesta ordem**:", align="j")
    table(doc, ["Ordem", "Condição", "Resultado"],
          [["1", "Algum item de identificação veicular não pôde ser verificado", "INCONCLUSIVO"],
           ["2", "Pelo menos um achado N4", "REPROVADO"],
           ["3", "Pelo menos um achado N2 ou N3 (N3 em destaque \"Atenção\")", "APROVADO COM APONTAMENTOS"],
           ["4", "Somente achados N0 ou N1", "APROVADO"]],
          [1.5, 9.5, 6.0], size=9, aligns=["c", "l", "l"])
    para(doc, "O laudo INCONCLUSIVO não conclui sobre o veículo: descreve o motivo e o que seria necessário para concluir.",
         size=9, color=TEXTO2, before=3)

    h1(doc, "11. Declaração e assinaturas")
    para(doc, "Declaro que realizei pessoalmente a vistoria descrita neste laudo, com independência técnica, segundo a metodologia e os "
              "limites da seção 9, e que a classificação reflete exclusivamente as condições observadas.", align="j")
    para(doc, "O solicitante declara ter recebido este laudo, lido a seção 9 (Metodologia e limites) e tido a oportunidade de "
              "esclarecer dúvidas sobre cada achado.", align="j")
    local_data(doc)
    assinaturas(doc, [["[NOME DO VISTORIADOR]", "Vistoriador responsável", "CPF [000.000.000-00]"],
                      ["[NOME DO SOLICITANTE]", "Solicitante – ciente do laudo", "CPF [000.000.000-00]"]])
    doc.save(OUT / "01-laudo-cautelar-modelo.docx")


# ======================================================================================
# 02 — TERMO DE CIÊNCIA E AUTORIZAÇÃO
# ======================================================================================
def termo():
    doc = novo_doc()
    timbre(doc, compacto=True)
    titulo(doc, "Autorização para vistoria e termo de ciência",
           "Parte A: autorização do proprietário/possuidor. Parte B: ciência do solicitante sobre escopo e limites. "
           "Parte C: ciência dos achados antes do fechamento do negócio.")
    form(doc, [
        [("Vistoria / laudo nº", "[0000/AAAA]"), ("Data", "[DD/MM/AAAA]")],
        [("Veículo", "[MARCA/MODELO/VERSÃO]"), ("Placa", "[ABC1D23]")],
        [("Chassi (VIN)", "[17 CARACTERES]"), ("Local", "[ENDEREÇO] – [CIDADE/UF]")],
    ])

    h1(doc, "Parte A – Autorização do proprietário ou possuidor")
    form(doc, [
        [("Nome", "[NOME COMPLETO]")],
        [("CPF / CNPJ", "[000.000.000-00]"), ("Telefone", "[(00) 00000-0000]")],
        [("Condição", "☐ proprietário  ☐ possuidor  ☐ representante da loja  ☐ procurador/autorizado")],
    ])
    spacer(doc)
    para(doc, "Autorizo [NOME DA EMPRESA / VISTORIADOR] a realizar vistoria cautelar no veículo acima, a pedido de "
              "[NOME DO SOLICITANTE], compreendendo:", align="j")
    autorizados = [
        "☐ acesso ao interior, ao compartimento do motor e ao porta-malas, com abertura de capô, portas e tampas;",
        "☐ remoção manual, sem ferramentas e sem desmontagem, de tapetes, tampas de acesso plásticas e do estepe, recolocados ao final;",
        "☐ medição de espessura de pintura com instrumento que não danifica a superfície;",
        "☐ partida do motor, funcionamento em marcha lenta e acionamento de comandos e acessórios;",
        "☐ conexão de scanner de diagnóstico à tomada OBD2 e teste de bateria/carga;",
        "☐ teste de rodagem curto (até [5] km), em trajeto combinado, por condutor habilitado nos termos do CTB – Lei 9.503/1997, "
        "acompanhado por [PROPRIETÁRIO / REPRESENTANTE]  ☐ não autorizo o teste de rodagem;",
        "☐ fotografias e vídeos do veículo, inclusive placa, chassi, número do motor, etiquetas e hodômetro, para uso exclusivo no laudo "
        "e nos registros do vistoriador, conforme o Aviso de Privacidade.",
    ]
    for a in autorizados:
        para(doc, a, align="j", indent=0.6, hanging=0.6, after=1, size=9.5)
    para(doc, "Declaro ser proprietário ou possuidor do veículo, ou estar autorizado por quem o seja, e que:", align="j", before=4, size=9.5)
    bullets(doc, size=9.5, items=[
        "a vistoria é não destrutiva e não inclui desmontagem; eventuais itens que exijam desmontagem ficarão como não verificados;",
        "durante o teste de rodagem, a responsabilidade por danos e infrações será tratada da seguinte forma: [DEFINIR – EX.: SEGURO DO "
        "VEÍCULO / SEGURO DE RESPONSABILIDADE CIVIL DO VISTORIADOR / CONDUTOR RESPONSÁVEL]; [VERIFICAR com seu advogado e sua seguradora];",
        "informo, por minha iniciativa, as seguintes condições conhecidas do veículo (opcional): [REPAROS, SINISTROS, TROCA DE PEÇAS, OUTROS].",
    ])
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    assinaturas(doc, [["[NOME DO PROPRIETÁRIO/POSSUIDOR]", "Proprietário / possuidor / autorizado"],
                      ["[NOME DO VISTORIADOR]", "Vistoriador"]])

    page_break(doc)
    h1(doc, "Parte B – Termo de ciência do solicitante")
    form(doc, [
        [("Solicitante", "[NOME COMPLETO]")],
        [("CPF / CNPJ", "[000.000.000-00]"), ("Telefone", "[(00) 00000-0000]")],
        [("Serviço contratado", "☐ [ESSENCIAL]  ☐ [COMPLETA]  ☐ [PREMIUM]  ☐ [OUTRO]")],
    ])
    spacer(doc)
    para(doc, "Declaro que fui informado(a), antes da vistoria, e estou ciente de que (rubrique cada item):", align="j")
    ciencia = [
        "A vistoria cautelar é um exame técnico **visual e não destrutivo**, sem desmontagem, realizado no estado em que o veículo foi apresentado.",
        "O resultado vale para a **data e o momento** da vistoria e perde atualidade com o tempo e o uso do veículo.",
        "Áreas **não acessíveis** sem desmontagem, elevador ou remoção de acabamentos podem não ser examinadas e serão registradas como não verificadas.",
        "O laudo **não garante** a inexistência de vícios ocultos, defeitos intermitentes ou falhas futuras.",
        "Este serviço **não é** a vistoria oficial exigida pelo órgão de trânsito para transferência, nem perícia oficial ou judicial.",
        "Em identificação veicular, o vistoriador registra divergências e indícios, mas **não conclui** sobre adulteração; nesse caso, a orientação é procurar a autoridade competente.",
        "As consultas dependem das **bases consultadas** e refletem a data e a hora da consulta; recomenda-se repeti-las antes do fechamento do negócio.",
        "A classificação segue a Política de Classificação do vistoriador (N0 a N4 e matriz de resultado), que me foi apresentada.",
        "A **decisão de compra ou venda é exclusivamente minha**; o laudo é um instrumento de informação, não uma aprovação do negócio.",
        "Nenhuma parte envolvida na negociação (vendedor, loja ou intermediário) pode interferir na classificação.",
    ]
    t = table(doc, ["Nº", "Declaração", "Rubrica"], [[str(i), c, ""] for i, c in enumerate(ciencia, 1)],
              [0.9, 13.8, 2.3], size=9, row_h=0.8, aligns=["c", "l", "c"])
    local_data(doc)
    assinaturas(doc, [["[NOME DO SOLICITANTE]", "Solicitante", "CPF [000.000.000-00]"],
                      ["[NOME DO VISTORIADOR]", "Vistoriador"]])

    page_break(doc)
    h1(doc, "Parte C – Ciência dos achados relevantes antes do fechamento do negócio")
    para(doc, "Preencher após a vistoria. Relacione os achados N2, N3 e N4 e todo item não verificado (NV). "
              "O solicitante rubrica cada linha depois da explicação.", size=9.5, color=TEXTO2, align="j")
    table(doc, ["Código", "Achado (como consta no laudo)", "Nível", "Rubrica"],
          [["", "", "", ""] for _ in range(10)], [2.0, 10.6, 1.8, 2.6], size=9, row_h=0.8, aligns=["c", "l", "c", "c"])
    spacer(doc)
    para(doc, "Resultado informado:  ☐ APROVADO   ☐ APROVADO COM APONTAMENTOS   ☐ REPROVADO   ☐ INCONCLUSIVO", bold=True, color=CARBONO)
    para(doc, "Declaro que cada achado acima me foi explicado, que pude fazer perguntas e que recebi o laudo completo. Estou ciente "
              "de que a decisão sobre o negócio é minha e registro, apenas para controle, a minha intenção neste momento:", align="j")
    for o in ["☐ não pretendo prosseguir com o negócio;",
              "☐ pretendo buscar avaliação complementar (oficina, concessionária, perícia) antes de decidir;",
              "☐ pretendo prosseguir, ciente dos achados acima;",
              "☐ prefiro não informar."]:
        para(doc, o, indent=0.6, after=1)
    local_data(doc)
    assinaturas(doc, [["[NOME DO SOLICITANTE]", "Solicitante", "CPF [000.000.000-00]"],
                      ["[NOME DO VISTORIADOR]", "Vistoriador"]])
    doc.save(OUT / "02-termo-de-ciencia-e-autorizacao.docx")


# ======================================================================================
# 03 — PROPOSTA COMERCIAL
# ======================================================================================
def proposta():
    doc = novo_doc()
    timbre(doc)
    titulo(doc, "Proposta comercial – Vistoria cautelar", "Para lojas, revendas, frotas, locadoras e compradores particulares.")
    form(doc, [
        [("Proposta nº", "[000/AAAA]"), ("Data", "[DD/MM/AAAA]")],
        [("Cliente", "[NOME / RAZÃO SOCIAL]")],
        [("A/C", "[NOME DO CONTATO]"), ("Validade", "[15] dias")],
    ])

    h1(doc, "1. Apresentação")
    para(doc, "[NOME DA EMPRESA] realiza vistoria cautelar de veículos usados: um exame técnico visual e não destrutivo, com "
              "registro fotográfico, medições e classificação por critério publicado, para apoiar decisões de compra, venda e "
              "entrada de estoque. Atendemos em [CIDADE/REGIÃO] desde [ANO].", align="j")
    para(doc, "Trabalhamos com uma **escala única de níveis (N0 a N4)** e uma **matriz de resultado** pública (Aprovado, Aprovado "
              "com apontamentos, Reprovado e Inconclusivo). A classificação é técnica e independente: ninguém envolvido na "
              "negociação interfere no resultado.", align="j")

    h1(doc, "2. Níveis de serviço")
    para(doc, "Os nomes e o conteúdo de cada nível são editáveis. ● incluso · ○ opcional · — não incluso.", size=9, color=TEXTO2)
    linhas = [
        ["Documentação, histórico e conferência com o veículo", "●", "●", "●"],
        ["Identificação veicular (conferência de elementos)", "●", "●", "●"],
        ["Estrutura: inspeção visual das áreas acessíveis", "●", "●", "●"],
        ["Carroceria e pintura com mapa de espessura peça a peça", "○", "●", "●"],
        ["Vidros, iluminação e sinalização", "●", "●", "●"],
        ["Rodas, pneus, suspensão e direção (visual e sensitivo)", "●", "●", "●"],
        ["Motor, transmissão e freios (visual, partida e funcionamento)", "●", "●", "●"],
        ["Leitura de diagnóstico OBD2 e teste de bateria/carga", "—", "●", "●"],
        ["Teste de rodagem curto (com autorização do proprietário)", "—", "○", "●"],
        ["Interior, segurança passiva e sinais de enchente/incêndio", "●", "●", "●"],
        ["Consultas: restrições, gravame, leilão/sinistro, débitos, recall", "○", "●", "●"],
        ["Relatório fotográfico (quantidade mínima de fotos)", "[20]", "[40]", "[60+]"],
        ["Laudo em PDF com explicação dos achados", "●", "●", "●"],
        ["Conversa de explicação do resultado (telefone ou presencial)", "—", "●", "●"],
        ["Prazo de entrega do laudo após a vistoria", "[24 h]", "[12 h]", "[no dia]"],
        ["**Valor por veículo**", "**R$ [000]**", "**R$ [000]**", "**R$ [000]**"],
    ]
    table(doc, ["Item", "[ESSENCIAL]", "[COMPLETA]", "[PREMIUM]"], linhas, [9.2, 2.6, 2.6, 2.6], size=9,
          aligns=["l", "c", "c", "c"])
    para(doc, "Motos, utilitários com chassi, veículos pesados, blindados e elétricos/híbridos: valor sob consulta.",
         size=9, color=TEXTO2, before=3)

    h1(doc, "3. Condições para lojas, revendas e frotas")
    table(doc, ["Volume mensal", "Condição", "Faturamento"],
          [["Até [10] vistorias", "Tabela padrão", "Por vistoria ou mensal"],
           ["[11] a [30] vistorias", "Desconto de [00]% sobre a tabela", "Mensal, fechamento no dia [00]"],
           ["Acima de [30] vistorias", "Condição a combinar em convênio", "Mensal, fechamento no dia [00]"]],
          [4.5, 7.0, 5.5], size=9)
    bullets(doc, [
        "O desconto por volume é concedido ao contratante do serviço e não altera o critério nem o resultado da vistoria.",
        "Quando a vistoria for paga pela loja e entregue ao consumidor, o laudo é entregue completo, sem supressão de páginas.",
        "Convênios de volume são formalizados em contrato de parceria específico.",
    ])

    h1(doc, "4. Prazos e agendamento")
    bullets(doc, [
        "Agendamento com antecedência mínima de [00] horas, pelo WhatsApp [(00) 00000-0000].",
        "Duração média da vistoria: [60 a 120] minutos, conforme o nível e o veículo.",
        "Entrega do laudo em PDF conforme o prazo do nível contratado, contado do término da vistoria.",
        "Deslocamento incluso num raio de [00] km de [BAIRRO/CIDADE]; acima disso, R$ [0,00] por km [OU TAXA FIXA].",
    ])

    h1(doc, "5. O que não está incluído")
    bullets(doc, [
        "Desmontagem de peças, abertura de conjuntos mecânicos, uso de elevador (salvo quando disponível no local) e ensaios de laboratório.",
        "Vistoria oficial de identificação veicular exigida pelo órgão de trânsito para transferência, perícia oficial ou judicial.",
        "Garantia do veículo, avaliação de preço de mercado e intermediação do negócio.",
        "Orçamento de reparos (pode ser indicado como recomendação de avaliação complementar).",
    ])

    h1(doc, "6. Condições de pagamento")
    bullets(doc, [
        "Particulares: pagamento [NO AGENDAMENTO / NA ENTREGA DO LAUDO] via PIX, cartão ou transferência.",
        "Empresas conveniadas: faturamento mensal, vencimento em [00] dias após o fechamento, com nota fiscal [VERIFICAR regime tributário e emissão de NFS-e no seu município].",
        "Em caso de atraso: multa de [2]% e juros de [1]% ao mês, proporcionais ao período.",
    ])

    h1(doc, "7. Reagendamento e cancelamento")
    bullets(doc, [
        "Reagendamento sem custo com aviso de até [00] horas antes do horário marcado.",
        "Se o veículo não estiver disponível no local e horário combinados, ou se a vistoria não puder ser iniciada por motivo alheio ao "
        "vistoriador, poderá ser cobrada taxa de deslocamento de R$ [0,00].",
        "Veículo que não permita a verificação da identificação veicular gera laudo INCONCLUSIVO, com o valor [INTEGRAL / REDUZIDO – DEFINIR].",
    ])

    h1(doc, "8. Aceite")
    para(doc, "Esta proposta é válida por [15] dias a partir da data de emissão. O aceite pode ser feito por assinatura abaixo ou por "
              "mensagem escrita que faça referência ao número desta proposta.", align="j")
    local_data(doc)
    assinaturas(doc, [["[NOME DO VISTORIADOR / EMPRESA]", "Proponente"], ["[NOME DO CLIENTE]", "De acordo – cliente"]])
    doc.save(OUT / "03-proposta-comercial.docx")


# ======================================================================================
# 04 — CONTRATO DE PRESTAÇÃO DE SERVIÇOS
# ======================================================================================
def contrato():
    doc = novo_doc()
    titulo(doc, "Contrato de prestação de serviços de vistoria cautelar")
    partes_contrato(doc, "CONTRATADA", "CONTRATANTE",
                    "[NOME COMPLETO / RAZÃO SOCIAL], inscrito(a) no CPF/CNPJ sob o nº [000.000.000-00], residente/sediado(a) em "
                    "[ENDEREÇO COMPLETO], [CIDADE/UF], telefone [(00) 00000-0000], e-mail [E-MAIL]")
    clausula(doc, 1, "Do objeto", [
        "O presente contrato tem por objeto a prestação, pela CONTRATADA, de serviço de **vistoria cautelar** no veículo "
        "[MARCA/MODELO], placa [ABC1D23], chassi [17 CARACTERES] (\"Veículo\"), no nível [ESSENCIAL/COMPLETA/PREMIUM], com emissão "
        "de laudo técnico, conforme descrito na proposta nº [000/AAAA], que integra este contrato.",
        "Vistoria cautelar é o exame técnico visual, sensitivo e não destrutivo do Veículo, destinado a informar o CONTRATANTE sobre as "
        "condições observadas no momento da vistoria, para apoiar sua decisão de compra, venda ou recebimento do bem.",
    ])
    clausula(doc, 2, "Do escopo e do método", [
        ("A vistoria compreende, conforme o nível contratado:", [
            "conferência de documentação e dos elementos de identificação veicular com o documento apresentado;",
            "inspeção das áreas acessíveis de estrutura, carroceria e pintura, com medição de espessura quando prevista;",
            "verificação de vidros, iluminação, rodas, pneus, suspensão, direção, motor, transmissão, freios, elétrica e interior;",
            "leitura de diagnóstico eletrônico e teste de rodagem curto, quando previstos e autorizados pelo proprietário/possuidor;",
            "consultas em bases de informação, quando previstas, com registro de fonte, data e hora;",
            "registro fotográfico e emissão de laudo com classificação segundo a Política de Classificação da CONTRATADA.",
        ]),
        "A classificação final observa a escala de níveis N0 a N4 e a matriz de resultado (Aprovado, Aprovado com apontamentos, "
        "Reprovado e Inconclusivo), cuja cópia foi entregue ao CONTRATANTE antes da contratação.",
    ])
    clausula(doc, 3, "Do que não integra o escopo", [
        ("Não integram o objeto deste contrato:", [
            "desmontagem de peças ou conjuntos, remoção de componentes fixados por ferramentas e ensaios de laboratório;",
            "a vistoria de identificação veicular exigida pelo órgão executivo de trânsito, nos termos do CTB – Lei 9.503/1997 e das "
            "Resoluções do CONTRAN [VERIFICAR resolução vigente], bem como perícia oficial ou judicial;",
            "garantia do Veículo, avaliação de preço de mercado, intermediação do negócio ou orçamento de reparos;",
            "conclusão sobre adulteração de elementos de identificação, que depende de perícia oficial; a CONTRATADA registrará "
            "divergências ou indícios observados.",
        ]),
        "Itens que não puderem ser verificados em razão do estado de apresentação do Veículo, de condições do local ou da falta de "
        "autorização serão registrados como não verificados, com o motivo.",
    ])
    clausula(doc, 4, "Das obrigações da CONTRATADA", [
        ("São obrigações da CONTRATADA:", [
            "executar a vistoria pessoalmente ou por profissional qualificado sob sua responsabilidade, com técnica, diligência e independência;",
            "informar previamente ao CONTRATANTE, de forma clara, o escopo, os limites, o preço e o prazo do serviço, nos termos do "
            "art. 6º, III, do CDC – Lei 8.078/1990;",
            "registrar os achados de forma objetiva, com código, nível e referência fotográfica;",
            "entregar o laudo completo no prazo ajustado e esclarecer as dúvidas do CONTRATANTE sobre o seu conteúdo;",
            "manter o sigilo das informações e tratar dados pessoais conforme a Cláusula 11ª;",
            "não aceitar de terceiros (vendedor, loja, intermediário) qualquer vantagem condicionada ao resultado da vistoria.",
        ]),
    ])
    clausula(doc, 5, "Das obrigações do CONTRATANTE", [
        ("São obrigações do CONTRATANTE:", [
            "fornecer informações verdadeiras e os dados necessários ao agendamento;",
            "providenciar a disponibilidade do Veículo no local e horário combinados, com o documento e a autorização do "
            "proprietário/possuidor, quando não for ele próprio;",
            "efetuar o pagamento na forma da Cláusula 9ª;",
            "ler o laudo integralmente, em especial a seção de metodologia e limites, e solicitar esclarecimentos antes de concluir o negócio.",
        ]),
    ])
    clausula(doc, 6, "Da natureza da obrigação", [
        "A obrigação assumida pela CONTRATADA é **de meio**: consiste em empregar a técnica e a diligência adequadas para examinar o "
        "Veículo dentro do escopo contratado e relatar fielmente o que for observado. A CONTRATADA não garante resultado do negócio, "
        "desempenho futuro do Veículo nem a inexistência de vícios não detectáveis pelo método descrito na Cláusula 2ª.",
        "O resultado retrata as condições observadas na data, no horário e no local da vistoria, no estado em que o Veículo foi apresentado.",
    ])
    clausula(doc, 7, "Da responsabilidade", [
        "A CONTRATADA responde pelos danos causados ao CONTRATANTE por falha na execução do serviço contratado, dentro do escopo "
        "definido nas Cláusulas 2ª e 3ª, nos termos da legislação aplicável, inclusive do CDC – Lei 8.078/1990 quando houver relação de consumo.",
        "A CONTRATADA não responde por condições do Veículo situadas fora do escopo contratado, não acessíveis pelo método descrito "
        "ou surgidas após a vistoria, nem por informações constantes de bases de terceiros consultadas, cujo conteúdo não produz.",
        "Nenhuma disposição deste contrato exclui ou atenua direitos assegurados ao consumidor por lei, em observância ao art. 51, I, "
        "do CDC. A delimitação acima define o objeto do serviço, e não exonera a CONTRATADA de sua responsabilidade pela correta execução.",
        "A CONTRATADA [MANTÉM / NÃO MANTÉM] seguro de responsabilidade civil profissional [VERIFICAR com sua seguradora a cobertura aplicável].",
        "Constatada falha na execução, a CONTRATADA se compromete a reexaminar o caso e, se confirmada, emitir laudo corrigido, sem "
        "prejuízo dos demais direitos do CONTRATANTE.",
    ])
    clausula(doc, 8, "Dos prazos", [
        "A vistoria será realizada em [DD/MM/AAAA], às [00h00], em [LOCAL], ou em outra data acordada por escrito.",
        "O laudo será entregue em até [00] horas após o término da vistoria, em PDF, pelo e-mail ou WhatsApp informados.",
        "Reagendamentos solicitados com até [00] horas de antecedência não geram custo.",
    ])
    clausula(doc, 9, "Do preço e do pagamento", [
        "Pelo serviço, o CONTRATANTE pagará R$ [0,00] ([VALOR POR EXTENSO]), acrescido de taxa de deslocamento de R$ [0,00], quando aplicável.",
        "O pagamento será feito [NO AGENDAMENTO / NA ENTREGA DO LAUDO / POR FATURAMENTO], via [PIX / CARTÃO / TRANSFERÊNCIA], "
        "mediante recibo ou nota fiscal [VERIFICAR obrigação de emissão de NFS-e no seu município].",
        "O atraso sujeita o valor devido a multa de [2]% e juros de [1]% ao mês, pro rata die, observado o limite do art. 52, § 1º, do CDC "
        "quando houver relação de consumo.",
        "O resultado da vistoria não altera o preço: o valor é devido pelo serviço prestado, qualquer que seja a classificação.",
    ])
    clausula(doc, 10, "Do cancelamento e do direito de arrependimento", [
        "O CONTRATANTE pode cancelar a vistoria sem custo até [00] horas antes do horário agendado. Após esse prazo, ou se o Veículo não "
        "for disponibilizado, poderá ser cobrada apenas a taxa de deslocamento de R$ [0,00].",
        "Quando a contratação ocorrer fora do estabelecimento comercial (por telefone, internet ou aplicativo de mensagens), fica "
        "assegurado ao consumidor o direito de arrependimento previsto no art. 49 do CDC [VERIFICAR a aplicação quando o serviço for "
        "executado, a pedido do consumidor, dentro do prazo de 7 dias].",
    ])
    clausula(doc, 11, "Da proteção de dados pessoais", [
        "As partes tratarão dados pessoais em conformidade com a LGPD – Lei 13.709/2018. A CONTRATADA trata os dados do CONTRATANTE, do "
        "proprietário/possuidor e do Veículo (inclusive placa, chassi e fotografias) para executar este contrato, cumprir obrigações "
        "legais e exercer direitos em eventual processo, nas hipóteses do art. 7º da LGPD.",
        "Os dados não serão utilizados para finalidade diversa nem compartilhados, salvo com prestadores necessários à execução "
        "(como provedores de consulta e armazenamento), por determinação legal ou com autorização do titular.",
        "O titular pode exercer os direitos do art. 18 da LGPD pelo contato [E-MAIL DO ENCARREGADO/CANAL DE PRIVACIDADE], "
        "conforme o Aviso de Privacidade entregue ao CONTRATANTE.",
        "Fotografias de uso em divulgação dependem de autorização específica, com placa e dados identificadores ocultados.",
    ])
    clausula(doc, 12, "Do laudo e da confidencialidade", [
        "O laudo destina-se ao CONTRATANTE. Seu compartilhamento com terceiros é livre ao CONTRATANTE, desde que o documento seja "
        "apresentado completo e sem alterações.",
        "A CONTRATADA manterá cópia do laudo e das fotografias pelo prazo indicado no Aviso de Privacidade, para fins de comprovação "
        "do serviço e defesa de direitos.",
    ])
    clausula(doc, 13, "Das disposições gerais", [
        "Este contrato pode ser assinado de forma física ou eletrônica, e as comunicações entre as partes podem ser feitas pelos e-mails "
        "e telefones indicados na qualificação.",
        "Integram este contrato: a proposta comercial, a Política de Classificação e o Aviso de Privacidade da CONTRATADA.",
        "Aplicam-se subsidiariamente as regras do Código Civil – Lei 10.406/2002 sobre prestação de serviço (arts. 593 e seguintes) e, "
        "havendo relação de consumo, o CDC – Lei 8.078/1990.",
    ])
    clausula(doc, 14, "Do foro", [
        "Fica eleito o foro da comarca de [CIDADE/UF] para dirimir questões oriundas deste contrato. Havendo relação de consumo, "
        "o consumidor poderá propor a ação no foro de seu domicílio, nos termos do art. 101, I, do CDC.",
    ])
    para(doc, "E, por estarem de acordo, as partes assinam este instrumento em [2] vias de igual teor.", align="j", before=6)
    local_data(doc)
    assinaturas(doc, [["[NOME DA EMPRESA]", "CONTRATADA"], ["[NOME DO CONTRATANTE]", "CONTRATANTE"],
                      ["Testemunha 1: [NOME]", "CPF [000.000.000-00]"], ["Testemunha 2: [NOME]", "CPF [000.000.000-00]"]])
    doc.save(OUT / "04-contrato-prestacao-servicos.docx")


# ======================================================================================
# 05 — CONTRATO DE PARCERIA COM LOJISTA
# ======================================================================================
def parceria():
    doc = novo_doc()
    titulo(doc, "Contrato de parceria para vistorias cautelares", "Convênio entre vistoriador e loja/revenda de veículos")
    partes_contrato(doc, "VISTORIADORA", "PARCEIRA",
                    "[RAZÃO SOCIAL DA LOJA], inscrita no CNPJ sob o nº [00.000.000/0000-00], com sede em [ENDEREÇO COMPLETO], "
                    "[CIDADE/UF], neste ato representada por [NOME DO REPRESENTANTE], [CARGO]")
    clausula(doc, 1, "Do objeto", [
        "Este contrato estabelece as condições para a prestação, pela VISTORIADORA, de serviços de vistoria cautelar em veículos "
        "indicados pela PARCEIRA, com faturamento mensal e condições comerciais de volume.",
        ("As vistorias poderão ocorrer nas seguintes modalidades:", [
            "**entrada de estoque**: contratada e paga pela PARCEIRA, para avaliação de veículos que pretende adquirir;",
            "**vistoria de venda**: contratada pela PARCEIRA para disponibilizar o laudo ao comprador;",
            "**vistoria a pedido do comprador**: contratada diretamente pelo consumidor, com agendamento facilitado pela PARCEIRA.",
        ]),
    ])
    clausula(doc, 2, "Da independência técnica", [
        "A VISTORIADORA atua com **total independência técnica**. A classificação de cada item e o resultado final seguem "
        "exclusivamente a Política de Classificação da VISTORIADORA (Anexo II), aplicada igualmente a todos os clientes.",
        "A PARCEIRA não poderá solicitar, sugerir ou condicionar qualquer vantagem à alteração, omissão ou suavização de achados, "
        "níveis ou resultados, nem à supressão de páginas ou fotografias do laudo.",
        "O volume de vistorias, o desconto ou a continuidade desta parceria **não dependem do resultado** das vistorias.",
        "A tentativa de interferência na classificação autoriza a VISTORIADORA a rescindir este contrato de imediato, sem ônus.",
    ])
    clausula(doc, 3, "Do volume e dos preços", [
        "A PARCEIRA estima a solicitação de [00] vistorias por mês, sem obrigação mínima [OU: com mínimo de [00] vistorias/mês].",
        "Os preços por modalidade e nível constam do Anexo I e poderão ser reajustados anualmente por [ÍNDICE], ou mediante acordo escrito.",
    ])
    clausula(doc, 4, "Do agendamento e dos prazos", [
        "As solicitações serão feitas por [WHATSAPP / E-MAIL / SISTEMA] com antecedência mínima de [00] horas, informando placa, local "
        "e responsável pelo acompanhamento.",
        "A PARCEIRA disponibilizará o veículo com documento, chave e local adequado (coberto ou com boa iluminação), e autorizará a "
        "partida do motor e, quando previsto, o teste de rodagem.",
        "O laudo será entregue em até [00] horas após a vistoria, ao contratante do serviço em cada modalidade.",
    ])
    clausula(doc, 5, "Do faturamento mensal", [
        "As vistorias contratadas pela PARCEIRA serão consolidadas em relatório mensal com fechamento no dia [00] de cada mês, contendo "
        "data, placa, modalidade, nível e valor.",
        "A VISTORIADORA emitirá nota fiscal [VERIFICAR regime e emissão de NFS-e] com vencimento em [00] dias do fechamento.",
        "O atraso sujeita o valor a multa de [2]% e juros de [1]% ao mês. Após [30] dias de atraso, a VISTORIADORA poderá suspender "
        "novos agendamentos até a regularização.",
    ])
    clausula(doc, 6, "Do laudo e da comunicação ao consumidor", [
        "Quando o laudo for disponibilizado a comprador, a PARCEIRA o entregará **completo, sem alterações e sem supressão de páginas**, "
        "incluindo a seção de metodologia e limites.",
        "A PARCEIRA não poderá apresentar o laudo como garantia do veículo nem usar expressões como \"veículo certificado\" ou "
        "\"aprovado pelo vistoriador\" em anúncios, salvo reprodução fiel do resultado e da data da vistoria.",
        "O consumidor poderá contatar diretamente a VISTORIADORA para esclarecimentos sobre o laudo.",
        "A PARCEIRA reconhece que o laudo não substitui a vistoria oficial exigida pelo órgão de trânsito para transferência.",
    ])
    clausula(doc, 7, "Da comissão, do desconto e da transparência", [
        ("Condição comercial adotada [MARCAR UMA OPÇÃO E EXCLUIR AS DEMAIS]:", [
            "☐ não há comissão nem remuneração entre as partes além do preço dos serviços;",
            "☐ a PARCEIRA recebe desconto de [00]% sobre a tabela nas vistorias que contratar;",
            "☐ a VISTORIADORA concede ao consumidor indicado pela PARCEIRA desconto de [00]% sobre a tabela.",
        ]),
        "É vedada qualquer comissão, bonificação ou vantagem, de qualquer das partes, vinculada ao resultado da vistoria ou à conclusão da venda.",
        "Havendo qualquer benefício comercial entre as partes relacionado a vistoria apresentada ao consumidor, este será **informado de "
        "forma clara ao consumidor** no laudo ou no termo de ciência, em atenção ao dever de informação do art. 6º, III, do CDC – Lei 8.078/1990.",
    ])
    clausula(doc, 8, "Da confidencialidade", [
        "As partes manterão sigilo sobre informações comerciais, preços, dados de clientes e de veículos a que tiverem acesso, "
        "durante a vigência e por [2] anos após o término deste contrato.",
        "Não se considera quebra de sigilo a entrega do laudo ao contratante do serviço ou ao consumidor, na forma da Cláusula 6ª, "
        "nem a divulgação exigida por lei ou por autoridade.",
    ])
    clausula(doc, 9, "Da proteção de dados pessoais", [
        "Cada parte é controladora dos dados pessoais que trata para suas próprias finalidades, nos termos da LGPD – Lei 13.709/2018, e "
        "responde pelo tratamento que realizar.",
        "A PARCEIRA somente compartilhará com a VISTORIADORA os dados estritamente necessários ao agendamento e à execução da vistoria, "
        "com base legal adequada e informação ao titular.",
        "As partes adotarão medidas de segurança razoáveis, comunicarão uma à outra, em até [48] horas, incidentes que afetem dados "
        "compartilhados, e cooperarão no atendimento a pedidos de titulares.",
        "Fotografias de veículos só serão usadas em divulgação por qualquer das partes com placa e dados identificadores ocultados e, "
        "quando houver dados pessoais, com autorização do titular.",
    ])
    clausula(doc, 10, "Do uso de nome e marca", [
        "Cada parte poderá mencionar a existência da parceria em seus canais, sem sugerir que a VISTORIADORA garante ou certifica os "
        "veículos da PARCEIRA. O uso de logotipos depende de autorização prévia por escrito.",
    ])
    clausula(doc, 11, "Da não exclusividade e da natureza da relação", [
        "Esta parceria não é exclusiva: a VISTORIADORA pode atender outras lojas e consumidores, e a PARCEIRA pode contratar outros profissionais.",
        "Não há vínculo empregatício, societário ou de representação entre as partes, que respondem cada uma por seus encargos.",
    ])
    clausula(doc, 12, "Da vigência e da rescisão", [
        "Este contrato vigora por [12] meses a partir da assinatura e se renova automaticamente por iguais períodos, salvo aviso em contrário.",
        "Qualquer parte pode rescindi-lo sem multa mediante aviso escrito com [30] dias de antecedência, mantidas as obrigações de "
        "pagamento das vistorias realizadas e as de sigilo e proteção de dados.",
        "O descumprimento das Cláusulas 2ª, 6ª ou 7ª autoriza a rescisão imediata pela parte prejudicada.",
    ])
    clausula(doc, 13, "Do foro", [
        "Fica eleito o foro da comarca de [CIDADE/UF] para dirimir questões oriundas deste contrato. Aplica-se o Código Civil – Lei 10.406/2002.",
    ])
    para(doc, "E, por estarem de acordo, as partes assinam este instrumento em [2] vias de igual teor.", align="j", before=6)
    local_data(doc)
    assinaturas(doc, [["[NOME DA EMPRESA]", "VISTORIADORA"], ["[RAZÃO SOCIAL DA LOJA]", "PARCEIRA"],
                      ["Testemunha 1: [NOME]", "CPF [000.000.000-00]"], ["Testemunha 2: [NOME]", "CPF [000.000.000-00]"]])

    page_break(doc)
    h1(doc, "Anexo I – Tabela de preços da parceria")
    table(doc, ["Modalidade", "[ESSENCIAL]", "[COMPLETA]", "[PREMIUM]"],
          [["Entrada de estoque", "R$ [000]", "R$ [000]", "R$ [000]"],
           ["Vistoria de venda (laudo ao comprador)", "R$ [000]", "R$ [000]", "R$ [000]"],
           ["A pedido do comprador (pago pelo consumidor)", "R$ [000]", "R$ [000]", "R$ [000]"],
           ["Motos / utilitários com chassi", "R$ [000]", "R$ [000]", "R$ [000]"],
           ["Deslocamento fora do raio de [00] km", "R$ [0,00]/km", "R$ [0,00]/km", "R$ [0,00]/km"]],
          [8.0, 3.0, 3.0, 3.0], size=9, aligns=["l", "c", "c", "c"])
    h1(doc, "Anexo II – Política de Classificação")
    para(doc, "Anexar a Política de Classificação vigente da VISTORIADORA (escala N0 a N4 e matriz de resultado), rubricada pelas partes.",
         align="j")
    doc.save(OUT / "05-contrato-parceria-lojista.docx")


# ======================================================================================
# 06 — RECIBO
# ======================================================================================
def recibo():
    doc = novo_doc()

    def via(nome_via):
        timbre(doc, compacto=True)
        t = doc.add_table(rows=1, cols=2)
        set_widths(t, [11.0, 6.0])
        tbl_borders(t, val="nil")
        p = para(t.cell(0, 0), "RECIBO", size=17, bold=True, color=CARBONO, after=0)
        para(t.cell(0, 0), f"Nº [0000/AAAA]  ·  {nome_via}", size=9, color=TEXTO2, after=0)
        vc = t.cell(0, 1)
        shade(vc, PAPEL2)
        vc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para(vc, "VALOR", size=8, bold=True, color=TEXTO2, align="c", after=0)
        para(vc, "R$ [0.000,00]", size=15, bold=True, color=CARBONO, align="c", after=0)
        pr = para(doc, "", after=4)
        p_border(pr, "bottom", sz=8, color=AMBAR, space=2)
        para(doc, "Recebi(emos) de **[NOME COMPLETO / RAZÃO SOCIAL]**, CPF/CNPJ [000.000.000-00], a importância de "
                  "**R$ [0.000,00]** ([VALOR POR EXTENSO]), referente à prestação de serviço de **vistoria cautelar** no veículo "
                  "[MARCA/MODELO], placa [ABC1D23], realizada em [DD/MM/AAAA], laudo nº [0000/AAAA], nível [NÍVEL DO SERVIÇO].",
             align="j", before=4)
        para(doc, "Forma de pagamento:  ☐ PIX   ☐ dinheiro   ☐ cartão de débito   ☐ cartão de crédito   ☐ transferência   ☐ faturado",
             size=9.5)
        para(doc, "Pelo que dou(damos) plena quitação do valor acima.", size=9.5)
        local_data(doc)
        assinaturas(doc, [["[NOME DO VISTORIADOR / EMPRESA]", "CPF/CNPJ [00.000.000/0000-00]"]], cols=1)
        para(doc, "Este recibo não substitui a nota fiscal quando sua emissão for obrigatória [VERIFICAR regras de NFS-e do seu município].",
             size=7.5, italic=True, color=TEXTO3, before=4, after=0)

    via("1ª via – cliente")
    corte = para(doc, "✂  recorte aqui", size=7.5, color=TEXTO3, align="c", before=8, after=8)
    p_border(corte, "top", sz=4, color="9AA3AD", space=4, val="dashed")
    via("2ª via – emitente")
    doc.save(OUT / "06-recibo.docx")


# ======================================================================================
# 07 — AVISO DE PRIVACIDADE
# ======================================================================================
def privacidade():
    doc = novo_doc()
    timbre(doc, compacto=True)
    titulo(doc, "Aviso de privacidade", "Como tratamos os dados pessoais de clientes e de proprietários nas vistorias cautelares – "
                                        "LGPD, Lei 13.709/2018. Versão [1.0] – vigente desde [DD/MM/AAAA].")
    h1(doc, "1. Quem somos")
    para(doc, "[NOME DA EMPRESA], CNPJ/CPF [00.000.000/0000-00], com endereço em [ENDEREÇO COMPLETO], [CIDADE/UF], é a "
              "**controladora** dos dados pessoais tratados na prestação de serviços de vistoria cautelar.", align="j")

    h1(doc, "2. Quais dados coletamos")
    table(doc, ["Categoria", "Exemplos", "Origem"],
          [["Identificação e contato", "Nome, CPF ou CNPJ, telefone/WhatsApp, e-mail, endereço", "Você, ao agendar ou contratar"],
           ["Dados do veículo", "Placa, chassi (VIN), RENAVAM, nº do motor, marca/modelo, ano, cor, quilometragem",
            "Documento do veículo e vistoria"],
           ["Fotografias e vídeos", "Imagens do veículo, inclusive placa, chassi, etiquetas e hodômetro", "Vistoria"],
           ["Proprietário/possuidor", "Nome, CPF e assinatura na autorização de vistoria", "Proprietário ou possuidor"],
           ["Consultas", "Restrições, gravame, leilão/sinistro, débitos e recall vinculados ao veículo",
            "Provedores e bases de consulta"],
           ["Pagamento", "Forma de pagamento, comprovantes, dados para nota fiscal", "Você e meios de pagamento"]],
          [3.8, 8.2, 5.0], size=9)
    para(doc, "Placa, chassi e demais dados do veículo são tratados como dados pessoais sempre que puderem ser associados a uma pessoa. "
              "Não coletamos dados pessoais sensíveis para a vistoria. Se aparecerem pessoas ou placas de terceiros nas fotos, "
              "elas são recortadas ou ocultadas.", size=9.5, align="j", before=4)

    h1(doc, "3. Para que usamos e com qual base legal")
    table(doc, ["Finalidade", "Base legal (LGPD, art. 7º)"],
          [["Agendar, realizar a vistoria, emitir e entregar o laudo", "Execução de contrato e procedimentos preliminares (inciso V)"],
           ["Emitir recibo ou nota fiscal e manter registros contábeis", "Cumprimento de obrigação legal ou regulatória (inciso II)"],
           ["Guardar laudo, fotos e termos para comprovar o serviço e defender direitos",
            "Exercício regular de direitos em processo judicial, administrativo ou arbitral (inciso VI)"],
           ["Segurança, prevenção a fraudes e melhoria do serviço", "Legítimo interesse (inciso IX), respeitados seus direitos"],
           ["Enviar novidades e ofertas; usar fotos (sem placa) em divulgação", "Consentimento (inciso I), revogável a qualquer momento"]],
          [9.0, 8.0], size=9)

    h1(doc, "4. Com quem compartilhamos")
    bullets(doc, [
        "Com o **solicitante da vistoria**, que recebe o laudo (e com quem ele escolher compartilhar).",
        "Com **prestadores necessários** ao serviço: provedores de consulta veicular, armazenamento em nuvem, aplicativo de vistoria, "
        "meios de pagamento e contabilidade, sob dever de confidencialidade.",
        "Com a **loja parceira**, somente quando ela for a contratante da vistoria ou quando você autorizar.",
        "Com **autoridades**, quando houver obrigação legal ou ordem judicial.",
        "Alguns prestadores podem armazenar dados fora do Brasil; nesse caso, adotamos as salvaguardas previstas nos arts. 33 e seguintes "
        "da LGPD [VERIFICAR mecanismo aplicável ao seu fornecedor].",
        "**Não vendemos** dados pessoais.",
    ])

    h1(doc, "5. Por quanto tempo guardamos")
    table(doc, ["Dado", "Prazo de retenção", "Motivo"],
          [["Laudo, fotos, termos e autorizações", "[5] anos após a vistoria",
            "Prazos de reclamação e de ação (ex.: art. 27 do CDC – Lei 8.078/1990; art. 206, § 3º, V, do Código Civil)"],
           ["Documentos fiscais e contábeis", "[PRAZO] [VERIFICAR com seu contador]", "Obrigação legal"],
           ["Contato para ofertas", "Até a revogação do consentimento", "Consentimento"],
           ["Mensagens de agendamento", "[12] meses", "Atendimento e histórico"]],
          [5.0, 4.2, 7.8], size=9)
    para(doc, "Ao fim do prazo, os dados são eliminados ou anonimizados, salvo as hipóteses de conservação do art. 16 da LGPD.",
         size=9.5, before=4)

    h1(doc, "6. Como protegemos")
    para(doc, "Adotamos medidas técnicas e administrativas razoáveis: acesso restrito aos arquivos, senhas e autenticação em dois "
              "fatores nas contas usadas, cópias de segurança e envio do laudo somente aos contatos informados na contratação.", align="j")

    h1(doc, "7. Seus direitos como titular (LGPD, art. 18)")
    bullets(doc, [
        "confirmação da existência de tratamento e acesso aos dados;",
        "correção de dados incompletos, inexatos ou desatualizados;",
        "anonimização, bloqueio ou eliminação de dados desnecessários, excessivos ou tratados em desconformidade com a lei;",
        "portabilidade, nos termos da regulamentação;",
        "eliminação dos dados tratados com consentimento, ressalvadas as hipóteses legais de conservação;",
        "informação sobre com quem compartilhamos os dados e sobre a possibilidade de não consentir e suas consequências;",
        "revogação do consentimento;",
        "petição à Autoridade Nacional de Proteção de Dados (ANPD).",
    ])
    para(doc, "Responderemos em até [15] dias. Podemos solicitar informações para confirmar sua identidade.", size=9.5)

    h1(doc, "8. Encarregado e contato")
    form(doc, [[("Encarregado (DPO)", "[NOME DO ENCARREGADO OU CANAL RESPONSÁVEL]")],
               [("E-mail", "[E-MAIL DE PRIVACIDADE]"), ("Telefone", "[(00) 00000-0000]")]])
    para(doc, "Agentes de tratamento de pequeno porte podem ser dispensados de indicar encarregado, devendo manter canal de comunicação "
              "com o titular (Resolução CD/ANPD nº 2/2022) [VERIFICAR enquadramento e norma vigente].", size=9, color=TEXTO2, before=4)

    h1(doc, "9. Atualizações")
    para(doc, "Este aviso pode ser atualizado. A versão vigente fica disponível em [SITE OU LINK] e pode ser solicitada pelo WhatsApp.",
         align="j")
    doc.save(OUT / "07-aviso-de-privacidade-lgpd.docx")


# ======================================================================================
# 08 — POLÍTICA DE CLASSIFICAÇÃO (pública, sem aviso jurídico no rodapé)
# ======================================================================================
def politica():
    doc = novo_doc(aviso=False, rodape_esq="[NOME DA EMPRESA]  ·  Política de Classificação  ·  versão [1.0]")
    timbre(doc, compacto=True)
    titulo(doc, "Política de classificação dos laudos",
           "Como classificamos cada achado e o resultado de cada vistoria cautelar. Publicada para que você saiba, antes de contratar, "
           "exatamente qual critério será usado.")
    h1(doc, "1. Por que publicamos este critério")
    para(doc, "Uma vistoria só é útil se o critério for o mesmo para todos os veículos e todos os clientes. Por isso usamos uma "
              "escala fixa de níveis e uma regra fixa de resultado, iguais para compradores, vendedores e lojas. Ninguém paga para "
              "mudar uma classificação.", align="j")
    h1(doc, "2. Os níveis de cada achado")
    exemplos = {
        "N0": "Item sem anormalidade observável.",
        "N1": "Riscos superficiais; pneu com metade da vida útil; pastilha de freio a meia vida.",
        "N2": "Porta repintada; para-lama substituído; vazamento leve.",
        "N3": "Painel frontal substituído; ponteira de longarina reparada; pneu abaixo do indicador de desgaste; luz de airbag acesa.",
        "N4": "Reparo em estrutura de sustentação (longarina, torre, coluna, soleira, assoalho, painel corta-fogo); sinais de enchente ou "
              "incêndio; divergência na identificação veicular.",
    }
    t = table(doc, ["Nível", "Nome", "O que significa", "Exemplos"],
              [[n["id"], n["nome"], n["definicao"], exemplos[n["id"]]] for n in DADOS["niveis"]],
              [1.3, 2.9, 5.6, 7.2], size=9, aligns=["c", "l", "l", "l"], first_bold=True)
    for i, n in enumerate(DADOS["niveis"]):
        c = t.rows[i + 1].cells[0]
        shade(c, n["cor"].lstrip("#"))
        for r in c.paragraphs[0].runs:
            r.font.color.rgb = rgb("FFFFFF")
    para(doc, "NA = não se aplica ao veículo. NV = não verificado, sempre com o motivo.", size=9, color=TEXTO2, before=3)

    h1(doc, "3. A regra do resultado final")
    para(doc, "Depois que todos os itens recebem um nível, o resultado é definido por esta matriz, **sempre nesta ordem**:", align="j")
    t = doc.add_table(rows=len(DADOS["classificacao"]), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_widths(t, [0.35, 1.0, 5.6, 10.05])
    tbl_borders(t)
    tbl_cell_margins(t, top=0.14, bottom=0.14)
    for i, c in enumerate(DADOS["classificacao"]):
        row = t.rows[i]
        row_flags(row)
        shade(row.cells[0], c["cor"].lstrip("#"))
        for cc in row.cells:
            cc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        para(row.cells[1], f"{i + 1}º", size=11, bold=True, color=TEXTO2, align="c", after=0)
        para(row.cells[2], c["id"], size=11, bold=True, color=CARBONO, after=0)
        para(row.cells[3], c["regra"], size=9.5, after=0)
    spacer(doc)
    bullets(doc, [
        "**INCONCLUSIVO** vem primeiro: se não foi possível verificar a identificação do veículo, o laudo não conclui e explica o motivo "
        "e o que seria necessário para concluir.",
        "**N3** aparece em destaque \"Atenção\" no laudo, mesmo quando o resultado é Aprovado com apontamentos.",
        "Divergência na identificação veicular é registrada como fato observado. Não concluímos sobre adulteração: orientamos a procurar "
        "a perícia oficial ou a autoridade competente.",
    ])

    h1(doc, "4. O que \"Aprovado\" significa – e o que não significa")
    t = table(doc, ["Significa", "Não significa"],
              [["O veículo não apresentou achados N2, N3 ou N4 no exame técnico visual realizado naquele momento.",
                "Que a compra está aprovada, que o veículo tem garantia ou que não há vícios ocultos."],
               ["Os itens foram examinados segundo o método e os limites descritos no laudo.",
                "Que o laudo substitui a vistoria oficial de transferência ou uma perícia."]],
              [8.5, 8.5], size=9.5)

    h1(doc, "5. Como registramos")
    bullets(doc, [
        "Cada achado tem código do item, nível, descrição objetiva e, quando aplicável, fotografia.",
        "Valores de medição (espessura de pintura, bateria, pneus) são orientativos e variam por fabricante; a conclusão considera o conjunto de evidências.",
        "As consultas informam fonte, data e hora, e dependem das bases consultadas.",
    ])
    h1(doc, "6. Independência e revisão")
    bullets(doc, [
        "Lojas parceiras, vendedores e intermediários não interferem na classificação. Descontos ou volumes não dependem do resultado.",
        "Se você discordar de um achado, peça uma revisão pelo [WHATSAPP/E-MAIL]. Reexaminamos as evidências e respondemos por escrito.",
        "Se houver erro, emitimos laudo corrigido, com nova versão numerada e referência à anterior.",
    ])
    para(doc, "[NOME DA EMPRESA]  ·  [CIDADE/UF]  ·  [WHATSAPP]  ·  Vigente desde [DD/MM/AAAA]", size=9, color=TEXTO2, align="c", before=12)
    doc.save(OUT / "08-politica-de-classificacao.docx")


# ======================================================================================
# 09 — MENSAGENS DE WHATSAPP
# ======================================================================================
def mensagens():
    doc = novo_doc(aviso=False, rodape_esq="Roteiros de mensagens – uso interno")
    titulo(doc, "Roteiros de mensagens para WhatsApp",
           "Modelos para copiar, adaptar e enviar. Troque tudo que está entre colchetes. Tom profissional e humano, sem promessas.")
    boxed(doc, [
        "**Antes de enviar, confira:**",
        "• Destinatário correto: o laudo vai só para quem contratou (ou para quem ele indicar por escrito).",
        "• Nada de \"compra segura\", \"carro garantido\" ou \"pode fechar\". Quem decide é o cliente; você informa.",
        "• Não envie foto de placa ou documento para terceiros. Em divulgação, sempre com placa oculta e autorização.",
        "• Emojis: no máximo um, e só em mensagens de relacionamento. Em resultado de laudo, nenhum.",
    ], size=9.5, bar=AMBAR)

    grupos = [
        ("1. Agendamento", [
            ("Confirmação de agendamento",
             ["Olá, [NOME]! Aqui é [SEU NOME], da [NOME DA EMPRESA].",
              "Sua vistoria cautelar está confirmada:",
              "Veículo: [MARCA/MODELO] – placa [ABC1D23]",
              "Data e hora: [DD/MM], às [00h00]",
              "Local: [ENDEREÇO]",
              "Serviço: [NÍVEL] – R$ [000] ([FORMA DE PAGAMENTO])",
              "Para a vistoria, precisamos do veículo com documento, chave e autorização do proprietário para a partida do motor"
              "[ e o teste de rodagem]. A vistoria leva em média [00] minutos e o laudo sai em até [00] horas.",
              "Antes, vou te enviar o termo de ciência e o aviso de privacidade para leitura. Qualquer dúvida, é só chamar."]),
            ("Lembrete – véspera",
             ["Oi, [NOME]! Passando para lembrar da vistoria amanhã, [DD/MM], às [00h00], em [LOCAL].",
              "Se possível, peça para o veículo estar com o motor frio (sem ter rodado antes) – isso ajuda na avaliação da partida.",
              "Se precisar remarcar, me avise até [00h00] de hoje, por favor."]),
            ("A caminho",
             ["[NOME], estou a caminho. Previsão de chegada: [00h00]. Te aviso quando chegar."]),
            ("Chegada",
             ["Cheguei ao local, [NOME]. Estou [PONTO DE REFERÊNCIA]. Vou me apresentar ao responsável pelo veículo e iniciar pela "
              "conferência do documento. Te atualizo ao final."]),
        ]),
        ("2. Entrega e explicação do resultado", [
            ("Entrega do laudo",
             ["[NOME], seu laudo está pronto e segue em PDF (laudo nº [0000/AAAA]).",
              "Resultado: [CLASSIFICAÇÃO].",
              "Peço que leia com atenção o resumo dos achados e a seção \"Metodologia e limites\". Posso te explicar cada ponto por "
              "ligação agora ou no horário que preferir."]),
            ("Resultado: APROVADO",
             ["O veículo foi classificado como APROVADO: no exame técnico visual de hoje, encontramos apenas itens conformes ou "
              "observações de desgaste natural (níveis N0 e N1), que estão descritos no laudo.",
              "Isso não é garantia do veículo nem cobre vícios ocultos, e não substitui a vistoria oficial de transferência. "
              "Se for concluir o negócio mais adiante, recomendo repetir as consultas no dia do fechamento."]),
            ("Resultado: APROVADO COM APONTAMENTOS",
             ["O veículo foi classificado como APROVADO COM APONTAMENTOS. Encontramos [00] apontamento(s), listados na página [00]:",
              "[EX.: PORTA DIANTEIRA DIREITA REPINTADA (N2); PNEUS TRASEIROS PRÓXIMOS AO LIMITE (N3 – ATENÇÃO)].",
              "Apontamento não quer dizer que o carro é ruim: quer dizer que existe algo que você precisa conhecer antes de decidir – "
              "e que pode pesar na negociação ou exigir manutenção. Quer que eu explique cada item por ligação?"]),
            ("Resultado: REPROVADO",
             ["[NOME], o veículo foi classificado como REPROVADO, porque encontramos [00] achado(s) de nível crítico (N4):",
              "[DESCRIÇÃO OBJETIVA, EX.: REPARO NO CORPO DA LONGARINA DIANTEIRA ESQUERDA – FOTOS 12 A 15].",
              "Pelo nosso critério, achados desse tipo afetam estrutura, segurança ou identificação do veículo. A decisão é sua; "
              "minha recomendação técnica é não concluir o negócio sem uma avaliação especializada complementar. Estou à disposição "
              "para explicar."]),
            ("Resultado: INCONCLUSIVO",
             ["[NOME], o laudo ficou INCONCLUSIVO. Não foi possível verificar [ITEM DE IDENTIFICAÇÃO] porque [MOTIVO OBJETIVO].",
              "Isso não significa que exista um problema: significa que, sem essa verificação, não é possível concluir. "
              "Para concluir, seria necessário [EX.: APRESENTAR O VEÍCULO LIMPO NA REGIÃO / ACESSO COM ELEVADOR / CONSULTA À AUTORIDADE].",
              "Recomendo não fechar negócio antes de resolver esse ponto. Posso te orientar sobre os próximos passos."]),
        ]),
        ("3. Financeiro", [
            ("Cobrança educada – 1º contato",
             ["Oi, [NOME], tudo bem? Passando para lembrar do pagamento da vistoria do dia [DD/MM] (laudo nº [0000/AAAA]), "
              "no valor de R$ [000]. Segue a chave PIX: [CHAVE]. Se já pagou, desconsidere e me envie o comprovante, por favor. Obrigado!"]),
            ("Cobrança – 2º contato",
             ["[NOME], ainda não localizei o pagamento da vistoria de [DD/MM] (R$ [000]). Consegue verificar para mim? Se preferir, "
              "podemos combinar outra data ou forma de pagamento."]),
        ]),
        ("4. Relacionamento", [
            ("Pedido de avaliação",
             ["[NOME], obrigado por confiar na [NOME DA EMPRESA]. Se o atendimento foi útil, sua avaliação ajuda muito outros "
              "compradores a nos encontrar: [LINK]. Leva um minuto. E se algo puder melhorar, me conte diretamente."]),
            ("Pós-venda (30 dias)",
             ["Oi, [NOME]! Faz um mês da vistoria do [MARCA/MODELO]. Espero que esteja tudo certo. Lembrando que as manutenções "
              "indicadas no laudo ([ITENS]) merecem atenção. Se precisar de uma vistoria para outro veículo, é só chamar."]),
        ]),
        ("5. Lojas e revendas", [
            ("Abordagem inicial a loja",
             ["Olá, [NOME DO GESTOR]. Sou [SEU NOME], vistoriador cautelar em [CIDADE]. Faço vistoria de entrada de estoque e laudo "
              "para comprador, com critério de classificação publicado e laudo em PDF com fotos em até [00] horas.",
              "Posso te mandar um laudo de exemplo e a nossa política de classificação? Trabalho com faturamento mensal para lojas."]),
            ("Follow-up com loja",
             ["[NOME], conseguiu ver o laudo de exemplo? Se fizer sentido, posso fazer uma vistoria de teste em um veículo de entrada "
              "de estoque, na condição da tabela de parceria. Qual o melhor dia para conversarmos 10 minutos?"]),
        ]),
        ("6. Situações delicadas", [
            ("Cliente pergunta \"posso fechar?\"",
             ["Entendo a pergunta, [NOME]. A decisão é sua, e o laudo existe para você decidir com informação. O que posso te dizer "
              "tecnicamente é: [RESUMO DOS ACHADOS]. Se quiser, repasso cada ponto e o que ele pode representar em custo ou risco."]),
            ("Pedido para alterar o laudo",
             ["[NOME], não posso alterar a classificação ou retirar achados: o critério é o mesmo para todos os veículos e está "
              "publicado. Se você identificar algum erro de fato (ex.: dado do veículo, peça trocada), me envie a informação e eu "
              "reexamino as evidências e respondo por escrito."]),
        ]),
    ]
    for gtit, msgs in grupos:
        h1(doc, gtit)
        for mt, linhas in msgs:
            h2(doc, mt)
            boxed(doc, linhas, fill="F7F8FA", size=10)
            spacer(doc, 4)
    doc.save(OUT / "09-mensagens-whatsapp.docx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    for f in (laudo, termo, proposta, contrato, parceria, recibo, privacidade, politica, mensagens):
        f()
    for p in sorted(OUT.glob("*.docx")):
        print("ok", p.relative_to(BASE))


if __name__ == "__main__":
    main()
